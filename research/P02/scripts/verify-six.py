"""Actual artifact checks and stale-source probe, no synthetic reviews."""
import hashlib,json,re,shutil,subprocess,sys,zipfile
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET
from pptx import Presentation
base=Path('/workspace/research/P02');workspace=base/'outputs/official-six'
sys.path.insert(0,'/artifacts/main-skill/scripts')
from pptagent_runtime import core
from pptagent_runtime.config import node_environment
final=json.loads((workspace/'final-report.json').read_text());assert final['complete'] and all(final['checks'].values())
state=core.load_state(workspace)
assert all(r['ok'] and r['review_mode']=='host' for r in state['reviews']['slides'].values()) and len(state['reviews']['slides'])==6
assert state['reviews']['deck']['ok'] and state['reviews']['deck']['review_mode']=='host'
for p in core.slide_files(workspace):assert core.render_current(workspace,state['renders']['slides'][str(p.relative_to(workspace))],core.input_snapshot(workspace,[p]))
answer=workspace/'answer.pptx';assert hashlib.sha256(answer.read_bytes()).hexdigest()==final['artifact_sha256']
prs=Presentation(str(answer));assert len(prs.slides)==6
info=[]
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
class Texts(HTMLParser):
 def __init__(self):super().__init__();self.depth=0;self.texts=[]
 def handle_starttag(self,tag,attrs):
  if tag in ['p','h1','h2','h3','h4','li']:self.depth+=1
 def handle_endtag(self,tag):
  if tag in ['p','h1','h2','h3','h4','li']:self.depth-=1
 def handle_data(self,s):
  if self.depth>0 and s.strip():self.texts.append(s)
compact=lambda s:re.sub(r'\s+','',s)
with zipfile.ZipFile(answer) as z:
 assert z.testzip() is None
 for i,s in enumerate(prs.slides,1):
  xml=ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'));text=''.join(n.text or '' for n in xml.findall('.//a:t',ns))
  parsed=Texts();parsed.feed((workspace/'slides'/f'slide{i:02}.html').read_text(encoding='utf-8'))
  missing=[t for t in parsed.texts if compact(t) not in compact(text)];assert not missing,(i,missing)
  info.append({'slide':i,'shapes':len(s.shapes),'native_text_shapes':sum(bool(sh.has_text_frame and sh.text.strip()) for sh in s.shapes),'picture_shapes':len(xml.findall('.//p:pic',ns)),'all_HTML_text_in_native_a_t':True,'text_items_checked':len(parsed.texts)})
stale=base/'outputs/workflow-probes/stale';shutil.copytree(workspace,stale)
p=stale/'slides/slide01.html';p.write_text(p.read_text(encoding='utf-8').replace('首条链路交付计划','首条链路修订计划'),encoding='utf-8')
args=['/artifacts/main-skill/.venv/bin/python','/artifacts/main-skill/scripts/pptagent.py','--config',str(base/'config-host-review.yaml'),'finalize','--workspace',str(stale)]
r=subprocess.run(args,capture_output=True,text=True,timeout=120);result=json.loads(r.stdout)
assert r.returncode==1 and not result['complete'] and not result['checks']['build_current'] and not result['checks']['slide_reviews_current']
(base/'validation/workflow-stale-source.txt').write_text(r.stdout+'\nSTDERR:\n'+r.stderr,encoding='utf-8')
# This is a rejected copied task; authoritative normal artifact is untouched.
assert core.finalize(workspace,{'delivery':{'mode':'strict'}})['complete']
artifacts=[]
for p in sorted(workspace.rglob('*')):
 if p.is_file():artifacts.append({'path':str(p.relative_to(base)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size})
report={'requirement_id':'RES-P02-01','task_ids':['TASK-RES-P02-01','VERIFY-RES-P02-01'],'verification_state':'completed_candidate_baseline_with_explicit_external_review_failure','source_commit':'833cda553b343be0e486a93b0b57cac962cdd566','normal':{'real_model_authoring':'validation/author-run.json','workflow':'official init→review-slides→record-host-review→build→review-deck→record-host-review→finalize','config':'config-host-review.yaml (official default multimodal/strict)','complete':True,'slides':6,'aspect_ratio':'16:9','PPTX_sha256':final['artifact_sha256'],'native_text_structure':info,'visual_review':'validation/host-slide-review.json; validation/host-deck-review.json; record-host-*.json','rendering':'Chromium HTML; LibreOffice/PyMuPDF exported deck; not PowerPoint acceptance'},'revision':{'change':'footer660→620px on6HTML sources after0.5in margin rejection','evidence':'validation/footer-revision.json; validation/official-six-build-host.json','fresh_render_review_rebuild':'validation/official-six-slides-host-v2.json; validation/official-six-build-host-v2.json','scope':'source HTML edit then full rebuild; not in-place preserving PPTX editing'},'boundary':{'stale_source':{'returncode':r.returncode,'result':result,'evidence':'validation/workflow-stale-source.txt','method':'copy actual complete task then change one HTML; never edit managed evidence'},'zero/one_count':'validation/workflow-probes.json'},'failure':{'visual_interface':'Two original text-mode attempts exit2; actual public response issue.slide="01" violates positive integer validator; preserved without normalizing or passing','visual_evidence':'validation/official-six-review-slides.json; validation/official-six-slides-traced.json; validation/visual-slides-traced-01.json','fallback':'Explicit switch to original default host multimodal review; external JSON integration remains Partial, not passed','build_margin':'first build exit2,0.40in footer; source change fixes it without editing converter','entry_and_gates':'zero input/overwrite/missing renders/incomplete finalize rejected with2/1 and real errors, validation/workflow-probes.json'},'runtime_observation_only':{'resolved_dependency':'pptagent1.1.37, validation/converter-source.json','node_used_by_official_build':subprocess.check_output(['node','--version'],env=node_environment(),text=True).strip(),'storage':'F:/DockerDesktopWSL/disk/docker_data.vhdx, F: bind mounts/cache/tmp','product_runtime_selection':False},'artifact_inventory':artifacts,'completion_boundary':'P02-01 official host workflow reproducible with real DeepSeek HTML authoring, observed source revision and retained failures; no fully automated external-review success, PowerPoint/full object QA, historical tag runtime, or product AC claims. P02-02/03/04/05 remain to map source and capabilities.'}
(base/'validation/verify-official-six.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'state':'passed','complete':True,'slides':6,'native_text_shapes':[s['native_text_shapes'] for s in info],'stale_source_rejected':True,'node_used':report['runtime_observation_only']['node_used_by_official_build']},ensure_ascii=False),flush=True)
