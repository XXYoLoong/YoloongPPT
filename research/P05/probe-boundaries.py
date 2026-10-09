"""Bounded candidate probes; run in existing P05 container, no network/model."""
import dataclasses, hashlib, json, subprocess, sys, traceback, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
sys.path.insert(0,'/app')
from pptx import Presentation
from pptx.util import Inches
from python_backend import smart_layer as sm
from python_backend.template_engine import parse_template,load_theme
from python_backend.pptx_renderer import render_deck_with_template

out=Path('/evidence/outputs/maps-probes');out.mkdir(exist_ok=True)
cases={}
def save(name,value):
 p=out/name;p.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');return p
def run(name,args):
 r=subprocess.run(args,cwd='/app',text=True,capture_output=True,timeout=180)
 log=Path('/evidence/validation')/(name+'.txt');log.write_text(r.stdout+'\nSTDERR:\n'+r.stderr,encoding='utf-8')
 cases[name]={'command':args,'returncode':r.returncode,'log':'validation/'+log.name}
 return r
normal=json.loads(Path('/evidence/outputs/normal/py-generated-deck.json').read_text())
mock12=sm.build_mock_deck('Create a 12-slide technical deck',[],[],[])
cases['page_budget']={'requested':12,'actual':mock12['slideCount'],'language_chinese_prompt':sm.infer_language('中文演示文稿'),'mechanism':'9 predefined pages sliced after 5..12 clamp'}
assert mock12['slideCount']==9
clarify=json.loads(json.dumps(normal));clarify['slides'][1]['bullets']=['X'*100+str(i) for i in range(8)]
sm.clarify_deck(clarify)
cases['clarify']={'before_items':8,'before_chars_each':101,'after_items':len(clarify['slides'][1]['bullets']),'after_chars':[len(s) for s in clarify['slides'][1]['bullets']],'reported_assumptions_added':clarify['assumptions']!=normal['assumptions']}
assert cases['clarify']['after_items']==6 and cases['clarify']['after_chars']==[80]*6
theme={};sm.apply_requested_theme(theme,'make it tech');resolved=load_theme(theme['theme'])
cases['theme_fallback']={'requested_theme':theme['theme'],'resolved_theme':resolved['name'],'reason':'tech-modern has no built-in asset; warning emitted by load_theme'}
save('mock12.json',mock12);save('clarify.json',clarify)

fixture=out/'brand-template.pptx';prs=Presentation()
for i in range(2):
 s=prs.slides.add_slide(prs.slide_layouts[0]);s.shapes.title.text='OLD-TEMPLATE-SENTINEL-'+str(i)
prs.save(fixture)
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with zipfile.ZipFile(fixture) as z: parts={n:z.read(n) for n in z.namelist()}
root=ET.fromstring(parts['ppt/theme/theme1.xml'])
root.find('.//a:clrScheme/a:accent1/a:srgbClr',ns).set('val','123456')
root.find('.//a:majorFont/a:latin',ns).set('typeface','Liberation Serif')
parts['ppt/theme/theme1.xml']=ET.tostring(root,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(fixture,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
config=parse_template(fixture)
parsed=dataclasses.asdict(config);parsed['source_path']=str(config.source_path)
save('template-parsed.json',parsed)
cases['template_theme']={'fixture_accent1':'123456','fixture_heading':'Liberation Serif','parsed_primary':config.theme_colors.primary,'parsed_heading':config.font_heading,'layouts':len(config.layouts),'placeholders':sum(len(l.placeholders) for l in config.layouts),'mapping_keys':list(config.layout_mapping),'scope':'default python-pptx master/layouts plus custom theme relationship; not commercial template'}
assert config.theme_colors.primary!='123456' and config.font_heading!='Liberation Serif'
r=run('template-cli',['/scratch/venv/bin/python','/app/auto-ppt','generate','--mock','--prompt','Create a 5-slide deck','--template',str(fixture),'--output-dir',str(out/'template-cli')])
cases['template-cli']['input_existing_slides']=2
cases['template-cli']['output_exists']=(out/'template-cli/py-generated-deck.pptx').exists()
assert r.returncode==1 and not cases['template-cli']['output_exists']

deck=json.loads(json.dumps(normal));chart=next(s for s in deck['slides'] if s['layout']=='chart')
table=json.loads(json.dumps(chart));table.update(layout='table',title='Native table probe',table={'columns':['Item','Value'],'rows':[['Alpha','1'],['Beta','2']]})
deck['slides']=[chart,table];deck['slideCount']=2
for i,s in enumerate(deck['slides'],1):s['page']=i;s['speakerNotes']=['P05-OBJECT-PROBE-NOTE']
deck['needsSpeakerNotes']=True
inp=save('native-objects.json',deck)
r=run('native-objects',['node','/app/generate-ppt.js',str(inp),str(out/'native-objects.pptx'),'--native-charts']);assert r.returncode==0
pydeck=json.loads(json.dumps(deck));kpi=json.loads(json.dumps(table));kpi.update(page=3,layout='kpi',title='Unsupported template KPI probe',bullets=[],kpis=[{'label':'KPI-SENTINEL','value':'99%'}])
pydeck['slides'].append(kpi);pydeck['slideCount']=3
# Direct renderer can be called outside planner; capture its zero-padding behavior.
pydeck['slides'][0]['chart']['categories']=['A','B','C'];pydeck['slides'][0]['chart']['series']=[{'name':'Short series','data':[7]}]
try:
 render_deck_with_template(pydeck,out/'template-objects.json',out/'template-objects.pptx',config)
 raise AssertionError('template failure no longer reproduced')
except (KeyError,ValueError) as e:
 Path('/evidence/validation/template-direct.txt').write_text(traceback.format_exc(),encoding='utf-8')
 cases['template_objects']={'exception_type':type(e).__name__,'message':str(e),'log':'validation/template-direct.txt','output_exists':(out/'template-objects.pptx').exists(),'reason':'slide deletion uses .get("r:id") without expanded XML namespace; rId=None and drop_rel raises KeyError'}
 assert not cases['template_objects']['output_exists']
empty=out/'empty-template.pptx';Presentation().save(empty);empty_config=parse_template(empty)
render_deck_with_template(pydeck,out/'empty-template-objects.json',out/'empty-template-objects.pptx',empty_config)
p=Presentation(str(out/'empty-template-objects.pptx'))
cases['empty_template_objects']={'chart_values':list(next(sh.chart for sh in p.slides[0].shapes if sh.has_chart).series[0].values),'kpi_text_present':any('KPI-SENTINEL' in sh.text for sh in p.slides[2].shapes if sh.has_text_frame),'kpi_layout_name':p.slides[2].slide_layout.name,'scope':'separate zero-slide template direct renderer input; no patch/workaround to upstream; bypasses planner validation'}
assert cases['empty_template_objects']['chart_values']==[7.0,0.0,0.0] and not cases['empty_template_objects']['kpi_text_present']
bad=out/'corrupt.pptx';bad.write_text('not a zip',encoding='utf-8')
for name,path in [('missing_template',out/'absent.pptx'),('corrupt_template',bad),('wrong_extension',out/'mock12.json')]:
 try:parse_template(path);raise AssertionError('expected failure')
 except (FileNotFoundError,RuntimeError,ValueError) as e:cases[name]={'exception_type':type(e).__name__,'message':str(e),'expected_rejection':True}
invalid=json.loads(json.dumps(deck));invalid['slides'][0]['layout']='not-supported'
inp=save('invalid-layout.json',invalid)
r=run('invalid-layout',['node','/app/generate-ppt.js',str(inp),str(out/'invalid-layout.pptx')]);assert r.returncode!=0 and not (out/'invalid-layout.pptx').exists()
inventory=[]
for f in sorted(out.rglob('*.pptx')):
 if f==bad:continue
 with zipfile.ZipFile(f) as z:
  assert z.testzip() is None
  charts=[n for n in z.namelist() if n.startswith('ppt/charts/chart') and n.endswith('.xml')]
 p=Presentation(str(f));inventory.append({'path':str(f.relative_to(Path('/evidence'))),'sha256':hashlib.sha256(f.read_bytes()).hexdigest(),'slides':len(p.slides),'native_chart_parts':len(charts),'native_tables':sum(sh.has_table for s in p.slides for sh in s.shapes),'notes_probe_present':any('P05-OBJECT-PROBE-NOTE' in s.notes_slide.notes_text_frame.text for s in p.slides),'zip_integrity':True})
js=next(a for a in inventory if a['path'].endswith('/native-objects.pptx'));assert js['native_chart_parts']==1 and js['native_tables']==1 and js['notes_probe_present']
report={'requirement_ids':['RES-P05-03','RES-P05-04','RES-P05-05'],'source_commit':'5ae0670747885c464aa8063329a902d80a251877','mode':'candidate mock + actual renderer + source inspection; no model/network','cases':cases,'artifacts':inventory,'reused_normal_qa':'validation/verify-res-p05-01-02.json','scope':'OOXML structure/readability tested; no new visual/PowerPoint/full object editing acceptance'}
Path('/evidence/validation/maps-probes.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False),flush=True)
