"""Run official P05 mock CLI and bounded research probes inside its container."""
import hashlib, json, os, subprocess, sys, zipfile
from pathlib import Path
from pptx import Presentation

root=Path('/app')
evidence=Path('/evidence')
outputs=evidence/'outputs'
outputs.mkdir(exist_ok=True)
python='/scratch/venv/bin/python'
cases={}
commands={
 'normal':['generate','--mock','--prompt','Create an 8-slide AI strategy deck for executives','--source','examples/inputs/sample-source-brief.md','--output-dir',str(outputs/'normal')],
 'boundary':['generate','--mock','--prompt','Create a 1-slide AI strategy deck','--output-dir',str(outputs/'boundary')],
 'failure':['generate','--mock','--prompt','Create an 8-slide AI strategy deck','--source','examples/inputs/does-not-exist.md','--output-dir',str(outputs/'failure')],
 'revision':['revise','--mock','--deck',str(outputs/'normal/py-generated-deck.json'),'--prompt','Compress to 6 slides, make it more conclusion-driven','--output-dir',str(outputs/'revision')],
 'visual_qa':['qa-visual',str(outputs/'normal/py-generated-deck.pptx'),'--strict','--output-dir',str(outputs/'normal-qa')],
}
for name,args in commands.items():
    command=[python,str(root/'auto-ppt'),*args]
    result=subprocess.run(command,cwd=root,text=True,capture_output=True,timeout=180)
    (evidence/'validation'/f'{name}.txt').write_text(result.stdout+'\nSTDERR:\n'+result.stderr,encoding='utf-8')
    cases[name]={'command':command,'returncode':result.returncode,'log':f'validation/{name}.txt'}
    print(name,result.returncode,flush=True)
artifacts=[]
for path in sorted(outputs.rglob('*.pptx')):
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        charts=sum(n.startswith('ppt/charts/chart') and n.endswith('.xml') for n in z.namelist())
    prs=Presentation(str(path))
    shapes=[len(s.shapes) for s in prs.slides]
    texts=[sum(bool(s.has_text_frame and s.text.strip()) for s in slide.shapes) for slide in prs.slides]
    artifacts.append({'path':str(path.relative_to(evidence)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size,'slides':len(prs.slides),'shape_counts':shapes,'text_shape_counts':texts,'native_chart_parts':charts,'pptx_readable':True,'zip_integrity':True})
report={'requirement_ids':['RES-P05-01','RES-P05-02'],'source_commit':'5ae0670747885c464aa8063329a902d80a251877','mode':'official --mock; not real AI generation','network':'disconnected for these probes','cases':cases,'artifacts':artifacts}
(evidence/'validation/baseline-run.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
assert cases['normal']['returncode']==0 and cases['boundary']['returncode']==0 and cases['failure']['returncode']!=0 and cases['revision']['returncode']==0,report
print(json.dumps(report,ensure_ascii=False),flush=True)
