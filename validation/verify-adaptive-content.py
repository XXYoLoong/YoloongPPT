"""Real component PPTX/render evidence for complete literal adaptive capacity.

Explicitly authored, approved fixture text; no model call and no AI/AC claim.
"""
import copy
import hashlib
import json
import traceback
import sys
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from yoloongppt.adaptive_content import adapt, TABLE_PAGINATION_NOTICE
from yoloongppt.adaptive_layout import fit_elements
from yoloongppt.artifacts import entity
from yoloongppt.atomic import AtomicRegistry
from yoloongppt.context import normalize
from yoloongppt.errors import TaskError, trace_id
from yoloongppt.narrative import prepare, finalize
from yoloongppt.planning import DEFAULT_STYLE, compile_deck
from yoloongppt.quality import check as quality_check
from yoloongppt.rendering import render
from yoloongppt.schemas import SchemaRegistry, ROOT
from yoloongppt.writer import execute

out = Path('/runtime/adaptive-content-validation') / datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')
out.mkdir(parents=True, exist_ok=False)
schemas = SchemaRegistry(); registry = AtomicRegistry(schemas); checks = []; reports = []


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def check(name, condition, observed=None):
    if not condition:
        raise AssertionError((name, observed))
    checks.append({'name': name, 'passed': True, 'observed': observed})
    save(out/'checks-progress.json', checks)


def slide(title, body, table=None):
    return {'kind': 'table' if table else 'text', 'title': title, 'body': body,
            'goal': '完整展示已批准的字面夹具', 'rationale': '按原文顺序保留，容量调整不改变事实',
            'notes': '原备注只用于解释，核心正文必须在页面可见。', 'evidence_refs': [],
            'table': table, 'chart': None}


def fixture(name, slides, kind='soft', keys=None):
    samples = json.loads(Path('/runtime/visual-validation/node-samples.json').read_text('utf-8'))
    task = copy.deepcopy(next(s['input']['task'] for s in samples if s['name'] == 'DEC-022_normal_3'))
    task['task_id'] = entity('task'); task['template_ref'] = None
    task['runtime_preferences'] = {'page_count': len(slides), 'adaptive_policy': {'table_key_columns': keys or {}}}
    task['constraints'] = {'hard': [], 'soft': [], 'defaults': [], 'conflicts': []}
    if kind == 'hard':
        task['constraints']['hard'].append({'constraint_id': 'fixture.exact-pages', 'field': 'page_count',
            'value': len(slides), 'source': {'kind': 'user_instruction', 'reference': 'explicit component fixture'}})
    elif kind == 'default':
        task['runtime_preferences'].pop('page_count')
        task['constraints']['defaults'].append({'constraint_id': 'fixture.default-pages', 'field': 'page_count',
            'value': len(slides), 'source': {'kind': 'system_default', 'reference': 'explicit component fixture'}, 'used': False})
    task.pop('raw_request', None)
    sid = entity('source'); eid = entity('evidence')
    raw = '\n'.join(s['title'] + '\n' + '\n'.join(s['body']) + '\n' + s['notes'] + '\n' +
        ('\n'.join(['\t'.join(s['table']['columns']), *['\t'.join(r) for r in s['table']['rows']]]) + '\n' + s['table']['units'] if s['table'] else '') for s in slides)
    path = out/(name+'-approved-source.txt'); path.write_text(raw, encoding='utf-8')
    task['sources'] = [{'source_id': sid, 'kind': 'text', 'locator': str(path), 'metadata': {'source_role': 'content_material'}}]
    evidence = {'evidence_id': eid, 'source_id': sid, 'source_hash': hashlib.sha256(path.read_bytes()).hexdigest(),
        'parser_version': 'explicit-literal-fixture-1', 'segment_index': 0,
        'anchor': {'source_id': sid, 'asset_id': None, 'native_locator': str(path), 'locator_status': 'located', 'reason': None,
            'locator': {'kind': 'text_span', 'segment_index': 0, 'start_char': 0, 'end_char': len(raw), 'offset_unit': 'unicode_code_points'}},
        'raw_text': raw, 'content': raw, 'raw_asset_refs': [], 'parse_note': 'Explicit approved component fixture; no semantic inference.',
        'data': None, 'asset_ref': None, 'confidence': None, 'metadata': {'fixture_scope': 'manual component evidence'}}
    schemas.validate('source-evidence.schema.json', evidence)
    sources = {'evidence': [evidence], 'approved_texts': {eid: raw},
        'fact_boundary': {'status': 'ready', 'assumption_values': {}, 'actions': []}}
    for page in slides:
        page['evidence_refs'] = [eid]
        for aid in page.get('assumption_refs', []):
            sources['fact_boundary']['assumption_values'][aid] = {'raw_text': '目标转化率为12.5%'}
    proposal = {'title': name, 'main_takeaway': '全部原文在可读页面展示',
        'storyline': {'pattern': 'linear', 'rationale': '保留批准原文順序'}, 'slides': slides}
    context = normalize(task, schemas, trace_id())['context']
    check(name+':normalized_constraints', not context['issues'] and not context['constraints']['conflicts'], context['issues'])
    return task, sources, context, proposal


def boundary(name, operation, expected):
    try:
        operation()
    except TaskError as error:
        save(out/(name+'-failure.json'), error.result(trace_id()))
        check(name, error.code == expected, error.code)
    else:
        raise AssertionError(name+' did not fail explicitly')


def run(name, slides, kind='soft', keys=None):
    task, sources, context, proposal = fixture(name, copy.deepcopy(slides), kind, keys)
    original = copy.deepcopy(proposal)
    packet = prepare(task, sources, context, schemas)
    check(name+':hard_soft_default_source', packet['content_policy']['adaptive_policy']['page_budget']['kind'] == kind)
    final, narrative = finalize(proposal, packet, sources, schemas)
    adaptation = narrative['section_budget']['adaptive_content']
    case = out/name; case.mkdir()
    for filename, value in [('task.json',task),('source-result.json',sources),('context.json',context),
        ('proposal.json',proposal),('narrative-context.json',packet),('narrative-result.json',narrative),
        ('final-proposal.json',final),('adaptive-content-result.json',adaptation)]:
        save(case/filename, value)
    check(name+':input_unmutated', proposal == original)
    check(name+':no_core_in_notes', all(p['notes'] == original['slides'][m['original_page']]['notes']
        for m in adaptation['page_mappings'] for e in m['outputs'] for p in [final['slides'][e['output_page']]]))
    check(name+':body_character_and_offset_conservation', all(''.join(seg['text'] for e in m['outputs'] for seg in e['body_segments'] if seg['body_item']==i)==value
        for m in adaptation['page_mappings'] for i,value in enumerate(original['slides'][m['original_page']]['body'])))
    check(name+':refs_and_title_exact', all(final['slides'][e['output_page']]['title'] == original['slides'][m['original_page']]['title'] and
        final['slides'][e['output_page']]['evidence_refs'] == original['slides'][m['original_page']]['evidence_refs'] and
        final['slides'][e['output_page']].get('assumption_refs',[]) == original['slides'][m['original_page']].get('assumption_refs',[])
        for m in adaptation['page_mappings'] for e in m['outputs']))
    check(name+':actual_narrative_allocation', narrative['total'] == len(final['slides']) and
        sum(s['pages'] for s in narrative['section_budget']['sections']) == len(final['slides']))
    check(name+':all_original_narrative_nodes_consumed', {t['node_id'] for t in narrative['decision_traces']} == {f'DEC-{i:03d}' for i in range(6,22)})
    deck, execution = compile_deck(final, DEFAULT_STYLE, context['trace_id'], registry=registry)
    save(case/'deck-spec.json',deck); save(case/'execution-plan.json',execution)
    object_map, result = execute(deck,execution,case/'deck.pptx',schemas)
    save(case/'object-map.json',object_map); save(case/'execution-result.json',result)
    native = Presentation(case/'deck.pptx')
    check(name+':real_editable_pptx_pages', len(native.slides)==len(final['slides']))
    check(name+':readable_font_preserved', all(e['font_size']>=16 for p in deck['slides'] for e in p['elements']))
    for mapping in adaptation['page_mappings']:
        for entry in mapping['outputs']:
            page = deck['slides'][entry['output_page']]; actual = native.slides[entry['output_page']]
            for e in page['elements']:
                if e['type']=='text':
                    text = next(s for s in actual.shapes if s.name==e['object_id']).text
                    check(name+':native_text:'+e['object_id'], text=='\n'.join(e['text']))
            if entry['table_coordinates']:
                table = next(s.table for s in actual.shapes if s.has_table)
                data = final['slides'][entry['output_page']]['table']
                check(name+':native_cells:'+str(entry['output_page']), [[c.text for c in row.cells] for row in table.rows]==[data['columns'],*data['rows']])
                coords = entry['table_coordinates']; source = original['slides'][mapping['original_page']]['table']
                check(name+':original_cells:'+str(entry['output_page']), data['columns']==[source['columns'][c] for c in coords['columns']] and
                    data['rows']==[[source['rows'][r][c] for c in coords['columns']] for r in coords['rows']])
                check(name+':native_units_visible:'+str(entry['output_page']), any(data['units'] in s.text for s in actual.shapes if s.has_text_frame))
    renders = render(case/'deck.pptx',deck,case); save(case/'render-report.json',renders)
    check(name+':actual_pdf_png_all_pages',sum(a['type']=='png' for a in renders['artifacts'])==len(final['slides']) and any(a['type']=='pdf' for a in renders['artifacts']))
    check(name+':all_render_hashes',all(hashlib.sha256((case/a['path']).read_bytes()).hexdigest()==a['hash'] for a in renders['artifacts']))
    qa = quality_check(deck,case/'deck.pptx',object_map,sources,renders,case);save(case/'quality-report.json',qa)
    check(name+':no_render_or_geometry_p0',qa['p0_issue_count']==0,[{'code':i['code'],'detail':i['detail']} for i in qa['issues']])
    reports.append({'name':name,'input_pages':len(original['slides']),'output_pages':len(final['slides']),
        'budget_kind':kind,'p0':qa['p0_issue_count'],'warnings':[i['code'] for i in qa['issues'] if i['severity']!='P0']})
    return task,sources,context,proposal,packet,final,adaptation,deck


def verify_pagination_notice_only():
    """Incremental checks on the real saved model seed; no render/model replay."""
    parent = Path('/runtime/runs/ad03ceec-a393-48a3-840e-4e2f52c50613')
    proposal = json.loads((parent/'model-response.json').read_text('utf-8'))
    context = json.loads((parent/'narrative-context.json').read_text('utf-8'))
    policy = context['content_policy']['adaptive_policy']; original = copy.deepcopy(proposal)
    adapted, report = adapt(proposal, policy, schemas=schemas)
    save(out/'pagination-notice-input.json', original)
    save(out/'pagination-notice-output.json', adapted)
    save(out/'pagination-notice-report.json', report)
    pages = [p for p in adapted['slides'] if p['kind'] == 'table']
    source_table = next(p for p in original['slides'] if p['kind'] == 'table')
    mapping = next(m for m in report['page_mappings'] if m['original_table_sha256'])
    check('notice:real_fourteen_rows_split', len(source_table['table']['rows']) == 14 and len(pages)>1)
    check('notice:all_actual_tiles_visible_non_numeric_metadata', all(TABLE_PAGINATION_NOTICE in p['body'] for p in pages)
          and not any(c.isdigit() for c in TABLE_PAGINATION_NOTICE)
          and all(e['visible_pagination_notice']['business_fact'] is False and e['visible_pagination_notice']['source_kind']=='execution_metadata'
                  for e in mapping['outputs'] if e['table_coordinates']))
    check('notice:original_title_characters_untouched', all(p['title']==source_table['title'] for p in pages))
    check('notice:all_original_cells_in_exact_order', [row for p in pages for row in p['table']['rows']]==source_table['table']['rows'])
    check('notice:references_units_and_notes_preserved', all(p['evidence_refs']==source_table['evidence_refs'] and
          p['notes']==source_table['notes'] and p['table']['units']==source_table['table']['units'] for p in pages))
    check('notice:seed_not_mutated_and_explicit_mapping',proposal==original and len(report['execution_metadata_labels'])==1 and
          report['execution_metadata_labels'][0]['tile_count']==len(pages))
    from yoloongppt.adaptive_layout import content_geometry, fits_body
    check('notice:capacity_reserved_with_complete_original_text',all(fits_body(p['body'],content_geometry(p)) for p in adapted['slides']))
    single = copy.deepcopy(source_table); single['table']['rows']=single['table']['rows'][:1]
    single_proposal={**copy.deepcopy(original),'slides':[single]}
    single_policy={'page_budget':{'kind':'hard','target':1,'minimum':1,'maximum':1},'table_key_columns':{}}
    short,short_report=adapt(single_proposal,single_policy,schemas=schemas)
    check('notice:single_tile_unmodified', len(short['slides'])==1 and short_report['execution_metadata_labels']==[] and
          TABLE_PAGINATION_NOTICE not in short['slides'][0]['body'] and short['slides'][0]['title']==single['title'])
    comparison=adapt({**copy.deepcopy(original),'slides':[single]},single_policy,schemas=schemas)[0]
    check('notice:single_page_idempotent_and_original_cells', comparison==short and short['slides'][0]['table']==single['table'])
    hard_policy={'page_budget':{'kind':'hard','target':len(proposal['slides']),
        'minimum':len(proposal['slides']),'maximum':len(proposal['slides'])},'table_key_columns':policy['table_key_columns']}
    boundary('notice:hard_budget_still_blocks',lambda:adapt(proposal,hard_policy,schemas=schemas),'ADAPTIVE_PAGE_BUDGET_EXCEEDED')
    source_hashes={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [ROOT/'src/yoloongppt/adaptive_content.py',
                  ROOT/'contracts/adaptive-content-result.schema.json',ROOT/'validation/verify-adaptive-content.py']}
    result={'ok':True,'check_count':len(checks),'checks':checks,'source_hashes':source_hashes,
        'parent_real_run':'ad03ceec-a393-48a3-840e-4e2f52c50613','seed_sha256':hashlib.sha256((parent/'model-response.json').read_bytes()).hexdigest(),
        'input_pages':len(proposal['slides']),'output_pages':len(adapted['slides']),
        'table_tile_row_counts':[len(p['table']['rows']) for p in pages],
        'scope':'New incremental pagination-notice contract/capacity checks on real saved model seed; no model/render/native replay and no repeat of the historical 161 checks.'}
    save(out/'pagination-notice-checks.json', result)
    print(json.dumps({'ok':True,'checks':len(checks),'output_pages':len(adapted['slides']),'rows':result['table_tile_row_counts'],'artifact_root':str(out)}))


if '--pagination-notice-only' in sys.argv:
    try:
        verify_pagination_notice_only()
    except Exception as error:
        save(out/'pagination-notice-failure.json',{'type':type(error).__name__,'error':str(error),
            'structured':error.result(trace_id()) if isinstance(error,TaskError) else None,'traceback':traceback.format_exc()})
        raise
    raise SystemExit(0)


try:
    aid = entity('assumption')
    marker = '【假设】目标转化率为12.5%，这是测试假设而非实际统计。'
    blocks = [marker] + [f'语义块{n}：完整正文保留数字3.5 kg、12.5%及中文 English units。'+
        '该段包含明确句界，跨页时不删除任何字符或单位。'*3 for n in range(1,11)]
    sentence_block = ''.join(f'句子{n}记录原文3.5 kg与12.5%，保留英文 Values and units。' for n in range(1,13))
    body_slide = slide('完整正文与假设标记', blocks+[sentence_block]);body_slide['assumption_refs']=[aid]
    body = run('normal-complete-body',[body_slide])
    check('body:actual_expansion',body[6]['output_pages']>1)
    check('body:visible_assumption_every_continuation',all(any(marker in t for t in p['body']) for p in body[5]['slides']))
    mixed_title = ('复杂标题 Mixed English Title 保留数字3.5 kg与12.5%及单位，完整展示中英文字形。'*3)
    title = run('normal-mixed-long-title',[slide(mixed_title,['批准原文保持完整。标题允许多行，正文不移入备注。'])],kind='default')
    check('title:actual_title_bounds_expansion',any(e['role']=='title' and e['bounds'][3]>1.25 for p in title[7]['slides'] for e in p['elements']))
    table_one = {'columns':['关键编号','月份','金额（万元）','增长率（%）'],
        'rows':[[f'ID{r:02d}',f'月{r:02d}',f'{r}.5',f'{r}%'] for r in range(1,27)],'units':'万元；%'}
    table_two = {'columns':['明确关键列']+[f'字段{c:02d}' for c in range(1,25)],
        'rows':[[f'K{r:02d}']+[f'原单元格{r:02d}列{c:02d}完整文字与单位' for c in range(1,25)] for r in range(1,3)],'units':'原文字与单位'}
    tables = run('normal-large-tables',[slide('纵向大表全部行保留',['表格按原行顺序展示。'],table_one),
        slide('多列表格明确重复关键列',['关键列由调用方明确指定。'],table_two)],keys={'1':[0]})
    first,second=tables[6]['page_mappings']
    check('table:vertical_rows_continue',len(first['outputs'])>1)
    check('table:horizontal_columns_continue',len({tuple(e['table_coordinates']['columns']) for e in second['outputs'] if e['table_coordinates']})>1)
    check('table:explicit_keys_repeated',all(0 in e['table_coordinates']['columns'] and e['table_coordinates']['repeated_key_columns']==[0] for e in second['outputs'] if e['table_coordinates']))
    hard = run('boundary-hard-short',[slide('硬页预算可容纳原文',['完整短正文与数字3.5 kg保留。'])],kind='hard')
    check('hard:exact_pages',hard[6]['input_pages']==hard[6]['output_pages']==1)
    hard_task,hard_sources,hard_context,hard_proposal=fixture('boundary-hard-overflow',copy.deepcopy([body_slide]),'hard')
    hard_packet=prepare(hard_task,hard_sources,hard_context,schemas)
    boundary('hard_overflow',lambda:finalize(hard_proposal,hard_packet,hard_sources,schemas),'ADAPTIVE_PAGE_BUDGET_EXCEEDED')
    boundary('hard_cannot_be_changed_to_soft',lambda:finalize(hard_proposal,hard_packet,hard_sources,schemas,
        adaptive_policy={'page_budget':{'kind':'soft','target':1,'minimum':1,'maximum':30}}),'ADAPTIVE_HARD_BUDGET_CHANGED')
    boundary('unknown_horizontal_key',lambda:adapt(tables[3],{**tables[4]['content_policy']['adaptive_policy'],'table_key_columns':{'1':[99]}},schemas=schemas),'ADAPTIVE_TABLE_KEY_POLICY_INVALID')
    boundary('missing_horizontal_key',lambda:adapt(tables[3],{**tables[4]['content_policy']['adaptive_policy'],'table_key_columns':{}},schemas=schemas),'ADAPTIVE_TABLE_KEY_COLUMNS_REQUIRED')
    indivisible=copy.deepcopy(hard[3]);indivisible['slides'][0]['body']=['无明确句界的超长内容'*300]
    boundary('indivisible_sentence',lambda:adapt(indivisible,{'page_budget':{'kind':'soft','target':1,'minimum':1,'maximum':30},'table_key_columns':{}},schemas=schemas),'ADAPTIVE_INDIVISIBLE_CONTENT')
    boundary('locked_template_geometry',lambda:adapt(hard[3],templates={'geometry_locked':True},schemas=schemas),'ADAPTIVE_TEMPLATE_GEOMETRY_LOCKED')
    body_without_marker=copy.deepcopy(body[3]);body_without_marker['slides'][0]['body']=body_without_marker['slides'][0]['body'][1:]
    boundary('assumption_marker_missing',lambda:adapt(body_without_marker,body[4]['content_policy']['adaptive_policy'],schemas=schemas),'ADAPTIVE_ASSUMPTION_VISIBILITY_REQUIRED')
    gallery = copy.deepcopy(body[3]); page = gallery['slides'][0]
    page['kind'] = 'image'; page.pop('assumption_refs', None)
    page['images'] = [{'asset_id':entity('asset'),'role':'information','alt_text':'显式图库声明',
                       'fit':'contain','required':True} for _ in range(2)]
    gallery_before = copy.deepcopy(gallery)
    gallery_adapted,gallery_report = adapt(gallery,body[4]['content_policy']['adaptive_policy'],schemas=schemas)
    save(out/'gallery-continuation-boundary.json',{'input':gallery,'output':gallery_adapted,'report':gallery_report,
        'scope':'Declaration preservation boundary only; images unresolved, no gallery render/AI acceptance claim.'})
    retained = [p for p in gallery_adapted['slides'] if p['kind']=='image']
    check('gallery:declarations_preserved_once_without_polluting_text',gallery==gallery_before and len(retained)==1 and
          retained[0]['images']==page['images'] and any(p['kind']=='text' for p in gallery_adapted['slides']) and
          all('images' not in p and 'image' not in p and 'diagram' not in p for p in gallery_adapted['slides'] if p['kind']=='text'))
    # Schema and fit_elements interfaces are genuinely consumed, not draft files.
    elements=title[7]['slides'][0]['elements']
    fitted,style,witness=fit_elements(title[5]['slides'][0],elements,DEFAULT_STYLE)
    schemas.validate('adaptive-layout-result.schema.json',witness)
    check('layout:contract_and_literal_elements_preserved',all(a.get('text')==b.get('text') and a['font_size']==b['font_size'] for a,b in zip(elements,fitted)))
    hashes={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [
        ROOT/'src/yoloongppt/adaptive_content.py',ROOT/'src/yoloongppt/adaptive_layout.py',ROOT/'src/yoloongppt/narrative.py',
        ROOT/'src/yoloongppt/layouts.py',ROOT/'src/yoloongppt/planning.py',ROOT/'src/yoloongppt/writer.py',ROOT/'src/yoloongppt/quality.py',ROOT/'src/yoloongppt/rendering.py']}
    save(out/'adaptive-content-verification.json',{'ok':True,'check_count':len(checks),'checks':checks,'cases':reports,
        'source_hashes':hashes,'schema_hashes':schemas.versions,'scope':'Explicit approved literal fixtures → normalize/DEC006–021 → adaptive → native atomic PPTX → real LibreOffice/PDF/PNG → deterministic QA; zero model calls; full AI/system AC not claimed.'})
    print(json.dumps({'ok':True,'checks':len(checks),'cases':reports,'artifact_root':str(out)}))
except Exception as error:
    save(out/'validation-failure.json',{'type':type(error).__name__,'error':str(error),
        'structured':error.result(trace_id()) if isinstance(error,TaskError) else None,'traceback':traceback.format_exc(),'completed_checks':len(checks),'cases':reports})
    print(json.dumps({'ok':False,'completed_checks':len(checks),'artifact_root':str(out),'error_type':type(error).__name__}))
    raise
