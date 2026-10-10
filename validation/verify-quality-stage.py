"""Bounded DEC-039/040 and actual generated PPTX QA verification; no model calls."""
import copy,json,hashlib
from pathlib import Path
from zipfile import ZipFile
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches,Pt
from PIL import Image
from yoloongppt.writer import image_shape,diagram_shape,connector_shape,text_shape
from yoloongppt.artifacts import entity
from yoloongppt.quality import check
from yoloongppt.quality_rules import audit,package_audit,contrast,issue
from yoloongppt.quality_decisions import decide,run_node
from yoloongppt.revision_plans import build_plan,validate_plan,history_input
from yoloongppt.schemas import SchemaRegistry
from yoloongppt.errors import TaskError

root=Path('/workspace/validation/parallel-artifacts/generated');out=Path('/runtime/quality-stage-validation');out.mkdir(exist_ok=True)
load=lambda n:json.loads((root/n).read_text('utf-8'))
deck=load('deck-spec.json');mapping=load('object-map.json');sources=load('approved-source-result.json');renders=load('render-report.json');schemas=SchemaRegistry();checks=[];cases=[]
def test(name,condition,observed=None):
 checks.append({'name':name,'passed':bool(condition),'observation':observed});assert condition,name

def error(name,fn,code):
 try:fn()
 except TaskError as e:test(name,e.code==code,{'code':e.code});return
 raise AssertionError(name+' did not fail')
quality=check(deck,root/'deck.pptx',mapping,sources,renders,root)
(out/'actual-quality-report.json').write_text(json.dumps(quality,ensure_ascii=False,indent=2)+'\n','utf-8')
test('real_8_page_pptx_checked',len(deck['slides'])==8 and len(quality['checks'])>=10,{'slides':8,'checks':len(quality['checks']),'p0':quality['p0_issue_count'],'issues':[i['code'] for i in quality['issues']]})
test('quality_issues_explainable',all({'object_id','requirement_id','measurement','threshold','action'}<=set(i) for i in quality['issues']))
test('coverage_explicit_no_false_powerpoint',quality['coverage']['powerpoint_compatibility']=='not_executed' and 'partial' in quality['coverage']['accessibility'])
base={'quality_report':quality,'deck':deck,'object_map':mapping,'parent_run_id':load('result.json')['run_id']}
r=decide(base,schemas);cases.append({'name':'actual8_partial','input':base,'output':r});test('actual8_never_full_final',r['completion']['state'] in {'warning','fail'})
clean={'issues':[],'coverage':{c:'complete' for c in ['structure','editability','geometry','fact_numeric','visual_semantic','fact_semantic','accessibility','global_consistency','powerpoint_compatibility']},'p0_issue_count':0,'acceptance':{'AC-001':'passed','GOV-008':'passed'}}
# Synthetic node conditions validate decision behavior; they do not assert system acceptance.
for name,report,expected in [('synthetic_final',clean,'final'),('partial_warning',{**clean,'coverage':{'structure':'partial'}},'warning'),('nonblocking_warning',{**clean,'issues':[issue('ADVICE',deck['slides'][0]['slide_id'],'Unknown optimization','QA-020','P2')]},'warning')]:
 d={'quality_report':report,'deck':deck,'object_map':mapping,'prior_traces':[{'node_id':f'DEC-{i:03d}','error':None,'output':{'synthetic_unit_input':True}} for i in range(1,39)]};result=decide(d,schemas);cases.append({'name':name,'synthetic_acceptance_inputs':True,'input':d,'output':result});test(name,result['completion']['state']==expected)
 for node in ['DEC-039','DEC-040']:test(name+'_'+node,run_node(d,node,schemas)['node_id']==node)
slide=deck['slides'][0];target=next(e for e in slide['elements'] if e['type']=='text' and e['role']=='title');oid=target['object_id'];sid=slide['slide_id']
contrastissue=issue('TEXT_CONTRAST_LOW',sid,'Low readability','QA-009',object_id=oid,measurement={'ratio':1},threshold={'minimum_ratio':3},action={'node':'style','mode':'local_revision'})
report={**clean,'issues':[contrastissue],'p0_issue_count':1};d={'quality_report':report,'deck':deck,'object_map':mapping};result=decide(d,schemas);cases.append({'name':'contrast_local_repair','input':d,'output':result})
test('p0_fail_with_safe_retry',result['completion']['state']=='fail' and result['completion']['retry']['allowed']);test('contrast_plan_local',result['revision_plan']['actions'][0]['operation']=='format' and result['revision_plan']['scope']['object_ids']==[oid]);test('plan_native_stale_guard_valid',validate_plan(result['revision_plan'],deck,mapping)['valid'])
locked=decide({**d,'revision_policy':{'locked_object_ids':[oid]}},schemas);cases.append({'name':'locked_p0_conflict','input':{**d,'revision_policy':{'locked_object_ids':[oid]}},'output':locked});test('locked_p0_explicit_fail',locked['revision_plan']['status']=='fail' and not locked['revision_plan']['actions'] and not locked['completion']['retry']['allowed'])
error('locked_plan_not_consumable',lambda:validate_plan(locked['revision_plan'],deck,mapping),'REVISION_PLAN_BLOCKED')
limited=decide({**d,'completion_policy':{'retry_count':2,'max_retries':2}},schemas);test('retry_exhausted_no_loop',not limited['completion']['retry']['allowed'])
error('p0_count_cannot_hide_issue',lambda:decide({**d,'quality_report':{**report,'p0_issue_count':0}},schemas),'QUALITY_P0_COUNT_MISMATCH')
error('unknown_policy_fails',lambda:decide({**d,'completion_policy':{'ignore_p0':True}},None),'COMPLETION_POLICY_UNSUPPORTED')
error('unknown_lock_fails',lambda:build_plan(deck,mapping,[],{'locked_object_ids':['absent']}),'REVISION_LOCK_TARGET_UNKNOWN')
factual=issue('FACT_NUMBER_UNSUPPORTED',sid,'Unsupported quantity','QA-014',object_id=oid,action={'node':'fact','mode':'requires_user_input'})
factplan=build_plan(deck,mapping,[factual]);test('fact_not_guessed',factplan['status']=='needs_input' and not factplan['actions'])
stale=copy.deepcopy(deck);stale['slides'][0]['elements'][0]['bounds'][0]+=0.1
error('stale_spec_fails',lambda:validate_plan(result['revision_plan'],stale,mapping),'REVISION_PLAN_STALE')
history=history_input(result['revision_plan'],'a'*64,'b'*64,['ppt/slides/slide1.xml'],report,clean);test('history_does_not_invent_preservation',history['unchanged_parts_byte_preserved'] is None and history['before_specs'])
empty_policy=decide({'quality_report':clean,'deck':deck,'object_map':mapping,'completion_policy':{'required_nodes':[],'required_coverage':[],'acceptance_required':[]}},schemas)
test('empty_policy_cannot_remove_baseline',empty_policy['completion']['state']=='warning' and len(empty_policy['completion']['missing_nodes'])==38 and empty_policy['completion']['system_acceptance']=='not_passed')
# Actual native PPTX mutations check measured rules, not mocks.
prs=Presentation(root/'deck.pptx');shape=next(s for s in prs.slides[0].shapes if s.name==oid);shape.left=Inches(-1)
issues,_,_=audit(deck,prs,mapping);test('actual_geometry_injection_detected',any(i['code']=='OBJECT_GEOMETRY_MISMATCH' and i['object_id']==oid for i in issues))
prs=Presentation(root/'deck.pptx');e2=next(e for e in slide['elements'] if e['role']=='body');sh2=next(s for s in prs.slides[0].shapes if s.name==e2['object_id']);sh1=next(s for s in prs.slides[0].shapes if s.name==oid);sh2.left=sh1.left;sh2.top=sh1.top
# Collision test consumes the authoritative changed spec as it would after a patch.
collision=copy.deepcopy(deck);c2=next(e for e in collision['slides'][0]['elements'] if e['object_id']==e2['object_id']);c2['bounds'][0]=target['bounds'][0];c2['bounds'][1]=target['bounds'][1]
issues,_,_=audit(collision,prs,mapping);test('actual_overlap_injection_detected',any(i['code']=='OBJECT_OVERLAP' for i in issues))
low=copy.deepcopy(deck);le=next(e for e in low['slides'][0]['elements'] if e['object_id']==oid);le.setdefault('format',{})['color']=low['slides'][0]['style']['background']
issues,_,_=audit(low,Presentation(root/'deck.pptx'),mapping);test('contrast_injection_detected',any(i['code']=='TEXT_CONTRAST_LOW' for i in issues));test('contrast_black_white',abs(contrast('000000','FFFFFF')-21)<1e-9)
# Relationship target removed, package validator must detect it independently.
broken=out/'relationship-broken.pptx'
with ZipFile(root/'deck.pptx') as before,ZipFile(broken,'w') as after:
 for info in before.infolist():
  if info.filename!='ppt/theme/theme1.xml':after.writestr(info,before.read(info.filename))
issues,_=package_audit(broken);test('missing_relation_target_injection',any(i['code']=='RELATIONSHIP_TARGET_MISSING' for i in issues));test('content_type_orphan_injection',any(i['code']=='CONTENT_TYPE_PART_MISSING' for i in issues))
# New real native image/diagram/connector objects use the production writer.
pixels=out/'quality-image.png';Image.new('RGB',(1920,960),(90,120,140)).save(pixels)
style=deck['style'];visual_prs=Presentation();visual_prs.slide_width=Inches(deck['width_inches']);visual_prs.slide_height=Inches(deck['height_inches']);native_slide=visual_prs.slides.add_slide(visual_prs.slide_layouts[6])
vtitle=copy.deepcopy(target);vtitle['bounds']=[0.5,0.3,12,1];vtitle['font_size']=30;vtitle['text']=['验证标题']
image_e={'object_id':entity('object'),'type':'image','role':'asset','bounds':[4.4,1.5,4,2],'font_size':16,'source_refs':slide['source_refs'],'data':{'path':str(pixels),'width_px':1920,'height_px':960,'sha256':hashlib.sha256(pixels.read_bytes()).hexdigest(),'alt_text':'示意图片','crop':[0,0,0,0],'fit':'contain'}}
node_e={'object_id':entity('object'),'type':'shape','role':'diagram','bounds':[0.5,1.5,2,1],'font_size':18,'source_refs':slide['source_refs'],'data':{'node_id':'literal-node','shape_type':'rounded_rectangle','text':'明示节点'}}
edge_e={'object_id':entity('object'),'type':'connector','role':'diagram','bounds':[2.5,2,1.9,0],'font_size':16,'source_refs':slide['source_refs'],'data':{'start':[2.5,2],'end':[4.4,2],'label':'明示关系','relation':'sequence','connector_type':'straight'}}
v_elements=[vtitle,image_e,node_e,edge_e];vmap={'objects':[]}
for el,fn in [(vtitle,text_shape),(image_e,image_shape),(node_e,diagram_shape),(edge_e,connector_shape)]:
 shape,_=fn(native_slide,el,style);shape.name=el['object_id'];vmap['objects'].append({'logical_object_id':el['object_id'],'shape_id':shape.shape_id,'part':str(native_slide.part.partname)})
visual_deck=copy.deepcopy(deck);visual_deck['slides']=[copy.deepcopy(slide)];visual_deck['slides'][0]['elements']=v_elements
visissues,_,_=audit(visual_deck,visual_prs,vmap);test('native_image_diagram_connector_positive',not any(i['severity']=='P0' for i in visissues),[i['code'] for i in visissues])
picture=next(sh for sh in native_slide.shapes if sh.name==image_e['object_id']);picture._element.nvPicPr.cNvPr.set('descr','')
visissues,_,_=audit(visual_deck,visual_prs,vmap);test('actual_alt_text_fault_detected',any(i['code']=='ALT_TEXT_MISSING' for i in visissues));picture._element.nvPicPr.cNvPr.set('descr','示意图片')
node_shape=next(sh for sh in native_slide.shapes if sh.name==node_e['object_id']);node_shape.text='未知999%'
visissues,_,_=audit(visual_deck,visual_prs,vmap);test('actual_diagram_node_label_fault_detected',any(i['code']=='DIAGRAM_LABEL_CHANGED' for i in visissues))
native_slide.notes_slide.notes_text_frame.text=slide['content']['notes']+'\n'+'\n'.join(slide['source_refs']);native_path=out/'native-visual-fault.pptx';visual_prs.save(native_path)
visqa=check(visual_deck,native_path,vmap,sources,renders,root);test('actual_diagram_fact_boundary_enforced',any(i['code']=='VISUAL_FACT_NUMBER_UNSUPPORTED' for i in visqa['issues']))
node_shape.text=node_e['data']['text'];edge_shape=next(sh for sh in native_slide.shapes if sh.name==edge_e['object_id']);edge_shape._element.nvCxnSpPr.cNvPr.set('descr','其他关系')
visissues,_,_=audit(visual_deck,visual_prs,vmap);test('actual_connector_label_fault_detected',any(i['code']=='DIAGRAM_CONNECTOR_LABEL_CHANGED' for i in visissues))
# Actual native color overrides are checked, not merely the planned color.
prs=Presentation(root/'deck.pptx');native_title=next(sh for sh in prs.slides[0].shapes if sh.name==oid);native_title.text_frame.paragraphs[0].font.color.rgb=RGBColor.from_string(slide['style']['background'])
visissues,_,_=audit(deck,prs,mapping);test('actual_native_color_fault_detected',any(i['code']=='ACTUAL_TEXT_CONTRAST_LOW' for i in visissues))

(out/'cases.json').write_text(json.dumps(cases,ensure_ascii=False,indent=2)+'\n','utf-8')
paths=[Path('/workspace/src/yoloongppt')/n for n in ['quality.py','quality_rules.py','quality_decisions.py','revision_plans.py']]+[Path('/workspace/contracts')/n for n in ['quality-stage-input.schema.json','quality-stage-output.schema.json','revision-plan-runtime-output.schema.json']]
result={'passed':all(c['passed'] for c in checks),'checks':checks,'requirement_ids':['DEC-039','DEC-040','SYS-015','SYS-016','QA-001','QA-003','QA-004','QA-005','QA-006','QA-008','QA-009','QA-010','QA-015','QA-016','QA-018','QA-019','QA-020','REV-006','REV-010'],'source_hashes':{str(p.relative_to('/workspace')):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},'real_pptx_sha256':hashlib.sha256((root/'deck.pptx').read_bytes()).hexdigest(),'real_qa_p0':quality['p0_issue_count'],'node_cases':str(out/'cases.json'),'model_calls':0,'system_acceptance':'not_passed','remaining':['Full original QA acceptance','PowerPoint native open validation','Semantic truth/visual approval','All RevisionPlan actions/locks/rollback/end-to-end']}
(out/'quality-stage-runtime.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n','utf-8');print(json.dumps({'passed':result['passed'],'checks':len(checks),'real_qa_p0':quality['p0_issue_count'],'report':str(out/'quality-stage-runtime.json')},ensure_ascii=False))
