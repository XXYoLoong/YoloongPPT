"""Actual Docker verification of directly reused P01 native export and editing."""
from pathlib import Path
import hashlib
import json
import sys
import zipfile
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.util import Inches

from yoloongppt.errors import TaskError
from yoloongppt.upstream_native import (audit_native_deck, canvas_guard, export_native_roundtrip,
                                       export_svg_deck, import_native_deck, insert_native_svg,
                                       relationship_guard, upstream_version, chart_edit_guard,
                                       table_merge_guard)

ROOT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('/runtime/upstream-native-validation')
ROOT.mkdir(parents=True, exist_ok=True)
checks = []


def check(name, condition, detail=None):
    checks.append({'name': name, 'passed': bool(condition), 'detail': detail})
    if not condition:
        raise AssertionError(name)


def blocked(name, operation, codes):
    try:
        operation()
    except TaskError as error:
        check(name, error.code in codes, error.result('validation'))
    else:
        check(name, False)


def entries(path):
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


SVG = '''<svg xmlns="http://www.w3.org/2000/svg" width="960" height="540" viewBox="0 0 960 540">
<defs><linearGradient id="grad"><stop offset="0" stop-color="#2563eb"/><stop offset="1" stop-color="#60a5fa"/></linearGradient></defs>
<rect x="0" y="0" width="960" height="540" fill="#ffffff"/>
<g id="vector-unit"><rect id="panel" x="48" y="100" width="400" height="230" rx="12" fill="url(#grad)"/>
<path id="trend" d="M 90 280 L 200 180 L 310 240 L 400 140" fill="none" stroke="#0f172a" stroke-width="4"/>
<circle id="point" cx="200" cy="180" r="8" fill="#ffffff"/></g>
<text id="heading" x="48" y="65" font-family="Noto Sans CJK SC" font-size="32" fill="#0f172a">原生矢量来源</text>
<text x="48" y="400" font-family="Noto Sans CJK SC" font-size="24" fill="#0f172a">来源数值：100 kWh</text>
</svg>'''


try:
    one, two = ROOT / 'slide-one.svg', ROOT / 'slide-two.svg'
    one.write_text(SVG, encoding='utf-8')
    two.write_text(SVG.replace('原生矢量来源', '第二页保留内容'), encoding='utf-8')
    check('fixed_upstream_commit', upstream_version()['commit'] == '2d72da616cf9fa40d4dcaf59fd4c980ecf534b7d')
    check('canvas_real_upstream_guard', canvas_guard(one)['size_emu'] == [9144000, 5143500])
    check('opaque_relationship_QName_detected', relationship_guard('<a xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="rId2"/>')['has_relationship_attributes'])
    check('ordinary_id_not_relationship', not relationship_guard('<a id="rId2"/>')['has_relationship_attributes'])
    chart_namespace = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
    scatter = chart_edit_guard(f'<c:chartSpace xmlns:c="{chart_namespace}"><c:chart><c:plotArea><c:scatterChart/></c:plotArea></c:chart></c:chartSpace>')
    check('upstream_scatter_edit_blocked_before_mutation', not scatter['supported'] and scatter['code']=='chart_edit_scatter_unsupported', scatter)
    table_namespace = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    merged = table_merge_guard(f'<a:tbl xmlns:a="{table_namespace}"><a:tr><a:tc gridSpan="2"/><a:tc hMerge="1"/></a:tr></a:tbl>')
    check('upstream_table_merge_anchor_slave_detected', merged['has_merges'] and merged['anchors'][0]['col_span']==2 and merged['slave_cells'][0]['anchor']=={'row':0,'col':0}, merged)
    deck = ROOT / 'exported.pptx'
    result = export_svg_deck([one, two], deck, title='直接复用原生导出',
                             notes={'slide-one': '人工验证来源：100 kWh', 'slide-two': '保持第二页'}, resource_root=ROOT)
    (ROOT / 'export-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    prs = Presentation(deck)
    check('two_real_native_slides', len(prs.slides) == 2)
    check('exported_text_editable', any('原生矢量来源' in s.text for s in prs.slides[0].shapes if s.has_text_frame))
    check('native_custom_geometry_not_raster', any(row['custom_geometry'] for row in result['native_objects']))
    check('native_group_not_bitmap', any(row['kind'] == 'grpSp' for row in result['native_objects']))
    check('speaker_notes_preserved', '100 kWh' in prs.slides[0].notes_slide.notes_text_frame.text)
    check('title_native_core_property', prs.core_properties.title == '直接复用原生导出', prs.core_properties.title)
    check('real_upstream_delivery_audit', not result['delivery_audit']['errors'])
    check('no_silent_skipped_object', all(not row.get('summary', {}).get('skipped') for row in result['conversion']['slides']))
    check('read_only_audit_callable', not audit_native_deck(deck)['audit']['errors'])
    wrong = ROOT / 'wrong-origin.svg'
    wrong.write_text(SVG.replace('0 0 960 540', '5 5 960 540'), encoding='utf-8')
    blocked('invalid_canvas_rejected', lambda: canvas_guard(wrong), {'UPSTREAM_NATIVE_FAILED'})
    remote = ROOT / 'external-image.svg'
    remote.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 960 540"><image href="https://example.com/image.png"/></svg>', encoding='utf-8')
    blocked('external_image_not_fetched', lambda: export_svg_deck([remote], ROOT/'unsafe.pptx', title='拒绝外链'), {'SVG_EXTERNAL_ASSET_BLOCKED'})
    entity = ROOT / 'entity.svg'
    entity.write_text('<!DOCTYPE svg [<!ENTITY e "active">]><svg viewBox="0 0 960 540"/>', encoding='utf-8')
    blocked('DTD_not_expanded', lambda: canvas_guard(entity), {'XML_ENTITY_BLOCKED'})
    # Actual existing native target retains its other slide and structural parts.
    target = Presentation()
    slide = target.slides.add_slide(target.slide_layouts[6])
    original = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    original.text = '现有文字保持'
    untouched = target.slides.add_slide(target.slide_layouts[6])
    untouched.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = '非目标页保持'
    before = ROOT / 'insert-before.pptx'
    target.save(before)
    target = Presentation(before)
    before_shape_count = len(target.slides[0].shapes)
    blocked('insertion_failure_keeps_target', lambda: insert_native_svg(target.slides[0], remote, 'vector-bad', [1,2,5,3]), {'SVG_EXTERNAL_ASSET_BLOCKED'})
    check('no_target_mutation_on_failure', len(target.slides[0].shapes) == before_shape_count)
    mapping = insert_native_svg(target.slides[0], one, 'object-vector-01', [1,2,5,2.8125], resource_root=ROOT)
    inserted = ROOT / 'insert-after.pptx'
    target.save(inserted)
    (ROOT / 'insert-map.json').write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding='utf-8')
    check('exact_stable_group_id_map', mapping['shape_name'] == 'object-vector-01' and len(mapping['shape_ids']) >= 5, mapping)
    check('contain_scale_explicit', mapping['placement']['fit']=='contain' and mapping['placement']['scale']==0.5)
    check('native_absolute_font_and_line_units_scaled', any(r['attribute']=='sz' and r['after']==r['before']/2 for r in mapping['placement']['scaled_absolute_units']) and any(r['tag']=='ln' and r['attribute']=='w' and r['after']==r['before']/2 for r in mapping['placement']['scaled_absolute_units']))
    native_ids = [int(node.get('id')) for node in target.slides[0]._element.iter() if node.tag.rsplit('}',1)[-1]=='cNvPr']
    check('remapped_native_ids_unique', len(native_ids) == len(set(native_ids)))
    check('target_text_unchanged', target.slides[0].shapes[0].text == '现有文字保持')
    a, b = entries(before), entries(inserted)
    check('non_target_slide_bytes_preserved', a['ppt/slides/slide2.xml'] == b['ppt/slides/slide2.xml'])
    structural = [name for name in a if name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/'))]
    check('native_master_layout_theme_bytes_preserved', all(a[name] == b[name] for name in structural))
    check('target_relationship_bytes_preserved', a['ppt/slides/_rels/slide1.xml.rels'] == b['ppt/slides/_rels/slide1.xml.rels'])
    check('inserted_file_reopens_native_group', any(s.name == 'object-vector-01' for s in Presentation(inserted).slides[0].shapes))
    check('inserted_file_passes_upstream_audit', not audit_native_deck(inserted)['audit']['errors'])
    # Direct upstream import/export source-preserving editor, with explicit local text edit.
    workspace = ROOT / 'native-workspace'
    imported = import_native_deck(deck, workspace)
    check('real_source_preserving_import', imported['slide_count'] == 2)
    first = workspace / 'authoring-svg-flat' / 'slide_01.svg'
    data = first.read_text(encoding='utf-8')
    check('authoring_selected_text_locatable', '原生矢量来源' in data)
    first.write_text(data.replace('原生矢量来源', '局部修订已应用'), encoding='utf-8')
    revised = ROOT / 'roundtrip-revised.pptx'
    revised_result = export_native_roundtrip(workspace, revised)
    (ROOT / 'roundtrip-result.json').write_text(json.dumps(revised_result, ensure_ascii=False, indent=2), encoding='utf-8')
    src_parts, dst_parts = entries(deck), entries(revised)
    check('native_revision_text_applied', '局部修订已应用' in dst_parts['ppt/slides/slide1.xml'].decode('utf-8'))
    check('roundtrip_unselected_slide_byte_identity', src_parts['ppt/slides/slide2.xml'] == dst_parts['ppt/slides/slide2.xml'])
    check('roundtrip_master_layout_theme_byte_identity', all(src_parts[n] == dst_parts[n] for n in src_parts if n.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/'))))
    check('roundtrip_source_value_preserved', '100 kWh' in ''.join(t.text or '' for t in ET.fromstring(dst_parts['ppt/slides/slide1.xml']).iter() if t.tag.endswith('}t')))
    blocked('existing_workspace_not_overwritten', lambda: import_native_deck(deck, workspace), {'NATIVE_WORKSPACE_EXISTS'})
    from yoloongppt.rendering import render
    for name, pptx in [('export', deck), ('insert', inserted), ('revision', revised)]:
        render_root = ROOT / (name + '-render')
        render_root.mkdir(exist_ok=False)
        report = render(pptx, {'slides':[{'slide_id':'slide-1'},{'slide_id':'slide-2'}]}, render_root)
        (render_root/'render-result.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        check(name+'_real_PDF_and_two_PNGs', len(report['artifacts'])==3 and all((render_root/row['path']).stat().st_size > 0 for row in report['artifacts']))
except Exception as error:
    (ROOT / 'failure.json').write_text(json.dumps({'type': type(error).__name__, 'message': str(error),
                                                   'checks': checks}, ensure_ascii=False, indent=2), encoding='utf-8')
    raise
finally:
    (ROOT / 'verification.json').write_text(json.dumps({'upstream': upstream_version(), 'checks': checks,
                                                       'passed': sum(c['passed'] for c in checks),
                                                       'total': len(checks),
                                                       'scope': 'component integration, not all PPT/AC acceptance'},
                                                      ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'passed': sum(c['passed'] for c in checks), 'total': len(checks), 'evidence': str(ROOT)}, ensure_ascii=False))
