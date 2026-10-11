"""Actual API multi-page revisions; reuse fixed real parents, never synthetic models."""
import argparse
import copy
import hashlib
import json
import urllib.request
from pathlib import Path
from zipfile import ZipFile

from lxml import etree

ROOT=Path('/runtime/batch-revision-validation');ROOT.mkdir(exist_ok=True)
parser=argparse.ArgumentParser();parser.add_argument('--case',choices=['text','chart','image','beautify','negative'],required=True)
parser.add_argument('--result',help='Reuse the saved output for evidence-only checks')
parser.add_argument('--parent-run',help='Use an explicitly saved parent for a repeated revision')
parser.add_argument('--label',help='Keep repeated validation artifacts separate from the original case')
args=parser.parse_args();checks=[]
label=args.label or args.case
load=lambda p:json.loads(p.read_text('utf-8'))
def test(name,condition,detail=None):
    checks.append({'name':name,'passed':bool(condition),'detail':detail})
    if not condition:raise AssertionError(name)

def call(run_id,request):
    raw=json.dumps(request,ensure_ascii=False).encode('utf-8')
    req=urllib.request.Request('http://127.0.0.1:8000/revise/'+run_id,data=raw,headers={'Content-Type':'application/json'})
    try:
        with urllib.request.urlopen(req,timeout=600) as response:return json.load(response)
    except urllib.error.HTTPError as error:return json.load(error)

def persist(value,name):
    (ROOT/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n','utf-8')

parent_id='468ac35a-eef5-424f-9585-2fa45ad6722a' if args.case=='image' else '39a2dbf3-c19b-4da2-99a6-27fe7b82ef7a'
parent_id=args.parent_run or parent_id
parent=Path('/runtime/runs')/parent_id;deck=load(parent/'deck-spec.json')
if args.case=='text':
    first=deck['slides'][0];last=deck['slides'][-1]
    request={'instruction':'只修改第1页标题为“虚构测试：周照明用电演示”，第6页标题为“结论：核对资料并按原文流程执行”。其他页、正文、对象和事实不变。',
             'scope':{'slide_ids':[first['slide_id'],last['slide_id']]}}
elif args.case=='chart':
    s=deck['slides'][4];e=next(e for e in s['elements'] if e['type']=='chart')
    value={'categories':e['data']['categories'],'series':list(reversed(e['data']['series']))}
    request={'instruction':'仅将原柱状图两系列的顺序改为本期、基线，保持类别、数值和千瓦时单位；其余对象保持。',
             'scope':{'object_ids':[e['object_id']]},'edits':[{'object_id':e['object_id'],'operation':'chart_data','value':value}]}
elif args.case=='image':
    s=deck['slides'][0];images=[e for e in s['elements'] if e['type']=='image'];a,b=images[:2]
    captions=[next(e for e in s['elements'] if e['slot_key']=='caption.'+image['slot_key'].split('.',1)[1]) for image in [a,b]]
    request={'instruction':'互换前两张已提供图片及其原文图注和替代文字；使用contain保持全部像素，保留其余对象。',
             'scope':{'object_ids':[a['object_id'],b['object_id'],*[e['object_id'] for e in captions]]},'edits':[
                 {'object_id':a['object_id'],'operation':'image','value':{'asset_id':b['data']['asset_id']}},
                 {'object_id':b['object_id'],'operation':'image','value':{'asset_id':a['data']['asset_id']}},
                 {'object_id':a['object_id'],'operation':'alt_text','value':b['data']['alt_text']},
                 {'object_id':b['object_id'],'operation':'alt_text','value':a['data']['alt_text']},
                 {'object_id':captions[0]['object_id'],'operation':'text','value':captions[1]['text']},
                 {'object_id':captions[1]['object_id'],'operation':'text','value':captions[0]['text']}]}
elif args.case=='beautify':
    # Use the previously saved real-parent fault injection, whose identity is explicit.
    fixture=load(Path('/workspace/validation/revision-plan-artifacts/revision-plan-execution.json')) if Path('/workspace/validation/revision-plan-artifacts/revision-plan-execution.json').is_file() else None
    candidates=[p for p in Path('/runtime/runs').iterdir() if (p/'fixture-identity.json').is_file()]
    candidates=[p for p in candidates if load(p/'fixture-identity.json').get('kind')=='artificial_fault_injection']
    if not candidates:raise RuntimeError('Saved explicit fault injection not found; do not invent a successful beautify input')
    parent=sorted(candidates,key=lambda p:p.stat().st_mtime)[-1];parent_id=parent.name;deck=load(parent/'deck-spec.json');s=deck['slides'][0]
    reference=load(Path('/runtime/runs')/load(parent/'fixture-identity.json')['source_run_id']/'deck-spec.json')
    orig={e['object_id']:e for e in reference['slides'][0]['elements']};ids=load(parent/'fixture-identity.json')['injected_objects']
    locked=[e['object_id'] for p in deck['slides'] for e in p['elements'] if e['object_id'] not in ids]
    request={'instruction':'美化第一页两个指定对象：恢复正文的原始可读布局并恢复标题与背景的对比度。所有文字、数据、图片保持，其他对象锁定。',
             'mode':'beautify','scope':{'object_ids':ids},'locked_object_ids':locked,'edits':[
                 {'object_id':ids[0],'operation':'format','value':{'color':orig[ids[0]].get('format',{}).get('color',reference['slides'][0]['style']['accent'])}},
                 {'object_id':ids[1],'operation':'geometry','value':orig[ids[1]]['bounds']}]}
else:
    s=deck['slides'][0];e=next(e for e in s['elements'] if e['role']=='title')
    base={'instruction':'明确局部修改验证','scope':{'object_ids':[e['object_id']]},'edits':[{'object_id':e['object_id'],'operation':'text','value':['无来源数值987654321']} ]}
    cases=[('unselected_numeric',base,'REVISION_FACT_UNSUPPORTED'),
           ('locked',base|{'locked_object_ids':[e['object_id']]},'REVISION_OBJECT_LOCKED'),
           ('scope',base|{'scope':{'object_ids':[deck['slides'][1]['elements'][0]['object_id']]}},'REVISION_SCOPE_EXCEEDED'),
           ('beautify_content',base|{'mode':'beautify'},'REVISION_BEAUTIFY_CONTENT_CHANGE'),
           ('duplicate',base|{'edits':base['edits']*2},'REVISION_EDIT_CONFLICT')]
    before=hashlib.sha256((parent/'deck.pptx').read_bytes()).hexdigest()
    for name,req,code in cases:
        result=call(parent_id,req);persist(result,'negative-'+name+'.json');test(name,result.get('error',{}).get('code')==code,result)
    test('parent_unchanged',hashlib.sha256((parent/'deck.pptx').read_bytes()).hexdigest()==before)
    persist({'case':args.case,'checks':checks,'passed':all(c['passed'] for c in checks)},'negative-evidence.json');print(json.dumps({'case':args.case,'checks':len(checks),'passed':True}));raise SystemExit()

persist({'parent_run_id':parent_id,'request':request},label+'-request.json')
result=load(Path(args.result)) if args.result else call(parent_id,request);persist(result,label+'-result.json')
try:
    test('api_success',result.get('ok'),result)
    output=Path('/runtime/runs')/result['run_id'];history=load(output/'revision-history.json');qa=load(output/'quality-report.json')
    targets=set(history['scope']['object_ids']);orders=[s['order'] for s in deck['slides'] if s['slide_id'] in history['scope']['slide_ids']]
    test('precise_target_scope',targets=={a['object_id'] for a in history['actions']})
    test('real_fresh_review_selected_pages',qa['model_review']['fresh_checked_slide_orders']==orders and load(output/'review-model-call.json')['real_model_call'])
    test('independent_review_preserved_unselected_issues',all(i in qa['model_review']['issues'] for i in load(parent/'quality-report.json').get('model_review',{}).get('issues',[]) if i['slide_order'] not in orders))
    test('all_pages_actual_render',len([a for a in load(output/'render-report.json')['artifacts'] if a['type']=='png'])==len(deck['slides']))
    test('mapping_preserved',load(output/'object-map.json')==load(parent/'object-map.json'))
    test('selected_fact_snapshots_preserved',all((output/n).read_bytes()==(parent/n).read_bytes() for n in ['approved-source-result.json','fact-boundary.json','evidence-resolution.json','fact-interpretation.json']))
    test('history_has_before_after_qa',history['quality_before'] and history['quality_after'] and (output/history['before_spec']).is_file() and (output/history['after_spec']).is_file())
    test('final_gate_still_original_baseline',result['completion']['final_allowed'] is False and result['system_acceptance']=='not_passed')
    prior_traces=[load(p) for p in parent.glob('decision-*.json') if p.name not in {'decision-DEC-039.json','decision-DEC-040.json'}]
    prior_traces += [{'node_id':d['node_id'],'output':d['result'],'error':None} for s in deck['slides'] for d in s['layout'].get('selection_trace',{}).get('decisions',[])]
    evidenced={t.get('node_id') for t in prior_traces if t.get('error') is None and t.get('output') is not None}
    expected_missing=[f'DEC-{n:03d}' for n in range(1,39) if f'DEC-{n:03d}' not in evidenced]
    test('decision_chain_ancestry_retained',result['completion']['missing_nodes']==expected_missing,result['completion'].get('missing_nodes'))
    with ZipFile(parent/'deck.pptx') as before,ZipFile(output/'deck.pptx') as after:
        test('package_members_preserved',before.namelist()==after.namelist())
        changed=[n for n in before.namelist() if before.read(n)!=after.read(n)]
        from yoloongppt.existing_deck import inspect_deck
        native_before=inspect_deck(parent/'deck.pptx')
        unselected_parts=[s['part'] for s in native_before['slides'] if s['slide_index'] not in orders]
        test('unselected_slides_byte_preserved',len(unselected_parts)==len(deck['slides'])-len(orders) and all(before.read(p)==after.read(p) for p in unselected_parts))
        ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
        def objects(blob):
            root=etree.fromstring(blob)
            return {n.xpath('.//p:cNvPr',namespaces=ns)[0].get('name'):etree.tostring(n) for n in root.xpath('./p:cSld/p:spTree/*',namespaces=ns) if n.xpath('.//p:cNvPr',namespaces=ns)}
        for order in orders:
            part='ppt/slides/slide'+str(order+1)+'.xml';a=objects(before.read(part));b=objects(after.read(part))
            test('unselected_objects_page_'+str(order),all(b[k]==v for k,v in a.items() if k not in targets))
        test('actual_changed_parts_recorded',set(changed)==set(history['changed_parts']),changed)
    if args.case=='text':test('natural_language_real_intent',load(output/'revision-intent-model-call.json')['real_model_call'])
    if args.case=='chart':
        from pptx import Presentation
        native_chart=next(s.chart for s in Presentation(output/'deck.pptx').slides[4].shapes if s.has_chart)
        expected=request['edits'][0]['value']
        test('actual_native_chart_data_matches',[c.label for c in native_chart.plots[0].categories]==expected['categories'] and [{'name':s.name,'values':list(s.values)} for s in native_chart.series]==expected['series'])
    if args.case=='beautify':test('facts_and_data_preserved',all(a['before_spec'].get('text')==a['after_spec'].get('text') and a['before_spec'].get('data')==a['after_spec'].get('data') for a in history['actions']))
    persist({'case':args.case,'parent_run_id':parent_id,'run_id':result['run_id'],'checks':checks,'passed':all(c['passed'] for c in checks),'qa':qa,'completion':result['completion']},label+'-evidence.json')
    print(json.dumps({'case':args.case,'run_id':result['run_id'],'checks':len(checks),'passed':True,'p0':qa['p0_issue_count']},ensure_ascii=False))
finally:
    persist({'case':args.case,'parent_run_id':parent_id,'result':result,'checks':checks},label+'-checks.json')
