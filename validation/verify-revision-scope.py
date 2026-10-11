"""REV-003/004/009 changed-component checks; fixtures do not prove AC acceptance."""
import copy
import json
import sys
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, '/workspace/src')

from PIL import Image
from pptx import Presentation
from yoloongppt.artifacts import RunArtifacts, entity
from yoloongppt.errors import TaskError
from yoloongppt.existing_deck import apply_patch, inspect_deck
from yoloongppt.native_objects import add_object
from yoloongppt.revision_scope import resolve_revision
from yoloongppt.schemas import SchemaRegistry

schemas = SchemaRegistry()
artifacts = RunArtifacts('/runtime/revision-scope-validation')
folder = artifacts.path
checks = []


def check(name, condition):
    assert condition, name
    checks.append(name)


def reject(name, code, fn):
    try:
        fn()
    except TaskError as error:
        check(name, error.code == code)
    else:
        raise AssertionError(name)


image = folder / 'registered.png'
replacement = folder / 'replacement.png'
Image.new('RGB', (32, 32), 'blue').save(image)
Image.new('RGB', (32, 32), 'red').save(replacement)
prs = Presentation(); prs.slide_width = 12192000; prs.slide_height = 6858000
deck = {'deck_id': entity('deck'), 'title': '范围测试', 'width_inches': 13.333333, 'height_inches': 7.5, 'slides': []}
object_map = {'objects': []}
ids = {}
for n in range(3):
    native_slide = prs.slides.add_slide(prs.slide_layouts[6]); sid = entity('slide')
    native_slide.notes_slide.notes_text_frame.text = '原始备注'
    elements = []
    specs = [('title', {'kind': 'text', 'name': '重复标题', 'text': f'第{n+1}页标题', 'bounds': [1, 0.5, 6, 0.8]})]
    if n == 0:
        specs += [('table', {'kind': 'table', 'name': '数据表', 'bounds': [1, 2, 3, 2], 'cells': [['项目', '金额'], ['A', '50']]}),
                  ('chart', {'kind': 'chart', 'name': '收入图', 'bounds': [5, 2, 4, 3], 'data': {'categories': ['A'], 'series': [{'name': '收入', 'values': [50]}]}}),
                  ('image', {'kind': 'image', 'name': '产品图', 'path': str(image), 'bounds': [10, 2, 1, 1]})]
    if n == 1:
        specs += [('body', {'kind': 'text', 'name': '正文', 'text': '正文保持事实', 'bounds': [1, 2, 5, 2]})]
    for key, spec in specs:
        obj = add_object(native_slide, {**spec, 'style': {'font_size': 24}})
        oid = entity('object'); ids[f'{n}-{key}'] = oid
        typ = spec['kind']; element = {'object_id': oid, 'type': typ, 'role': 'title' if key == 'title' else 'body' if key == 'body' else 'asset' if key == 'image' else 'data', 'bounds': spec['bounds'], 'font_size': 24, 'source_refs': []}
        if typ == 'text': element['text'] = spec['text'].split('\n')
        else: element['data'] = copy.deepcopy(spec.get('data', {}))
        if typ == 'image': element['data'].update(asset_id=entity('asset'), path=str(image))
        elements.append(element)
        object_map['objects'].append({'logical_object_id': oid, 'slide_id': sid, 'type': typ, 'shape_id': obj.shape_id, 'name': obj.name, 'part': f'ppt/slides/slide{n+1}.xml'})
    deck['slides'].append({'slide_id': sid, 'order': n, 'content': {'title': f'第{n+1}页标题'}, 'elements': elements, 'assets': [], 'layout': {}})
source = folder / 'fixture.pptx'; prs.save(source)
with ZipFile(source, 'a') as archive:
    archive.writestr('customXml/vendor-unknown.xml', b'<vendor xmlns="urn:vendor">KEEP</vendor>')
native = inspect_deck(source)
aid = entity('asset')
native['registered_assets'] = [{'asset_id': aid, 'path': str(replacement), 'usage_status': 'authorized', 'license': 'test_fixture', 'width_px': 32, 'height_px': 32}]
evidence_id = entity('evidence')
sources = {'evidence': [{'evidence_id': evidence_id, 'raw_text': 'REJECTED 99'}], 'approved_texts': {evidence_id: 'APPROVED 50'}, 'fact_boundary': {'status': 'ready', 'assumption_values': {}}, 'projection_scope': 'Only selected value supplied.'}
original_asset_id = deck['slides'][0]['elements'][3]['data']['asset_id']
deck['slides'][0]['elements'][3]['data']['source_refs'] = [evidence_id]
native['registered_assets'].append({'asset_id': original_asset_id, 'path': str(image), 'source_refs': []})
task = {'providers': {'text': {'provider': 'deepseek', 'model': 'deepseek-flash'}}}


def resolve(request, selected_deck=None, selected_native=None):
    return resolve_revision(request, selected_deck or deck, object_map, selected_native or native, sources, task, artifacts, schemas)


edits = [
    {'selector': {'slide_index': 0, 'name': '重复标题'}, 'operation': 'text', 'value': ['新第一页标题']},
    {'selector': {'slide_id': deck['slides'][1]['slide_id'], 'text': '第2页标题'}, 'operation': 'text', 'value': ['新第二页标题']},
    {'object_id': ids['0-table'], 'operation': 'table_cell', 'value': {'row': 1, 'column': 1, 'text': '50'}},
    {'object_id': ids['0-chart'], 'operation': 'chart_data', 'value': {'categories': ['A'], 'series': [{'name': '收入', 'values': [50]}]}},
    {'object_id': ids['0-image'], 'operation': 'image', 'value': {'asset_id': aid}},
    {'object_id': ids['1-body'], 'operation': 'geometry', 'value': [1.1, 2, 5, 2]},
    {'object_id': ids['1-body'], 'operation': 'format', 'value': {'color': '112233', 'bold': True}},
    {'object_id': ids['0-image'], 'operation': 'alt_text', 'value': '真实替换素材'},
    {'object_id': ids['0-title'], 'operation': 'notes', 'value': '范围测试备注'}]
request = {'instruction': '仅修改前两页明确登记对象', 'scope': {'slide_ids': [s['slide_id'] for s in deck['slides'][:2]]}, 'edits': edits}
with patch('yoloongppt.revision_scope.DeepSeek', side_effect=AssertionError('explicit edits must not invoke model')):
    result = resolve(request)
check('explicit_edits_do_not_call_model', result['intent_trace']['real_model_call'] is False)
check('exact_two_page_scope', result['scope']['slide_ids'] == request['scope']['slide_ids'])
check('selectors_resolve_to_registered_logical_ids', result['edits'][0]['object_id'] == ids['0-title'] and result['edits'][1]['object_id'] == ids['1-title'])
check('third_page_excluded', deck['slides'][2]['slide_id'] not in result['scope']['slide_ids'])
patches = []
for edit in result['edits']:
    mapping = next(m for m in object_map['objects'] if m['logical_object_id'] == edit['object_id'])
    obj = next(o for s in native['slides'] for o in s['objects'] if o['part'] == mapping['part'] and o['shape_id'] == mapping['shape_id'])
    value = '\n'.join(edit['value']) if edit['operation'] == 'text' else {'path': str(replacement)} if edit['operation'] == 'image' else edit['value']
    patches.append({'object_id': obj['object_id'], 'operation': 'image_replace' if edit['operation'] == 'image' else edit['operation'], 'value': value})
output = folder / 'edited.pptx'
report = apply_patch(source, output, {'expected_sha256': native['source_sha256'], 'patches': patches})
updated = Presentation(output)
check('resolved_edits_consumed_by_real_native_executor', updated.slides[0].shapes[0].text == '新第一页标题' and updated.slides[1].shapes[0].text == '新第二页标题')
check('actual_image_replaced', updated.slides[0].shapes[3].image.blob == replacement.read_bytes())
with ZipFile(source) as before, ZipFile(output) as after:
    check('unselected_page_preserved_byte_for_byte', before.read('ppt/slides/slide3.xml') == after.read('ppt/slides/slide3.xml'))
    check('master_layout_theme_and_unknown_parts_preserved', all(before.read(n) == after.read(n) for n in before.namelist() if any(s in n for s in ['slideMasters/', 'slideLayouts/', 'theme/', 'customXml/'])))


def single(edit, **options):
    return {'instruction': '明确编辑', 'edits': [edit], **options}


base_edit = {'object_id': ids['0-title'], 'operation': 'text', 'value': ['新标题']}
reject('ambiguous_name_selector_rejected', 'REVISION_SELECTION_AMBIGUOUS', lambda: resolve(single({'selector': {'name': '重复标题'}, 'operation': 'text', 'value': ['未修改']})))
scoped_selector = resolve(single({'selector': {'name': '重复标题'}, 'operation': 'text', 'value': ['在明确边界中定位']}, scope={'slide_ids': [deck['slides'][1]['slide_id']]}))
check('explicit_boundary_disambiguates_repeated_name', scoped_selector['scope']['object_ids'] == [ids['1-title']])
reject('missing_selector_rejected', 'REVISION_OBJECT_NOT_FOUND', lambda: resolve(single({'selector': {'shape_id': 999}, 'operation': 'text', 'value': ['未修改']})))
reject('scope_escape_rejected', 'REVISION_SCOPE_EXCEEDED', lambda: resolve(single(base_edit, scope={'slide_ids': [deck['slides'][1]['slide_id']]})))
reject('inconsistent_scope_rejected', 'REVISION_SCOPE_CONFLICT', lambda: resolve(single(base_edit, scope={'slide_ids': [deck['slides'][1]['slide_id']], 'object_ids': [ids['0-title']]})))
reject('unknown_object_rejected', 'REVISION_OBJECT_NOT_FOUND', lambda: resolve(single({**base_edit, 'object_id': entity('object')})))
reject('caller_lock_rejected', 'REVISION_OBJECT_LOCKED', lambda: resolve(single(base_edit, locked_object_ids=[ids['0-title']])))
locked_deck = copy.deepcopy(deck); locked_deck['slides'][0]['elements'][0]['locked'] = True
reject('parent_spec_lock_rejected', 'REVISION_OBJECT_LOCKED', lambda: resolve(single(base_edit), selected_deck=locked_deck))
reject('unknown_lock_rejected', 'REVISION_LOCK_OBJECT_UNKNOWN', lambda: resolve(single(base_edit, locked_object_ids=[entity('object')])))
reject('beautify_text_rejected', 'REVISION_BEAUTIFY_CONTENT_CHANGE', lambda: resolve(single(base_edit, mode='beautify')))
beautify = resolve(single({'object_id': ids['1-body'], 'operation': 'geometry', 'value': [1.2, 2, 5, 2]}, mode='beautify', locked_object_ids=[ids['0-title']]))
check('beautify_only_geometry_and_preserves_lock_union', beautify['locked_object_ids'] == [ids['0-title']] and beautify['edits'][0]['operation'] == 'geometry')
reject('duplicate_property_rejected', 'REVISION_EDIT_CONFLICT', lambda: resolve({'instruction': '冲突', 'edits': [base_edit, base_edit]}))
reject('operation_type_mismatch_rejected', 'REVISION_OPERATION_TYPE_MISMATCH', lambda: resolve(single({**base_edit, 'object_id': ids['0-chart']})))
reject('unknown_image_asset_rejected', 'REVISION_ASSET_UNKNOWN', lambda: resolve(single({'object_id': ids['0-image'], 'operation': 'image', 'value': {'asset_id': entity('asset')}})))
reject('raw_image_path_schema_rejected', 'TASK_SCHEMA_INVALID', lambda: resolve(single({'object_id': ids['0-image'], 'operation': 'image', 'value': {'path': '/runtime/secrets/private'}})))
reject('table_cell_bounds_rejected', 'REVISION_TABLE_CELL_OUT_OF_BOUNDS', lambda: resolve(single({'object_id': ids['0-table'], 'operation': 'table_cell', 'value': {'row': 5, 'column': 0, 'text': '不写入'}})))
reject('chart_series_mismatch_rejected', 'REVISION_CHART_DATA_MISMATCH', lambda: resolve(single({'object_id': ids['0-chart'], 'operation': 'chart_data', 'value': {'categories': ['A', 'B'], 'series': [{'name': '收入', 'values': [50]}]}})))
reject('unsupported_operation_schema_rejected', 'TASK_SCHEMA_INVALID', lambda: resolve(single({**base_edit, 'operation': 'animation'})))
reject('outside_slide_geometry_rejected', 'REVISION_GEOMETRY_OUT_OF_SLIDE', lambda: resolve(single({'object_id': ids['0-title'], 'operation': 'geometry', 'value': [12, 1, 6, 1]})))
stale = copy.deepcopy(native); stale['slides'][0]['objects'][0]['text'] = '实际文本已改'
reject('stale_actual_native_content_rejected', 'REVISION_OBJECT_MAP_STALE', lambda: resolve(single(base_edit), selected_native=stale))
branded = copy.deepcopy(deck)
branded['slides'][0]['elements'][0]['template_binding'] = {'native_layout_part': 'ppt/slideLayouts/slideLayout1.xml', 'placeholder_index': 0}
branded['slides'][0]['layout']['template_binding'] = {'locks': {'objects': [{'part': 'ppt/slideLayouts/slideLayout1.xml', 'placeholder_index': 0}], 'safe_areas': [[0, 0, 1, 1]], 'tokens': {'foreground': '112233'}}}
reject('actual_native_template_placeholder_lock_rejected', 'REVISION_OBJECT_LOCKED', lambda: resolve(single(base_edit), selected_deck=branded))
reject('brand_safe_area_rejected', 'REVISION_BRAND_SAFE_AREA_CONFLICT', lambda: resolve(single({'object_id': ids['0-table'], 'operation': 'geometry', 'value': [0, 0, 2, 2]}), selected_deck=branded))
token_deck = copy.deepcopy(branded); token_deck['slides'][0]['layout']['template_binding']['locks']['objects'] = []
reject('brand_token_rejected', 'REVISION_BRAND_TOKEN_CONFLICT', lambda: resolve(single({'object_id': ids['0-title'], 'operation': 'format', 'value': {'color': 'FFFFFF'}}), selected_deck=token_deck))

# Synthetic responses exercise response/context checks only; live model evidence
# is produced separately by the integrated entry and must never cite this as AC.
instruction = '把第一页标题改成明确标题'
response = {'status': 'resolved', 'scope_level': 'object', 'scope': {'slide_ids': [deck['slides'][0]['slide_id']], 'object_ids': [ids['0-title']]},
            'edits': [{**base_edit, 'value': ['明确标题']}], 'selections': [{'object_id': ids['0-title'], 'instruction_quote': instruction, 'candidate_object_ids': [ids['0-title']], 'reason': '明确页序和标题'}], 'ambiguities': [], 'reason': '唯一目标'}
captured = []


class FixtureProvider:
    def __init__(self, config):
        pass

    def complete(self, messages):
        captured.append(messages)
        return copy.deepcopy(response), {'real_model_call': False, 'synthetic_component_fixture': True}


with patch('yoloongppt.revision_scope.DeepSeek', FixtureProvider):
    resolved = resolve({'instruction': instruction})
    check('model_response_resolves_exact_registered_target_fixture_only', resolved['scope']['object_ids'] == [ids['0-title']])
    check('synthetic_provider_metadata_never_claims_live_model_call', resolved['intent_trace']['real_model_call'] is False)
    model_context = json.loads(captured[0][1]['content'])
    check('model_context_excludes_rejected_fact', 'REJECTED' not in json.dumps(model_context) and model_context['evidence'][0]['text'] == 'APPROVED 50')
    check('model_context_excludes_all_file_paths', str(image) not in json.dumps(model_context) and str(replacement) not in json.dumps(model_context))
    original_asset = next(asset for asset in model_context['registered_assets'] if asset['asset_id'] == original_asset_id)
    check('registration_and_use_evidence_both_preserved', original_asset['registered_source_refs'] == [] and original_asset['usages'][0]['source_refs'] == [evidence_id])
    response['selections'][0]['candidate_object_ids'].append(ids['1-title'])
    reject('model_multi_candidate_ambiguity_rejected_fixture_only', 'REVISION_INTENT_AMBIGUOUS', lambda: resolve({'instruction': instruction}))
    response['selections'][0]['candidate_object_ids'] = [ids['0-title']]
    response['scope']['slide_ids'] = [deck['slides'][1]['slide_id']]
    reject('model_scope_disagreement_rejected_fixture_only', 'REVISION_MODEL_SCOPE_MISMATCH', lambda: resolve({'instruction': instruction}))
    response['scope']['slide_ids'] = [deck['slides'][0]['slide_id']]
    response['edits'][0]['value'] = 'invalid string instead of paragraphs'
    reject('model_value_schema_rejected_fixture_only', 'REVISION_MODEL_RESPONSE_INVALID', lambda: resolve({'instruction': instruction}))

artifacts.json('component-validation.json', {'ok': True, 'checks': checks, 'requirement_ids': ['REV-003', 'REV-004', 'REV-009'],
                                           'acceptance': 'component_only; live_model_and_visual/fact_QA/AC-008/009/010_not_proven_by_this_script',
                                           'synthetic_model_calls': len(captured), 'native_report': report})
artifacts.json('manifest.json', {'run_id': artifacts.run_id, 'artifacts': artifacts.manifest()})
print(json.dumps({'ok': True, 'checks': len(checks), 'artifact_root': str(folder), 'system_acceptance': 'not_passed'}, ensure_ascii=False))
