"""Bounded native compiler/writer checks against an actual saved model proposal."""
import copy
import hashlib
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from pptx import Presentation
from yoloongppt.atomic import AtomicRegistry
from yoloongppt.errors import TaskError,trace_id
from yoloongppt.layouts import choose
from yoloongppt.planning import compile_deck,DEFAULT_STYLE
from yoloongppt.quality import check
from yoloongppt.rendering import render
from yoloongppt.schemas import ROOT,SchemaRegistry
from yoloongppt.writer import execute

s=SchemaRegistry();registry=AtomicRegistry(s);cases=[]
def record(id,ok,actual):
    assert ok,id
    cases.append({'id':id,'passed':True,'actual':actual})
def fails(id,fn,code):
    try:fn()
    except TaskError as e:record(id,e.code==code,e.code)
    else:raise AssertionError(id)
base=Path('/runtime/runs/f75a27dc-5e9d-43fd-90a9-07164f468fb5')
load=lambda path:json.loads(path.read_text(encoding='utf-8'))
proposal=load(base/('model-response-effective.json' if (base/'model-response-effective.json').exists() else 'model-response.json'));facts=load(base/'fact-source-result.json')
text=next(copy.deepcopy(x) for x in proposal['slides'] if x['kind']=='text')
text['title']='清晰表达每一页的核心信息';text['body']=['保留原始内容和证据。','保留单位和来源。','容量不足时明确报告。','输出可编辑原生对象。']
for id,order,count,last,expected in [('normal_opening',0,8,None,'cover'),('normal_cards',2,8,None,'cards'),('normal_variation',3,8,'cards','columns'),('normal_conclusion',7,8,None,'closing')]:
    elements,style,trace=choose(text,order,count,DEFAULT_STYLE,last,registry)
    record(id,trace['selected_layout']==expected,trace)
    record(id+'_content_preserved',[t for e in elements if e['role']=='body' for t in e['text']]==text['body'],'Original order and all body paragraphs retained.')
long=copy.deepcopy(text);long['body']=['过长文字'*1000]*6
fails('capacity_boundary_no_silent_cut',lambda:choose(long,2,8,DEFAULT_STYLE,None,registry),'LAYOUT_CAPACITY_UNSUPPORTED')
for kind in ['table','chart']:
    data=next((x for x in proposal['slides'] if x['kind']==kind),None)
    if data:
        elements,style,trace=choose(data,2,8,DEFAULT_STYLE,None,registry)
        record('normal_'+kind,trace['selected_layout']==kind and next(e['data'] for e in elements if e['type']==kind)==data[kind],trace)
deck,dag=compile_deck(proposal,DEFAULT_STYLE,trace_id(),registry=registry)
s.validate('deck-execution.schema.json',deck);s.validate('deck-execution.schema.json',dag)
record('compiler_all_original_content',all(a['content']==b for a,b in zip(deck['slides'],proposal['slides'])),len(deck['slides']))
record('catalog_binding_before_write',all(registry.check_call(call) for call in dag['calls']),len(dag['calls']))
recovered,_=compile_deck(proposal,DEFAULT_STYLE,trace_id(),deck,registry)
record('checkpoint_slots_ids_preserved',[e['object_id'] for x in recovered['slides'] for e in x['elements']]==[e['object_id'] for x in deck['slides'] for e in x['elements']],'Native slot keys keep stable objects across recovery.')
for slide in deck['slides']:
    elements=slide['elements']
    for i,a in enumerate(elements):
        for b in elements[i+1:]:
            x,y,w,h=a['bounds'];xx,yy,ww,hh=b['bounds']
            assert x+w<=xx or xx+ww<=x or y+h<=yy or yy+hh<=y,(slide['order'],a['slot_key'],b['slot_key'])
record('geometric_slots_no_overlap',True,'All compiler slots including footer are nonoverlapping.')
folder=Path('/runtime/layout-validation')/trace_id();folder.mkdir(parents=True,exist_ok=False)
objects,calls=execute(deck,dag,folder/'deck.pptx',s)
presentation=Presentation(folder/'deck.pptx')
record('actual_native_objects_mapped',len(objects['objects'])==sum(len(x.shapes) for x in presentation.slides),'Every actual native object mapped.')
cover=presentation.slides[0];title=cover.shapes[0]
record('actual_opening_typography',title.text_frame.paragraphs[0].font.size.pt==40 and str(title.text_frame.paragraphs[0].font.color.rgb)==DEFAULT_STYLE['background'],'Cover font and foreground observed in PPTX.')
rendered=render(folder/'deck.pptx',deck,folder)
qa=check(deck,folder/'deck.pptx',objects,facts,rendered,folder)
(folder/'quality-report.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf-8')
record('actual_render_source_QA',qa['p0_issue_count']==0,qa)
tree=ET.parse(folder/rendered['render_bbox_path']);page=tree.findall('.//{http://www.w3.org/1999/xhtml}page')[0]
page.remove(page.findall('{http://www.w3.org/1999/xhtml}word')[0]);tree.write(folder/'render/missing-word.xhtml',encoding='utf-8')
negative=check(deck,folder/'deck.pptx',objects,facts,{**rendered,'render_bbox_path':'render/missing-word.xhtml'},folder)
record('missing_PDF_word_still_blocks',any(i['code']=='RENDER_TEXT_MISSING' for i in negative['issues']),'Actual rendered word removed in isolated counterexample; original PDF/PPTX unchanged.')
record('system_acceptance_preserved',qa['acceptance']['AC-001']=='not_passed','Full QA/DEC/AC remain incomplete.')
(folder/'deck-spec.json').write_text(json.dumps(deck,ensure_ascii=False,indent=2),encoding='utf-8')
if len(sys.argv)>1:
    run=Path('/runtime/runs')/sys.argv[1];live=load(run/'deck-spec.json');result=load(run/'result.json');review=load(run/'review-model-call.json')
    record('live_HTTP_generation',result['state']=='draft_generated' and result['model']['real_model_call'] and review['real_model_call'],{'run_id':sys.argv[1],'qa':result['qa']})
    record('live_layouts_consumed',all(x['layout'].get('design_version')=='native-editorial-1' for x in live['slides']) and (run/'layout-selection.json').is_file(),[x['layout']['selected_layout'] for x in live['slides']])
    record('live_P0_checks',result['qa']['p0_issue_count']==0,result['qa'])
paths=[*sorted((ROOT/'src/yoloongppt').glob('*.py')),ROOT/'contracts/deck-execution.schema.json',ROOT/'contracts/atomic-registry.catalog.json',ROOT/'contracts/native-layout.catalog.json']
report={'passed':True,'cases':cases,'source_model_run':base.name,'live_run':sys.argv[1] if len(sys.argv)>1 else None,'scope':'Current native text/table/chart layout subset; DEC nodes remain partial and full system AC not passed. Saved real model proposal reused for compiler checks; no model called by this script.','consumed_file_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}}
Path('/runtime/layout-runtime.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'passed':True,'checks':len(cases),'layouts':[x['layout']['selected_layout'] for x in deck['slides']],'artifact_root':str(folder)}))
