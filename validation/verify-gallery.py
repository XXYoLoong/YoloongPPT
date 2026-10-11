"""Artificial local raster fixtures: 3 native normal cases plus explicit failures.

No AI generation or model review; this does not close AC-012/system acceptance.
"""
import copy
import hashlib
import json
import argparse
import re
from pathlib import Path
from uuid import uuid4
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from yoloongppt.artifacts import entity
from yoloongppt.asset_resolver import resolve
from yoloongppt.visual import finalize
from yoloongppt.gallery_layout import candidate, evaluate
from yoloongppt.layout_decisions import image_fit, geometry
from yoloongppt.writer import image_shape, text_shape
from yoloongppt.rendering import render
from yoloongppt.schemas import SchemaRegistry
from yoloongppt.errors import TaskError
from yoloongppt.quality_rules import audit, package_audit

args=argparse.ArgumentParser();args.add_argument('--resume-existing');args=args.parse_args()
resume=bool(args.resume_existing)
root = Path('/runtime/gallery-validation') / (args.resume_existing or str(uuid4()))
if resume:
    if not root.is_dir() or root.resolve().parent != Path('/runtime/gallery-validation').resolve():raise RuntimeError('Invalid existing gallery run')
else:root.mkdir(parents=True)
schema = SchemaRegistry(); checks = []
def save(name, value):
    (root/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n','utf-8')
def test(name, condition, observation):
    checks.append({'name':name,'passed':bool(condition),'observation':observation})
    save('incremental-checks.json',{'checks':checks,'resume_existing':resume})
    if not condition:save('incremental-failure.json',{'failed_check':name,'observation':observation,'checks':checks})
    assert condition,name
def fails(name, fn, codes):
    try: fn()
    except TaskError as error: test(name,error.code in codes,{'code':error.code,'details':error.details}); return
    raise AssertionError(name+' did not reject')

ratios = [(1800,900),(900,1800),(1400,1400),(2400,800),(800,2400),(1600,1200),(1200,1600),(2000,1100)]
requests = []
for i,size in enumerate(ratios):
    path = root/f'fixture-{i+1}.png'
    if not resume:
        image = Image.new('RGB',size,(30+i*22,90+i*12,120+i*10)); draw=ImageDraw.Draw(image)
        draw.rectangle([size[0]*.18,size[1]*.2,size[0]*.75,size[1]*.75],outline='white',width=20)
        draw.text((size[0]//3,size[1]//3),f'ARTIFICIAL {i+1}',fill='white'); image.save(path)
    requests.append({'asset_id':entity('asset'),'source_type':'user','path':str(path),'license':'fixture_owned','usage_status':'authorized'})
if resume:requests=json.loads((root/'asset-input-requests.json').read_text('utf-8'))['requests']
assets = resolve({'requests':requests}, schemas=schema)
legacy = Path('/runtime/runs/39a2dbf3-c19b-4da2-99a6-27fe7b82ef7a/visual-packet.json')
packet=json.loads(legacy.read_text('utf-8')); packet['resolved_assets']=assets
evidence_id=entity('evidence'); captions=[f'素材示例{chr(65+i)}' for i in range(8)]
sources={'evidence':[{'evidence_id':evidence_id,'raw_text':'原文含有未选数字999。'}],
    'fact_boundary':{'status':'ready'},'approved_texts':{evidence_id:'；'.join(captions)}}
style={'font':'Noto Sans CJK SC','foreground':'192C32','background':'F6F4EE','accent':'147C80','minimum_size':16}
declarations=[{'asset_id':item['asset_id'],'role':['information','decoration','logo','reference'][i%4],
    'alt_text':f'人工栅格样例{chr(65+i)}','fit':'contain','required':True,
    'caption':captions[i],'caption_source_refs':[evidence_id]} for i,item in enumerate(assets['entries'])]
base={'kind':'image','title':'不同宽高比的原生多图','body':['全部为人工验证素材'],'notes':'人工fixture，非AI生成结果。','evidence_refs':[evidence_id],'images':declarations}
if resume:
    sources=json.loads((root/'approved-caption-fixture.json').read_text('utf-8'))
    base=json.loads((root/'declared-gallery-fixture.json').read_text('utf-8'));declarations=base['images']
else:
    save('asset-input-requests.json',{'requests':requests,'fixture_identity':'artificial local raster files'})
    save('resolved-assets.json',assets)
    save('approved-caption-fixture.json',sources)
    save('declared-gallery-fixture.json',base)
def resolved(content):
    proposal,result=finalize({'slides':[content]},packet,sources,schemas=schema)
    return proposal['slides'][0],result
presentation=Presentation(root/'gallery.pptx') if resume else Presentation()
if not resume:presentation.slide_width=Inches(13.333333);presentation.slide_height=Inches(7.5)
original_pptx_hash=hashlib.sha256((root/'gallery.pptx').read_bytes()).hexdigest() if resume else None
saved_render=json.loads((root/'render-report.json').read_text('utf-8')) if resume else None
native_cases=[];specs=[]
def normal(name, content):
    item,result=resolved(content); elements,page_style=candidate(item,len(native_cases),3,style)
    trace=page_style.pop('gallery_selection'); geometry(elements)
    selected=next(record for record in trace['candidates'] if record['candidate']==trace['selected'])
    slide=presentation.slides[len(native_cases)] if resume else presentation.slides.add_slide(presentation.slide_layouts[6])
    saved_elements=json.loads((root/(name+'-layout.json')).read_text('utf-8'))['elements'] if resume else None
    if not resume:slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string(style['background'])
    fits=[]
    for element in elements:
        if element['type']=='image':
            fit=image_fit(element['data'],element['bounds']); expected=next(i['fit'] for i in selected['images'] if i['slot_key']==element['slot_key'])
            assert fit==expected
            element['bounds']=fit['bounds'];element['data']['crop']=fit['crop']; element['data']['fit_result']=fit
            if resume:
                saved=next(e for e in saved_elements if e['slot_key']==element['slot_key'])
                assert saved['bounds']==element['bounds'] and saved['data']==element['data']
                shape=next(s for s in slide.shapes if s.name==saved['object_id'])
            else:shape,_=image_shape(slide,element,page_style)
            assert shape.shape_type==MSO_SHAPE_TYPE.PICTURE
            ratio=(element['data']['width_px']*(1-shape.crop_left-shape.crop_right))/(element['data']['height_px']*(1-shape.crop_top-shape.crop_bottom))
            assert abs((shape.width/shape.height)/ratio-1)<.0001
            fits.append(fit)
        elif resume:
            saved=next(e for e in saved_elements if e['slot_key']==element['slot_key'])
            shape=next(s for s in slide.shapes if s.name==saved['object_id']);assert shape.text=='\n'.join(element['text'])
        else:shape,_=text_shape(slide,element,page_style)
        if not resume:shape.name=element['object_id']
    count=len(content['images']); test(name,len([s for s in slide.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE])==count and all(f['effective_dpi']>=96 for f in fits),{'image_count':count,'selected':trace['selected'],'dpis':[f['effective_dpi'] for f in fits]})
    save(('resume-' if resume else '')+name+'-visual.json',result)
    if not resume:save(name+'-layout.json',{'elements':elements,'trace':trace})
    slide_id=[a for a in saved_render['artifacts'] if a['type']=='png'][len(specs)]['slide_id'] if resume else entity('slide')
    native_cases.append({'name':name,'images':count});specs.append({'slide_id':slide_id,'order':len(specs),'elements':saved_elements if resume else elements,'style':copy.deepcopy(style)})
    return item,result,trace
item,result,trace=normal('normal-eight-ratios',base)
test('eight-stable-native-slots',len({i['slot_key'] for c in trace['candidates'] if c['eligible'] for i in c['images']})==8,{'asset_ids':[i['asset_id'] for i in declarations]})
test('every-input-asset-selected-and-inventory-preserved',len(result['asset_plan'][0]['selection_inventory'])==8 and all(i['selected'] for i in result['asset_plan'][0]['selection_inventory']) and len(result['resolved_assets']['entries'])==8,result['asset_plan'][0]['selection_inventory'])
cover=copy.deepcopy(base);cover['images']=cover['images'][:3]
for i,image in enumerate(cover['images']):image.update(fit='cover',focal_point=[.15+i*.3,.8-i*.2])
cover_item,cover_result,cover_trace=normal('normal-cover-explicit-focus',cover)
test('unused-assets-retained-with-reason',len(cover_result['resolved_assets']['entries'])==8 and sum(not e['selected'] for e in cover_result['asset_plan'][0]['selection_inventory'])==5,cover_result['asset_plan'][0]['selection_inventory'])
crop=copy.deepcopy(base);crop['images']=crop['images'][:2]
for image in crop['images']: image.update(fit='crop',crop=[.05,.1,.1,.05],focal_point=[.7,.2])
normal('normal-explicit-crop',crop)
single=copy.deepcopy(base);single['image']=single.pop('images')[0];single_item,_=resolved(single)
test('legacy-single-compatible','image' in single_item and 'images' not in single_item and single_item['image']['sha256']==assets['entries'][0]['sha256'],single_item['image'])
bad=copy.deepcopy(base);bad['images'][0]['fit']='cover'
fails('cover-missing-focus',lambda:resolved(bad),{'IMAGE_FOCAL_POINT_REQUIRED'})
bad=copy.deepcopy(base);bad['images'][0].update(fit='cover',focal_point=[1.2,.5])
fails('focus-outside-image',lambda:resolved(bad),{'IMAGE_FOCAL_POINT_INVALID'})
bad=copy.deepcopy(base);bad['images']=[]
fails('empty-gallery-no-text-fallback',lambda:resolved(bad),{'IMAGE_COUNT_INVALID'})
bad=copy.deepcopy(base);bad['images'][0]['asset_id']=entity('asset')
fails('missing-asset',lambda:resolved(bad),{'IMAGE_ASSET_UNKNOWN'})
bad=copy.deepcopy(base);bad['images'][1]['asset_id']=bad['images'][0]['asset_id']
fails('duplicate-asset-id',lambda:resolved(bad),{'IMAGE_ASSET_DUPLICATE'})
bad=copy.deepcopy(base);bad['images'][0]['caption']='增长达到999'
fails('caption-cannot-use-unselected-raw-text',lambda:resolved(bad),{'IMAGE_CAPTION_UNGROUNDED'})
bad=copy.deepcopy(base);bad['images'][0]['caption_source_refs']=[]
fails('caption-without-evidence',lambda:resolved(bad),{'IMAGE_CAPTION_UNGROUNDED'})
bad=copy.deepcopy(base);bad['images'][0].update(fit='crop',crop=[.8,0,.4,0],focal_point=[.5,.5])
fails('crop-removes-all-pixels',lambda:resolved(bad),{'IMAGE_CROP_INVALID'})
bad=copy.deepcopy(item);bad['images'][0]['width_px']=16;bad['images'][0]['height_px']=16
fails('low-dpi-no-replacement',lambda:candidate(bad,0,1,style),{'GALLERY_CAPACITY_UNSUPPORTED'})
bad=copy.deepcopy(item);bad['images'][0]['minimum_dpi']=10000
fails('caller-dpi-threshold-enforced',lambda:candidate(bad,0,1,style),{'GALLERY_CAPACITY_UNSUPPORTED'})
bad=copy.deepcopy(base);bad['image']=bad['images'][0]
fails('single-and-gallery-conflict',lambda:resolved(bad),{'IMAGE_DECLARATION_CONFLICT'})
prior=copy.deepcopy(packet);packet['resolved_assets']['entries'][0]['sha256']='0'*64
fails('stale-asset-hash',lambda:resolved(base),{'IMAGE_ASSET_CHANGED'});packet=prior
prior=copy.deepcopy(packet);packet['resolved_assets']['entries'][0]['width_px']=1
fails('stale-asset-pixels',lambda:resolved(base),{'IMAGE_DIMENSIONS_CHANGED'});packet=prior
prior=copy.deepcopy(packet);packet['resolved_assets']['entries'][0]['path']=str(root/'missing-never-created.png')
fails('real-asset-file-missing',lambda:resolved(base),{'ASSET_FILE_MISSING'});packet=prior
prior=copy.deepcopy(packet);packet['resolved_assets']['entries'][0]['usage_status']='denied'
fails('asset-usage-denied',lambda:resolved(base),{'IMAGE_ASSET_UNAUTHORIZED'});packet=prior
bad=copy.deepcopy(item);bad['images'][0]['sha256']='0'*64
fails('declared-stale-hash-not-overwritten',lambda:resolved(bad),{'IMAGE_RESOLVED_METADATA_CONFLICT'})
bad=copy.deepcopy(item);bad['images'][0]['source_refs']=[entity('evidence')]
fails('image-unknown-source-refs',lambda:resolved(bad),{'IMAGE_SOURCE_REFS_INVALID'})
prior=copy.deepcopy(packet);duplicate=copy.deepcopy(packet['resolved_assets']['entries'][0]);duplicate['asset_id']=declarations[1]['asset_id'];packet['resolved_assets']['entries'][1]=duplicate
fails('duplicate-bytes-unconfirmed',lambda:resolved(base),{'IMAGE_BYTES_DUPLICATE_UNCONFIRMED'})
reuse=copy.deepcopy(base);reuse['images'][0]['deliberate_reuse']=reuse['images'][1]['deliberate_reuse']=True
reuse_item,_=resolved(reuse);test('duplicate-bytes-explicit-reuse-retained',reuse_item['images'][0]['sha256']==reuse_item['images'][1]['sha256'],{'same_hash':reuse_item['images'][0]['sha256']});packet=prior
for image in item['images']:schema.validate('visual-model-fields.schema.json',image)
test('resolved-image-schema-valid',True,{'images':len(item['images'])})
if not resume:presentation.save(root/'gallery.pptx')
reloaded=Presentation(root/'gallery.pptx')
test('native-pptx-reopen-13-editable-pictures',sum(s.shape_type==MSO_SHAPE_TYPE.PICTURE for slide in reloaded.slides for s in slide.shapes)==13,{'slides':len(reloaded.slides),'picture_counts':[sum(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in slide.shapes) for slide in reloaded.slides]})
render_report=saved_render if resume else render(root/'gallery.pptx',{'slides':specs},root)
if not resume:save('render-report.json',render_report)
else:
    test('reuse-native-pptx-and-render-bytes-unchanged',hashlib.sha256((root/'gallery.pptx').read_bytes()).hexdigest()==original_pptx_hash and all(hashlib.sha256((root/a['path']).read_bytes()).hexdigest()==a['hash'] for a in render_report['artifacts']),{'pptx_sha256':original_pptx_hash,'render_artifacts':render_report['artifacts']})
test('three-pages-real-render',len([a for a in render_report['artifacts'] if a['type']=='png'])==3,render_report['artifacts'])
render_text=(root/'render/deck.txt').read_text('utf-8')
normalized_render=re.sub(r'\s+','',render_text)
test('every-original-caption-visible-in-render',all(re.sub(r'\s+','',caption) in normalized_render for caption in captions) and all(any(s.has_text_frame and s.text==caption for s in reloaded.slides[0].shapes) for caption in captions),{'captions':captions,'render_text_path':'render/deck.txt','comparison':'native caption exact; PDF extraction whitespace only normalized; actual text not rewritten'})
deck={'slides':specs,'style':style,'width_inches':13.333333,'height_inches':7.5}
mapping={'objects':[{'logical_object_id':e['object_id']} for s in specs for e in s['elements']]}
qa_issues,qa_checks,coverage=audit(deck,reloaded,mapping);package_issues,package_checks=package_audit(root/'gallery.pptx')
save('gallery-native-quality.json',{'issues':qa_issues+package_issues,'checks':qa_checks+package_checks,'coverage':coverage,'system_acceptance':'not_passed'})
test('native-product-quality-audit',not qa_issues and not package_issues,{'issue_codes':[i['code'] for i in qa_issues+package_issues],'checks':qa_checks+package_checks,'coverage':coverage})
hashes={str(p.relative_to('/workspace')):hashlib.sha256(p.read_bytes()).hexdigest() for p in [Path('/workspace/src/yoloongppt/visual.py'),Path('/workspace/src/yoloongppt/gallery_layout.py'),Path('/workspace/contracts/visual-model-fields.schema.json'),Path('/workspace/validation/verify-gallery.py')]}
report={'passed':all(c['passed'] for c in checks),'checks':checks,'fixture_identity':'artificial generated raster fixture, no user/AI E2E representation','native_cases':native_cases,'source_hashes':hashes,'model_calls':0,'render_calls':0 if resume else 1,'total_render_calls_for_fixture':1,'reused_existing_native_render':resume,'system_acceptance':'not_passed','remaining':['Independent image relevance/visual quality model review','Real 8-image AI full-pipeline AC-012','Face/object-aware crop and perceptual duplicate handling'],'output_path':str(root)}
save('gallery-runtime.json',report)
print(json.dumps({'passed':report['passed'],'checks':len(checks),'output_path':str(root)},ensure_ascii=False))
