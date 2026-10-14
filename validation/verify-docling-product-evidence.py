"""Validate saved real source-generation and scoped recovery, no new model calls."""
import hashlib,json
from pathlib import Path
from uuid import uuid4

root=Path('/runtime/docling-product-evidence')/uuid4().hex;root.mkdir(parents=True)
checks=[]
def load(folder,name):return json.loads((folder/name).read_text('utf-8'))
def check(name,passed,observation=None):
    checks.append({'check':name,'passed':bool(passed),'observed':observation})
    if not passed:raise AssertionError(name)
for kind,origin_id,final_id in [
    ('docx','361dcc5e-5276-4564-8157-f664ecaf46ec','22a5a93b-5a22-40e0-80e4-18a47e0c7eff'),
    ('xlsx','3f6fe040-dae7-4e17-92d7-c0184de48978','aa6b5914-83f7-4810-b81c-a3a6ffa5d594')]:
    origin=Path('/runtime/runs')/origin_id;final=Path('/runtime/runs')/final_id
    task=load(origin,'task.json');source=task['sources'][0];raw=Path(source['locator']).read_bytes()
    inputs=load(origin,'fact-source-result.json')['evidence']
    check(kind+'-real-original-file',source['kind']==kind and all(e['source_hash']==hashlib.sha256(raw).hexdigest() for e in inputs))
    check(kind+'-actual-docling-parser',all(e['parser_version'].startswith('docling-2.138.0:') for e in inputs))
    check(kind+'-real-fact-and-content-calls',all(load(origin,n)['real_model_call'] for n in ['fact-model-call.json','model-call.json']))
    check(kind+'-real-independent-final-review',load(final,'review-model-call.json')['real_model_call'])
    check(kind+'-pptx-pdf-png', (final/'deck.pptx').is_file() and any(a['type']=='pdf' for a in load(final,'render-report.json')['artifacts']) and any(a['type']=='png' for a in load(final,'render-report.json')['artifacts']))
    check(kind+'-qa-no-p0',load(final,'quality-report.json')['p0_issue_count']==0)
    check(kind+'-full-system-not-falsely-promoted',load(final,'result.json')['completion']['final_allowed'] is False)
    if kind=='docx':
        check('docx-qa-only-pptx-original-bytes',(origin/'deck.pptx').read_bytes()==(final/'deck.pptx').read_bytes())
        render_files=[p for p in (origin/'render').rglob('*') if p.is_file()]
        check('docx-qa-only-all-render-bytes',all(p.read_bytes()==(final/'render'/p.relative_to(origin/'render')).read_bytes() for p in render_files))
    else:
        raw_text=[e['raw_text'] for e in inputs]
        check('xlsx-actual-dates-percent-units-formula',all(x in raw_text for x in ['2026-01-01','2026-02-01','2026-03-01','8%','9%','10%','3750 千瓦时','SUM(B2:B4)']))
        old=load(origin,'model-response-effective.json');new=load(final,'reference-repair-response.json')
        for a,b in zip(old['slides'],new['slides']):
            a.pop('evidence_refs');b.pop('evidence_refs')
        check('xlsx-reference-revision-entire-copy-preserved',old==new)
        check('xlsx-actual-literal-selected-reference-binding',bool(load(final,'reference-repair-literal-bindings.json')['bindings']))
        from pptx import Presentation
        deck=Presentation(final/'deck.pptx')
        check('xlsx-native-table-chart-editable',any(s.has_table for p in deck.slides for s in p.shapes) and any(s.has_chart for p in deck.slides for s in p.shapes))
(root/'checks.json').write_text(json.dumps({'checks':checks,'passed':True,'fresh_model_calls':0,'new_renders':0,'scope':'Normal DOCX and XLSX real product generation/recovery evidence; complex PDF and all original AC remain incomplete.'},ensure_ascii=False,indent=2))
print(json.dumps({'passed':len(checks),'evidence':str(root)}))
