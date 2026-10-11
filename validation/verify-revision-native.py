"""REV-001/002/004/005/010: bounded multi-page native edits and atomic rejection.

Reuse the fixed native-object fixture; this is component evidence for AC-008/009,
not a fresh AI run, rendered review, or complete system acceptance.
"""
import copy
import hashlib
import io
import json
import posixpath
import time
import zipfile
from pathlib import Path
from unittest.mock import patch as mock_patch

from lxml import etree
from PIL import Image
from pptx import Presentation
from yoloongppt.errors import TaskError
from yoloongppt.existing_deck import NS, apply_patch, inspect_deck
from yoloongppt.native_objects import add_object


folder=Path('/runtime/revision-native-validation')/str(time.time_ns())
folder.mkdir(parents=True)
fixture=Path('/workspace/validation/native-object-artifacts/fixture.pptx')
checks=[]
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()


def check(name, condition, observation=None):
    checks.append({'name':name,'passed':bool(condition),'observation':observation})
    assert condition,name


def members(path):
    with zipfile.ZipFile(path) as archive:
        return {n:archive.read(n) for n in archive.namelist()}


def write_package(path,parts):
    with zipfile.ZipFile(path,'x',compression=zipfile.ZIP_DEFLATED) as archive:
        for name,blob in parts.items(): archive.writestr(name,blob)


def rejected(name, code, source_path, patches, extra=None):
    original_hash=sha(source_path)
    output=folder/(name+'.pptx')
    request={'expected_sha256':original_hash,'patches':patches,**(extra or {})}
    try: apply_patch(source_path,output,request)
    except TaskError as error:
        check(name,error.code==code and not output.exists() and not output.with_name(output.name+'.pending').exists(),{'code':error.code})
        check(name+'_source_unchanged',sha(source_path)==original_hash)
    else: raise AssertionError(name+' was not rejected')


fixture_hash=sha(fixture)
prs=Presentation(fixture)
second=prs.slides[1]
image2=folder/'image2.png';replace1=folder/'replace1.png';replace2=folder/'replace2.png'
Image.new('RGB',(80,60),'green').save(image2)
Image.new('RGB',(64,48),'red').save(replace1)
Image.new('RGB',(72,54),'yellow').save(replace2)
pic2=add_object(second,{'kind':'image','path':str(image2),'bounds':[1,3,1,1]})
chart2=add_object(second,{'kind':'chart','bounds':[5,1,4,3],'data':{'categories':['旧一','旧二'],'series':[{'name':'旧系列','values':[2,4]}]}})
table2=add_object(second,{'kind':'table','bounds':[3,4,3,1],'cells':[['甲','乙'],['丙','丁']]})
second.notes_slide.notes_text_frame.text='第二页原备注'
third=prs.slides.add_slide(prs.slide_layouts[6])
add_object(third,{'kind':'text','bounds':[1,1,4,1],'text':'第三页完全保留'})
third.notes_slide.notes_text_frame.text='第三页保留备注'
source=folder/'two-page-edit-fixture.pptx';prs.save(source)
# The reusable fixture intentionally contains an unreferenced vendor part.
# Keep its exact original bytes in the extended fixture used by the patcher.
with zipfile.ZipFile(fixture) as base,zipfile.ZipFile(source,'a') as archive:
    if 'customXml/vendor-unknown.xml' not in archive.namelist():
        archive.writestr('customXml/vendor-unknown.xml',base.read('customXml/vendor-unknown.xml'))
before=inspect_deck(source);source_hash=sha(source)
first,second_model,third_model=before['slides']
text1=next(o for o in first['objects'] if o['name']=='标题')
text2=next(o for o in second_model['objects'] if o['kind']=='sp')
pic1=next(o for o in first['objects'] if o['kind']=='pic')
chart1=next(o for o in first['objects'] if 'chart_part' in o)
table1=next(o for o in first['objects'] if 'table' in o)
group=next(o for o in first['objects'] if o['kind']=='grpSp')
group_child=first['objects'][first['objects'].index(group)+1]
lookup2={o['shape_id']:o for o in second_model['objects']}
pic2=lookup2[pic2.shape_id];chart2=lookup2[chart2.shape_id];table2=lookup2[table2.shape_id]
native_patch=lambda item,operation,value:{'object_id':item['object_id'],'operation':operation,'value':value}
data1={'categories':['一季','二季','三季'],'series':[{'name':'销售','values':[11,12,13]},{'name':'成本','values':[4,5,6]}]}
data2={'categories':['上期','本期'],'series':[{'name':'能耗','values':[7,9]}]}
patches=[
    native_patch(chart1,'chart_data',data1),
    native_patch(text1,'text','第一页修改后的文案'),
    native_patch(chart2,'chart_data',data2),
    native_patch(text2,'text','第二页修改后的文案'),
    native_patch(pic1,'image_replace',{'path':str(replace1)}),
    native_patch(pic2,'image_replace',{'path':str(replace2)}),
    native_patch(text1,'format',{'font_size':28,'bold':True,'color':'223344'}),
    native_patch(text2,'geometry',[1.2,1.1,4.1,1.2]),
    native_patch(table1,'table_cell',{'row':0,'column':1,'text':'第一表新值'}),
    native_patch(table2,'table_cell',{'row':0,'column':0,'text':'第二表新甲'}),
    native_patch(table2,'table_cell',{'row':1,'column':1,'text':'第二表新丁'}),
    native_patch(pic1,'alt_text','第一张替换图'),
    native_patch(text2,'notes','第二页修改后的备注'),
]
output=folder/'edited.pptx'
report=apply_patch(source,output,{'expected_sha256':source_hash,'patches':patches})
old,new=members(source),members(output)
after=report['after'];objects={o['object_id']:o for s in after['slides'] for o in s['objects']}
check('three_pages_and_stable_ids',[o['object_id'] for s in before['slides'] for o in s['objects']]==[o['object_id'] for s in after['slides'] for o in s['objects']] and len(after['slides'])==3)
check('first_page_text_changed',objects[text1['object_id']]['text']=='第一页修改后的文案')
check('second_page_text_changed',objects[text2['object_id']]['text']=='第二页修改后的文案')
check('selected_text_remains_native_editable',all(objects[t['object_id']]['kind']=='sp' for t in [text1,text2]))
edited_prs=Presentation(output)
edited_title=next(s for s in edited_prs.slides[0].shapes if s.shape_id==text1['shape_id'])
font=edited_title.text_frame.paragraphs[0].runs[0].font
check('native_format_applied',font.size.pt==28 and font.bold and str(font.color.rgb)=='223344')
check('alt_text_applied',objects[pic1['object_id']]['alt_text']=='第一张替换图')
check('separate_table_cells_applied',objects[table1['object_id']]['table'][0][1]=='第一表新值' and objects[table2['object_id']]['table']==[['第二表新甲','乙'],['丙','第二表新丁']])
check('geometry_applied',objects[text2['object_id']]['bounds_emu']==[round(v*914400) for v in [1.2,1.1,4.1,1.2]])
check('image1_bytes_applied',new[pic1['image_relationship']['target']]==replace1.read_bytes())
check('image2_bytes_applied',new[pic2['image_relationship']['target']]==replace2.read_bytes())
check('notes_changed_only_selected_slide',after['slides'][1]['notes'][0]['text'].endswith('第二页修改后的备注') and after['slides'][0]['notes']==before['slides'][0]['notes'] and after['slides'][2]['notes']==before['slides'][2]['notes'])
for index,(chart,data) in enumerate([(chart1,data1),(chart2,data2)],1):
    series=objects[chart['object_id']]['chart']['series']
    check(f'chart{index}_cache_applied',[s['categories'] for s in series]==[data['categories']]*len(data['series']) and [[float(v) for v in s['values']] for s in series]==[s['values'] for s in data['series']] and [s['name'] for s in series]==[[s['name']] for s in data['series']])
    relname=chart['chart_part'].rsplit('/',1)[0]+'/_rels/'+chart['chart_part'].rsplit('/',1)[1]+'.rels'
    relationship=next(r for r in etree.fromstring(new[relname]) if r.get('Type').endswith('/package'))
    workbook=posixpath.normpath(posixpath.join(posixpath.dirname(chart['chart_part']),relationship.get('Target')))
    with zipfile.ZipFile(io.BytesIO(new[workbook])) as archive:
        ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        strings=[''.join(s.itertext()) for s in etree.fromstring(archive.read('xl/sharedStrings.xml'))]
        cells=etree.fromstring(archive.read('xl/worksheets/sheet1.xml')).xpath('//x:c',namespaces=ns)
        values={c.get('r'):strings[int(c.find('x:v',ns).text)] if c.get('t')=='s' else float(c.find('x:v',ns).text) for c in cells if c.find('x:v',ns) is not None}
        check(f'chart{index}_workbook_matches_categories',[values[f'A{r+2}'] for r in range(len(data['categories']))]==data['categories'])
        check(f'chart{index}_workbook_matches_series',all(values[f'{chr(66+s)}1']==entry['name'] and [values[f'{chr(66+s)}{r+2}'] for r in range(len(data['categories']))]==entry['values'] for s,entry in enumerate(data['series'])))
check('complete_changed_parts_audit',set(report['changed_parts'])=={n for n in old if old[n]!=new[n]} and all(report['part_hashes_before'][n]==hashlib.sha256(old[n]).hexdigest() and report['part_hashes'][n]==hashlib.sha256(new[n]).hexdigest() for n in old))
check('untouched_members_byte_identical',all(old[n]==new[n] for n in report['preserved_parts']))
check('third_slide_byte_identical',old[third_model['part']]==new[third_model['part']])
check('masters_layouts_themes_unknown_bytes_preserved',all(old[n]==new[n] for n in old if any(prefix in n for prefix in ('slideMasters/','slideLayouts/','theme/','customXml/'))))
check('untargeted_group_preserved',objects[group['object_id']]==group and objects[group_child['object_id']]==group_child)
targeted={p['object_id'] for p in patches}
def shape_xml(blob):
    root=etree.fromstring(blob)
    result={}
    for node in root.xpath('./p:cSld/p:spTree//*[self::p:sp or self::p:pic or self::p:graphicFrame or self::p:cxnSp or self::p:grpSp]',namespaces=NS):
        identity=node.xpath('./p:nvSpPr/p:cNvPr|./p:nvPicPr/p:cNvPr|./p:nvGraphicFramePr/p:cNvPr|./p:nvCxnSpPr/p:cNvPr|./p:nvGrpSpPr/p:cNvPr',namespaces=NS)
        if identity:result[int(identity[0].get('id'))]=etree.tostring(node)
    return result
check('all_untargeted_shape_xml_identical',all(a[sid]==b[sid] for slide in before['slides'] for a,b in [(shape_xml(old[slide['part']]),shape_xml(new[slide['part']]))] for sid in a if f'pptx:s{slide["slide_id"]}:shape{sid}' not in targeted))
notes_part=second_model['notes'][0]['part']
notes_before,notes_after=shape_xml(old[notes_part]),shape_xml(new[notes_part])
notes_root=etree.fromstring(old[notes_part])
note_body=int(notes_root.xpath('//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]/p:nvSpPr/p:cNvPr/@id',namespaces=NS)[0])
check('untargeted_notes_shape_xml_identical',all(notes_before[sid]==notes_after[sid] for sid in notes_before if sid!=note_body))
check('history_contains_native_before_after',len(report['history'])==len(patches) and all(h['before']['object_id']==h['after']['object_id'] and h['after']==objects[h['object_id']] for h in report['history']))
notes_history=next(h for h in report['history'] if h['operation']=='notes')
check('notes_history_before_after',notes_history['before_notes']==before['slides'][1]['notes'] and notes_history['after_notes']==after['slides'][1]['notes'])
check('scope_exact_two_pages',report['scope']['slide_parts']==[first['part'],second_model['part']])
check('original_and_reused_fixture_unchanged',sha(source)==source_hash and sha(fixture)==fixture_hash)

# Invalid second operations cannot publish the successful first text patch.
invalid=folder/'damaged.png';invalid.write_bytes(b'not a png')
rejected('invalid_second_image_atomic','IMAGE_FILE_INVALID',source,[native_patch(text1,'text','不可发布'),native_patch(pic1,'image_replace',{'path':str(invalid)})])
rejected('invalid_second_chart_atomic','NATIVE_CHART_INVALID',source,[native_patch(text1,'text','不可发布'),native_patch(chart1,'chart_data',{'categories':['一','二'],'series':[{'name':'错长','values':[1]}]})])
rejected('duplicate_text_conflict','REVISION_PATCH_CONFLICT',source,[native_patch(text1,'text','一'),native_patch(text1,'text','二')])
rejected('duplicate_cell_conflict','REVISION_PATCH_CONFLICT',source,[native_patch(table2,'table_cell',{'row':0,'column':0,'text':'一'}),native_patch(table2,'table_cell',{'row':0,'column':0,'text':'二'})])
rejected('same_slide_notes_conflict','REVISION_PATCH_CONFLICT',source,[native_patch(text2,'notes','一'),native_patch(pic2,'notes','二')])
with mock_patch('yoloongppt.existing_deck.local_asset',side_effect=AssertionError('Locked batch must not read a replacement asset')):
    rejected('last_patch_lock_precedes_asset_read','REVISION_OBJECT_LOCKED',source,[native_patch(pic1,'image_replace',{'path':str(invalid)}),native_patch(text2,'text','锁定不可改')],{'locked_object_ids':[text2['object_id']]})
rejected('locked_group_protects_child','REVISION_OBJECT_LOCKED',source,[native_patch(group_child,'text','不可改')],{'locked_object_ids':[group['object_id']]})

def altered(name, fn):
    parts=copy.copy(old);fn(parts);path=folder/(name+'.pptx');write_package(path,parts);return path


def reuse_chart_rid(parts):
    root=etree.fromstring(parts[first['part']]);tree=root.find('p:cSld/p:spTree',NS)
    frame=next(n for n in tree if n.xpath('.//c:chart',namespaces=NS))
    duplicate=copy.deepcopy(frame);identity=duplicate.find('p:nvGraphicFramePr/p:cNvPr',NS)
    identity.set('id','900');identity.set('name','共享图表未选对象');tree.append(duplicate)
    parts[first['part']]=etree.tostring(root)


shared_chart=altered('shared-chart-rid-fixture',reuse_chart_rid)
rejected('shared_chart_same_rid_rejected','SHARED_CHART_UNSUPPORTED',shared_chart,[native_patch(text1,'text','不可发布'),native_patch(chart1,'chart_data',data1)])


def share_workbook(parts):
    names=[c['chart_part'].rsplit('/',1)[0]+'/_rels/'+c['chart_part'].rsplit('/',1)[1]+'.rels' for c in [chart1,chart2]]
    roots=[etree.fromstring(parts[n]) for n in names]
    source_target=next(r.get('Target') for r in roots[0] if r.get('Type').endswith('/package'))
    next(r for r in roots[1] if r.get('Type').endswith('/package')).set('Target',source_target)
    parts[names[1]]=etree.tostring(roots[1])


shared_book=altered('shared-workbook-fixture',share_workbook)
rejected('shared_chart_workbook_rejected','SHARED_CHART_WORKBOOK_UNSUPPORTED',shared_book,[native_patch(chart1,'chart_data',data1)])


def reuse_image_rid(parts):
    root=etree.fromstring(parts[first['part']]);tree=root.find('p:cSld/p:spTree',NS)
    picture=next(n for n in tree if etree.QName(n).localname=='pic')
    duplicate=copy.deepcopy(picture);identity=duplicate.find('p:nvPicPr/p:cNvPr',NS)
    identity.set('id','901');identity.set('name','共享图片未选对象');tree.append(duplicate)
    parts[first['part']]=etree.tostring(root)


shared_image=altered('shared-image-rid-fixture',reuse_image_rid)
rejected('shared_image_same_rid_rejected','SHARED_IMAGE_UNSUPPORTED',shared_image,[native_patch(pic1,'image_replace',{'path':str(replace1)})])


def shared_relationship(parts, relation_type):
    names=[s['part'].rsplit('/',1)[0]+'/_rels/'+s['part'].rsplit('/',1)[1]+'.rels' for s in [first,second_model]]
    roots=[etree.fromstring(parts[n]) for n in names]
    source_target=next(r.get('Target') for r in roots[0] if r.get('Type').endswith('/'+relation_type))
    next(r for r in roots[1] if r.get('Type').endswith('/'+relation_type)).set('Target',source_target)
    parts[names[1]]=etree.tostring(roots[1])


shared_notes=altered('shared-notes-fixture',lambda parts:shared_relationship(parts,'notesSlide'))
rejected('shared_notes_rejected','SHARED_NOTES_UNSUPPORTED',shared_notes,[native_patch(text2,'notes','不可发布')])
shared_image_pages=altered('shared-image-two-pages-fixture',lambda parts:shared_relationship(parts,'image'))
rejected('shared_image_two_pages_rejected','SHARED_IMAGE_UNSUPPORTED',shared_image_pages,[native_patch(pic1,'image_replace',{'path':str(replace1)})])

manifest={'passed':all(c['passed'] for c in checks),'checks':checks,'fixture':'extended fixed native-object fixture; synthetic component only',
          'reused_fixture_sha256':fixture_hash,'source_sha256':source_hash,'output_sha256':sha(output),
          'requirement_ids':['REV-001','REV-002','REV-004','REV-005','REV-010','AC-008','AC-009'],
          'task_ids':['TASK-REV-001','TASK-REV-002','TASK-REV-004','TASK-REV-005','TASK-REV-010'],
          'changed_parts':report['changed_parts'],'system_acceptance':'not_passed','qa':'not_executed','model_calls':0,
          'limitations':['Shared chart/media/workbook edits are explicitly rejected; no copy-on-write implementation.',
                         'Semantic selection, PowerPoint opening, rendered visual/factual QA and full AC-008/009 remain outside this component verification.'],
          'source_versions':{str(p.relative_to('/workspace')):sha(p) for p in [Path('/workspace/src/yoloongppt/existing_deck.py'),Path('/workspace/validation/verify-revision-native.py')]}}
(folder/'native-revision-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
(folder/'native-revision-checks.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'passed':manifest['passed'],'checks':len(checks),'evidence_path':str(folder),'changed_parts':report['changed_parts']},ensure_ascii=False))
