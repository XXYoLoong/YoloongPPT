"""Bounded changed-component checks; live model artifacts passed as run IDs."""
import copy
import hashlib
import json
import platform
import subprocess
import sys
import urllib.error
import urllib.request
from importlib.metadata import distribution, version
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from yoloongppt.errors import TaskError, trace_id
from yoloongppt.planning import preflight, compile_deck
from yoloongppt.revision import revise
from yoloongppt.generation import resume
from yoloongppt.evidence import EvidenceStore
from yoloongppt.providers import DeepSeek
from yoloongppt.schemas import SchemaRegistry, ROOT
from yoloongppt.writer import capacity

cases = []
def record(name, condition, observation):
    assert condition, name
    cases.append({'id': name, 'passed': True, 'observation': observation})

def fails(name, fn, code):
    try:fn()
    except TaskError as e:
        record(name, e.code == code, {'error_code': e.code})
    else:raise AssertionError(name)

schemas = SchemaRegistry()
task = json.loads((ROOT/'validation/generation-task.json').read_text(encoding='utf-8'))
count, style, warnings = preflight(task)
record('preflight_explicit_hard_page_budget', count == 10 and not warnings, 'Explicit ten-page hard constraint consumed.')
for name, mutation, code in [
    ('unknown_hard_constraint', lambda t: t['constraints']['hard'].append({'constraint_id':'unknown','field':'invented','value':1,'source':{'kind':'user_instruction','reference':'boundary'}}), 'HARD_CONSTRAINT_UNSUPPORTED'),
    ('conflicting_hard_constraint', lambda t: t['constraints']['hard'].append({**t['constraints']['hard'][0], 'constraint_id':'second','value':11}), 'HARD_CONSTRAINT_CONFLICT'),
    ('unsupported_template', lambda t: t.update(template_ref='not-imported'), 'GENERATION_TEMPLATE_UNSUPPORTED'),
    ('unsupported_font', lambda t: t.update(style={'font':'Uninstalled Font'}), 'GENERATION_STYLE_UNSUPPORTED'),
    ('unsupported_style', lambda t: t.update(style={'title_size':8}), 'GENERATION_STYLE_UNSUPPORTED'),
    ('unsupported_output', lambda t: t.update(output={'formats':['gif']}), 'OUTPUT_FORMAT_UNSUPPORTED'),
    ('unsupported_runtime_option', lambda t: t['runtime_preferences'].update(ignored_option=True), 'GENERATION_OPTION_UNSUPPORTED'),
    ('unsupported_generation_mode', lambda t: t['route'].update(mode='modify_existing_deck'), 'GENERATION_ROUTE_UNSUPPORTED')]:
    changed = copy.deepcopy(task); mutation(changed)
    fails(name, lambda: preflight(changed), code)
fails('text_capacity_no_truncation', lambda: capacity(['过长文字'*1000], [0,0,1,1], 22), 'TEXT_CAPACITY_EXCEEDED')
fails('checkpoint_path_rejected', lambda: resume('../secrets', schemas, EvidenceStore(), trace_id()), 'CHECKPOINT_ID_INVALID')
fails('provider_redirect_target_rejected', lambda: DeepSeek({'provider':'deepseek','model':'deepseek-flash','base_url':'https://example.com'}), 'PROVIDER_UNSUPPORTED')
base_id, revision_id = sys.argv[1:3]
base = Path('/runtime/runs')/base_id; revised = Path('/runtime/runs')/revision_id
load = lambda p: json.loads(p.read_text(encoding='utf-8'))
deck = load(base/'deck-spec.json'); proposal = load(base/'model-response-effective.json') if (base/'model-response-effective.json').exists() else load(base/'model-response.json')
compiled, execution = compile_deck(proposal, style, trace_id(), deck)
schemas.validate('deck-execution.schema.json', compiled);schemas.validate('deck-execution.schema.json', execution)
record('checkpoint_entity_ids_and_DAG', [s['slide_id'] for s in compiled['slides']] == [s['slide_id'] for s in deck['slides']] and [e['object_id'] for s in compiled['slides'] for e in s['elements']] == [e['object_id'] for s in deck['slides'] for e in s['elements']], 'IDs retained; DAG schema validates after recovery.')
presentation = Presentation(base/'deck.pptx')
record('actual_ten_slide_pptx', len(presentation.slides) == 10, {'slides': len(presentation.slides)})
kinds = {'tables':sum(s.has_table for slide in presentation.slides for s in slide.shapes),'charts':sum(s.has_chart for slide in presentation.slides for s in slide.shapes),'text_boxes':sum(s.has_text_frame for slide in presentation.slides for s in slide.shapes)}
record('native_table_chart_text', kinds['tables'] >= 1 and kinds['charts'] >= 1 and kinds['text_boxes'] >= 20, kinds)
with ZipFile(base/'deck.pptx') as z:
    record('embedded_chart_workbook', any(n.startswith('ppt/embeddings/') and n.endswith('.xlsx') for n in z.namelist()), 'Native chart has editable xlsx data.')
model = load(base/'model-call.json')
record('real_model_not_mock', model['real_model_call'] and model['requested_model'] == 'deepseek-flash' and model['usage']['total_tokens'] > 0, model)
record('real_visual_review', load(base/'review-model-call.json')['real_model_call'] and load(base/'quality-report.json')['model_review']['checked_slide_orders'] == list(range(10)), 'Actual ten rendered PNGs reviewed by configured image-capable model.')
renders = load(base/'render-report.json')['artifacts']
record('actual_render_hashes', len([a for a in renders if a['type']=='png']) == 10 and all(hashlib.sha256((base/a['path']).read_bytes()).hexdigest()==a['hash'] for a in renders), 'One PDF plus ten actual hashed PNG pages.')
objects = load(base/'object-map.json')
record('real_object_map', len(objects['objects']) == sum(len(s.shapes) for s in presentation.slides), {'mapped_objects':len(objects['objects'])})
history = load(revised/'revision-history.json')
with ZipFile(base/'deck.pptx') as a, ZipFile(revised/'deck.pptx') as b:
    changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
record('local_revision_preserves_other_parts', changed == ['ppt/slides/slide1.xml'] and history['entity_ids_preserved'], {'changed_parts':changed})
record('local_revision_object_map_stable', load(base/'object-map.json') == load(revised/'object-map.json'), 'Logical IDs/backend object IDs unchanged.')
record('revised_actual_QA', load(revised/'quality-report.json')['p0_issue_count']==0, 'Actual revised-page review, unchanged-page evidence reuse, deterministic checks. Not full AC acceptance.')
record('acceptance_not_inflated', load(base/'quality-report.json')['acceptance']['AC-001']=='not_passed' and load(revised/'quality-report.json')['status']=='partial', 'Full decision graph and additional gates remain unfinished.')
def http(path, body):
    req=urllib.request.Request('http://127.0.0.1:8000'+path,data=json.dumps(body).encode('utf-8'),headers={'Content-Type':'application/json'})
    try:r=urllib.request.urlopen(req,timeout=20)
    except urllib.error.HTTPError as e:r=e
    return r.status,json.loads(r.read()),r.headers
bad=copy.deepcopy(task);bad['template_ref']='unsupported'
status,body,headers=http('/generate',bad)
record('API_generate_same_preflight', status==422 and body['error']['code']=='GENERATION_TEMPLATE_UNSUPPORTED' and headers['X-Trace-Id']==body['trace_id'], 'Shared core rejects unsupported template before model call.')
status,body,_=http('/revise/'+base_id,{'object_id':'object_ffffffff-ffff-4fff-afff-ffffffffffff','replacement_text':['不存在'], 'reason':'boundary'})
record('API_revision_unknown_object', status==422 and body['error']['code']=='REVISION_OBJECT_NOT_FOUND', 'No mutation when target missing.')
cli=subprocess.run([sys.executable,'-m','yoloongppt','resume','../secrets'],capture_output=True,text=True,encoding='utf-8')
record('CLI_checkpoint_error_JSON', cli.returncode==1 and json.loads(cli.stdout)['error']['code']=='CHECKPOINT_ID_INVALID', 'Structured failure, no traceback/credentials.')
licenses=[]
for name in ['python-pptx','lxml','Pillow','XlsxWriter']:
    dist=distribution(name)
    found=[]
    for file in dist.files or []:
        if any(t in str(file).lower() for t in ['license','copying']):
            p=Path(dist.locate_file(file))
            if p.is_file():found.append({'path':str(file),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
    record('license_files_'+name,bool(found),{'version':version(name),'files':found})
    licenses.append({'name':name,'version':version(name),'files':found})
paths=[*sorted((ROOT/'src/yoloongppt').glob('*.py')),ROOT/'requirements.lock',ROOT/'Dockerfile.app',ROOT/'compose.yaml',*[ROOT/'contracts'/n for n in ['generation-model.schema.json','deck-execution.schema.json','revision-request.schema.json']],ROOT/'validation/generation-task.json',ROOT/'validation/fixtures/节能试点材料.md']
report={'passed':True,'cases':cases,'base_run_id':base_id,'revised_run_id':revision_id,'scope':'Changed generation/revision components and actual artifacts; full AC-001/DEC graph not passed.',
        'python':platform.python_version(),'licenses':licenses,'consumed_file_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}}
Path('/runtime/generation-runtime.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':True,'checks':len(cases),'base_run_id':base_id,'revised_run_id':revision_id}))
