"""DEC-001 fixed cases and actual CLI/API/registered writer execution, bounded scope."""
import copy
import hashlib
import json
import subprocess
import sys
import urllib.error
import urllib.request
from pathlib import Path

from pptx import Presentation
from yoloongppt.atomic import AtomicRegistry
from yoloongppt.errors import TaskError, trace_id
from yoloongppt.planning import compile_deck, preflight
from yoloongppt.routing import decide, runtime_snapshot
from yoloongppt.schemas import ROOT, SchemaRegistry
from yoloongppt.writer import execute

schemas = SchemaRegistry()
cases = []
traces = []
def record(name, condition, result):
    assert condition, name
    cases.append({'id': name, 'passed': True, 'actual': result})

def fails(name, action, code):
    try:action()
    except TaskError as e:record(name, e.code == code, e.code)
    else:raise AssertionError(name)

fixtures = json.loads((ROOT/'validation/route-cases.json').read_text(encoding='utf-8'))
for case in fixtures['cases']:
    result = decide(case['input'], schemas, trace_id())
    for key, value in case['expected'].items():assert result['route'][key] == value, (case['case_id'], key)
    decision = result['decision_trace']
    assert decision['input']['raw_request'] == case['input']['raw_request']
    assert all(c['reasons'] and c['constraints'] for c in decision['candidates'])
    assert (decision['selected'] is not None) == (result['route']['mode'] is not None)
    traces.append({'case_id': case['case_id'], 'trace': decision})
    record(case['case_id'], True, result['route'])

normal = fixtures['cases'][2]['input']
r = decide(normal, schemas, trace_id())['route']
record('unsupported_route_not_replaced', r['mode'] == 'create_from_reference_style' and not r['trace']['executable'], r)
changed = copy.deepcopy(fixtures['cases'][0]['input'])
changed['raw_request']['request']['requested_mode'] = 'unknown-mode'
fails('unknown_mode_rejected', lambda: decide(changed, schemas, trace_id()), 'REQUESTED_MODE_UNKNOWN')
bad = copy.deepcopy(normal);bad['runtime_snapshot'] = {'mode_support': {}}
fails('incomplete_snapshot_rejected', lambda: decide(bad, schemas, trace_id()), 'RUNTIME_SNAPSHOT_INVALID')
record('actual_runtime_probe', runtime_snapshot(schemas)['tools']['python-pptx'] == '1.0.2', runtime_snapshot(schemas)['provenance'])

atomic = AtomicRegistry(schemas)
task = json.loads((ROOT/'validation/generation-task.json').read_text(encoding='utf-8'))
_, style, _ = preflight(task)
old = Path('/runtime/runs/f40df117-e1fb-4bbf-b813-55b1ddf78603')
proposal = json.loads((old/'model-response-effective.json').read_text(encoding='utf-8'))
deck, plan = compile_deck(proposal, style, trace_id(), registry=atomic)
record('registered_atomic_ids', all(c['capability_id'].startswith('capability_') and c['implementation_id'].startswith('implementation_') for c in plan['calls']), {'calls':len(plan['calls']), 'definitions':len(atomic.catalog['capabilities'])})
folder = Path('/runtime/route-validation');folder.mkdir(exist_ok=True)
objects, executed = execute(deck, plan, folder/'registered-deck.pptx', schemas)
record('actual_registered_native_writer', len(Presentation(folder/'registered-deck.pptx').slides) == 10 and len(objects['objects']) == 32, {'objects':len(objects['objects']), 'calls':len(executed['calls'])})
for call in executed['calls']:
    binding = next(c for c in plan['calls'] if c['call_id'] == call['call_id'])
    assert binding['implementation_id'] == call['implementation_id'] and binding['capability_id'] == call['capability_id']
tampered = copy.deepcopy(plan);tampered['calls'][-1]['implementation_id'] = 'implementation_ffffffff-ffff-4fff-afff-ffffffffffff'
target = folder/'invalid-must-not-exist.pptx'
assert not target.exists()
fails('invalid_binding_before_write', lambda: execute(deck, tampered, target, schemas), 'IMPLEMENTATION_BINDING_INVALID')
record('invalid_binding_no_saved_mutation', not target.exists(), 'All registered bindings checked before write.')
badcall = copy.deepcopy(plan['calls'][0]);badcall['inputs']['order'] = -1
fails('atomic_input_schema_enforced', lambda: atomic.check_call(badcall), 'ATOMIC_CONTRACT_INVALID')
badtable = copy.deepcopy(next(c for c in plan['calls'] if c['capability'] == 'add_table'))
badtable['inputs']['data'] = {}
fails('atomic_table_shape_enforced', lambda: atomic.check_call(badtable), 'ATOMIC_CONTRACT_INVALID')
badchart = copy.deepcopy(next(c for c in plan['calls'] if c['capability'] == 'add_chart'))
badchart['inputs']['data']['series'][0]['values'].pop()
fails('atomic_chart_length_enforced', lambda: atomic.check_call(badchart), 'ATOMIC_DATA_INVALID')

def http(path, body):
    req = urllib.request.Request('http://127.0.0.1:8000'+path, data=json.dumps(body).encode('utf-8'), headers={'Content-Type':'application/json'})
    try:r=urllib.request.urlopen(req,timeout=20)
    except urllib.error.HTTPError as e:r=e
    return r.status, json.loads(r.read())
status, api = http('/route', fixtures['cases'][0]['input'])
record('HTTP_standalone_node', status == 200 and api['route']['mode'] == 'create_from_scratch' and Path(api['artifact_root'],'decision-route.json').is_file(), {'run_id':api['run_id']})
inputfile = folder/'node-input.json';inputfile.write_text(json.dumps(fixtures['cases'][8]['input'],ensure_ascii=False),encoding='utf-8')
cli = subprocess.run([sys.executable,'-m','yoloongppt','route',str(inputfile)],capture_output=True,text=True,encoding='utf-8')
cliresult = json.loads(cli.stdout)
record('CLI_conflict_trace_saved', cli.returncode == 0 and cliresult['route']['route_status'] == 'needs_clarification' and cliresult['decision_trace']['error']['code'] == 'MODE_CLARIFICATION_REQUIRED', {'run_id':cliresult['run_id']})
conflict = copy.deepcopy(task);conflict['raw_request'] = fixtures['cases'][8]['input']['raw_request']
status, error = http('/generate', conflict)
record('generation_conflict_stops_before_model', status == 422 and error['error']['code'] == 'MODE_CLARIFICATION_REQUIRED', error['error']['code'])
runinfo = next(d for d in error['error']['details'] if 'artifact_root' in d)
failfolder = Path(runinfo['artifact_root'])
record('conflict_preserves_decision_no_model', (failfolder/'decision-route.json').is_file() and not (failfolder/'model-call.json').exists() and not (failfolder/'deck.pptx').exists(), {'run_id':runinfo['run_id']})
mapping = json.loads((ROOT/'validation/route-project-map.json').read_text(encoding='utf-8'))['projects']
record('five_project_mapping', {p['project'] for p in mapping} == {'P01','P02','P03','P04','P05'}, 'Five source-indexed counterparts; P02 explicit command/visual branch only, automatic intent router Not found.')
if len(sys.argv) > 1:
    run = Path('/runtime/runs')/sys.argv[1]
    load = lambda name: json.loads((run/name).read_text(encoding='utf-8'))
    decision = load('decision-route.json');schemas.validate('route-trace.schema.json', decision)
    current = load('task.json');effective = load('task-effective.json')
    record('E2E_raw_request_route_preserved', decision['input']['raw_request'] == current['raw_request'] and decision['output'] == effective['route'] and effective['route']['mode'] == 'create_from_materials', {'run_id':sys.argv[1], 'trace_id':decision['trace_id']})
    savedregistry = load('atomic-registry.json');execution = load('execution-plan.json');calls = load('execution-trace.json')['calls']
    schemas.validate('deck-execution.schema.json', execution)
    for call in execution['calls']:atomic.check_call(call)
    record('E2E_registered_calls_executed', len(calls) == len(execution['calls']) and all(c['capability_id'] in {a['capability_id'] for a in savedregistry['capabilities']} and c['implementation_id'] in {i['implementation_id'] for a in savedregistry['capabilities'] for i in a['implementations']} for c in calls), {'calls':len(calls)})
    actualdeck = load('deck-spec.json');objects = load('object-map.json')['objects'];sources = load('source-result.json')
    known_evidence = {e['evidence_id'] for e in sources['evidence']}
    for obj in objects:
        slide = next(s for s in actualdeck['slides'] if s['slide_id'] == obj['slide_id'])
        spec = next(e for e in slide['elements'] if e['object_id'] == obj['logical_object_id'])
        call = next(c for c in calls if c['call_id'] == obj['execution_call_id'])
        assert spec == call['input'] and set(obj['source_ref']) <= known_evidence and actualdeck['trace_id'] == decision['trace_id']
    record('E2E_object_to_spec_execution_source_route', True, {'objects':len(objects), 'node':'DEC-001', 'remaining_nodes':'DEC-002–040 not proven'})
    model = load('model-call.json');qa = load('quality-report.json')
    record('E2E_real_model_native_render_QA', model['real_model_call'] and load('review-model-call.json')['real_model_call'] and len(Presentation(run/'deck.pptx').slides) == 10 and qa['p0_issue_count'] == 0, {'response_id':model['response_id'],'tokens':model['usage']['total_tokens'],'P0':qa['p0_issue_count']})
    record('E2E_full_acceptance_not_inflated', qa['acceptance']['AC-001'] == 'not_passed' and qa['status'] == 'partial', qa['acceptance'])
paths = [*sorted((ROOT/'src/yoloongppt').glob('*.py')), *[ROOT/'contracts'/name for name in ['atomic-registry.schema.json','atomic-registry.catalog.json','route-request.schema.json','route-trace.schema.json','task-route.schema.json','task-spec.schema.json','deck-execution.schema.json']], ROOT/'validation/route-cases.json',ROOT/'validation/route-project-map.json']
report = {'passed':True,'requirement_ids':['DEC-001','SYS-007','SYS-005','TST-003'],'cases':cases,'traces':traces,'scope':'DEC-001 node and current native registry integration, not DEC-002–040 or all downstream modes/AC acceptance.', 'consumed_file_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}}
(folder/'execution-plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
(folder/'execution-trace.json').write_text(json.dumps(executed,ensure_ascii=False,indent=2),encoding='utf-8')
(folder/'object-map.json').write_text(json.dumps(objects,ensure_ascii=False,indent=2),encoding='utf-8')
Path('/runtime/route-runtime.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':True,'checks':len(cases),'normal_routes':8,'boundary_cases':10}))
