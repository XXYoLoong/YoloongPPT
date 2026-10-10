"""Bounded Docker checks: structured parser + native template, no model invocation."""
import hashlib
import io
import json
import zipfile
from pathlib import Path
from uuid import uuid4
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import xlsxwriter

from yoloongppt.assets import inspect_image, inventory
from yoloongppt.errors import TaskError
from yoloongppt.input_formats import parse_input, NS
from yoloongppt.schemas import SchemaRegistry, ROOT
from yoloongppt.templates import parse_template, query_layouts, instantiate_template

output=Path('/runtime/input-format-validation'); output.mkdir(parents=True,exist_ok=True)
schemas=SchemaRegistry(); cases=[]; source_id='source_'+str(uuid4()); asset_id='asset_'+str(uuid4())
def check(name, ok, details=None):
    assert ok,name
    cases.append({'id':name,'passed':True,'observed':details})
def parse(kind,raw):
    result=parse_input(kind,raw,source_id=source_id,asset_id=asset_id)
    schemas.validate('input-format-result.schema.json',result)
    for item in result['segments']:schemas.validate('source-anchor.schema.json',item['anchor'])
    return result
def reject(name,kind,raw,expected):
    try:parse(kind,raw)
    except TaskError as e:check(name,e.code==expected,e.code)
    else:raise AssertionError(name+' unexpectedly accepted')
def office(part,payload,extra=None):
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml','<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr(part,payload)
        for name,body in (extra or {}).items():z.writestr(name,body)
    return buffer.getvalue()

docx=office('word/document.xml',f'''<w:document xmlns:w="{NS['w']}" xmlns:r="{NS['r']}" xmlns:a="{NS['a']}"><w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>学校节能试点</w:t></w:r></w:p>
<w:p><w:hyperlink r:id="rId1"><w:r><w:t>试点2026年能耗降低8%。</w:t></w:r></w:hyperlink></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>指标</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>8%</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
<w:p><w:r><w:drawing><a:blip r:embed="rId2"/></w:drawing></w:r></w:p><w:sectPr/></w:body></w:document>''',
{'word/_rels/document.xml.rels':'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="hyperlink" Target="https://example.invalid/source" TargetMode="External"/><Relationship Id="rId2" Type="image" Target="media/image1.png"/></Relationships>',
 'word/styles.xml':f'<w:styles xmlns:w="{NS["w"]}"><w:style w:styleId="Heading1"><w:name w:val="heading 1"/></w:style></w:styles>'})
(output/'school-source.docx').write_bytes(docx)
d=parse('docx',docx)
check('docx_heading_table_hyperlink_image_preserved',d['document']['stories'][0]['blocks'][0]['kind']=='heading' and any(x['data'].get('kind')=='table_cell' for x in d['segments']) and d['document']['stories'][0]['blocks'][1]['hyperlinks'][0]['text']=='试点2026年能耗降低8%。')
check('docx_native_anchor_without_fake_page',d['segments'][0]['anchor']['locator']['page_location']['status']=='pending')
check('docx_external_link_not_fetched',d['document']['stories'][0]['relationships'][0]['retrieval']=='not_fetched')
check('docx_raw_original_hash',d['source_sha256']==hashlib.sha256(docx).hexdigest())
original=parse('docx',(ROOT/'AI_PPT_PDR_完整需求定义_V0.3.docx').read_bytes())
check('original_requirement_docx_real_parse',len(original['segments'])>308,{'segments':len(original['segments']),'stories':len(original['document']['stories'])})

buffer=io.BytesIO(); workbook=xlsxwriter.Workbook(buffer,{'in_memory':True}); sheet=workbook.add_worksheet('数据')
sheet.write_row(0,0,['年度','节能比例']); sheet.write_number(1,0,2026); sheet.write_number(1,1,8)
sheet.write_formula(2,1,'=B2*2',None,16);sheet.merge_range('D1:E1','合并标题');sheet.add_table('A1:B2',{'columns':[{'header':'年度'},{'header':'节能比例'}]});workbook.close()
xlsx=buffer.getvalue();(output/'school-data.xlsx').write_bytes(xlsx);x=parse('xlsx',xlsx)
check('xlsx_cell_anchor_and_formula_cache',any(s['anchor']['locator']['cell_ref']=='B3' and s['data']['formula']=='B2*2' and s['raw_text']=='16' for s in x['segments']))
check('xlsx_table_and_merge_preserved',len(x['document']['sheets'][0]['tables'])==1 and x['document']['sheets'][0]['merged_ranges']==['D1:E1'])
with zipfile.ZipFile(io.BytesIO(xlsx)) as z:
    extras={n:z.read(n) for n in z.namelist() if n not in {'[Content_Types].xml','xl/worksheets/sheet1.xml'}}
    xml=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    for c in xml.findall('.//s:c',NS):
        if c.get('r')=='B3':c.remove(c.find('s:v',NS))
    missing=office('xl/worksheets/sheet1.xml',ET.tostring(xml),extras)
# Repack retains workbook path while office() supplies content types.
m=parse('xlsx',missing);check('xlsx_missing_formula_cache_explicit',any(i['code']=='XLSX_FORMULA_CACHE_MISSING' for i in m['warnings']))
original_x=parse('xlsx',(ROOT/'AI_PPT_完整需求与任务矩阵_V0.3.xlsx').read_bytes())
check('original_requirement_xlsx_real_parse',len(original_x['document']['sheets'])==20,{'sheets':len(original_x['document']['sheets']),'cells':len(original_x['segments'])})
csv=parse('csv','名称,内容\n"多行\n中文",8%\n少列\n'.encode())
check('csv_quoted_multiline_and_ragged_preserved',csv['document']['rows'][1][0]=='多行\n中文' and any(w['code']=='CSV_RAGGED_ROWS' for w in csv['warnings']))
reject('csv_corrupt_quotes','csv',b'a,b\n"unterminated','CSV_CORRUPT')

image=Image.new('RGBA',(80,40),(255,0,0,128));buffer=io.BytesIO();image.save(buffer,format='PNG');png=buffer.getvalue()
(output/'source-image.png').write_bytes(png);i=parse('image',png);schemas.validate('asset-image-inspection.schema.json',i['document'])
check('image_pixels_palette_without_fake_ocr',i['document']['width']==80 and i['document']['height']==40 and i['segments'][0]['raw_text']=='' and i['document']['ocr_status']=='unsupported')
reject('image_corrupt','image',b'not an image','IMAGE_CORRUPT_OR_UNSUPPORTED')

deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='可编辑模板';slide.placeholders[1].text='学校节能试点2026年降低8%。'
slide.shapes.add_picture(io.BytesIO(png),Inches(1),Inches(4),width=Inches(1));table=slide.shapes.add_table(2,2,Inches(4),Inches(3),Inches(4),Inches(1)).table;table.cell(0,0).text='指标';table.cell(1,1).text='8%'
chart=CategoryChartData();chart.categories=['2026'];chart.add_series('节能',[8]);slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(6),Inches(1),Inches(3),Inches(2),chart)
slide.notes_slide.notes_text_frame.text='引用来源：学校节能记录。';buffer=io.BytesIO();deck.save(buffer);pptx=buffer.getvalue();(output/'native-template.pptx').write_bytes(pptx)
p=parse('pptx',pptx);schemas.validate('template-native-pack.schema.json',p['document']['template'])
check('pptx_shapes_tables_charts_images_notes',len(p['document']['slides'][0]['objects'])==5 and len(p['document']['charts'])==1 and p['document']['slides'][0]['notes'][0]['text'])
pack=parse_template(pptx,source={'path':'native-template.pptx'});check('template_master_layout_theme_slots',len(pack['masters'])==1 and len(pack['layouts'])==11 and len(pack['themes'])>=1 and any(l['slots'] for l in pack['layouts']),{'masters':len(pack['masters']),'layouts':len(pack['layouts']),'themes':len(pack['themes'])})
candidates=query_layouts(pack,minimum_slots=2);check('template_structural_query_with_unmeasured_capacity',any(c['eligible'] for c in candidates) and all(c['capacity_status']=='unmeasured' for c in candidates))
selected=next(l for l in pack['layouts'] if l['type']=='obj');instantiated,report=instantiate_template(pptx,selected['part'],{'0':'实例化标题','1':'保留母版布局与主题'})
(output/'instantiated-template.pptx').write_bytes(instantiated)
loaded=Presentation(io.BytesIO(instantiated));check('template_actual_editable_instantiation',len(loaded.slides)==2 and loaded.slides[-1].shapes.title.text=='实例化标题' and report['master_layout_theme_preserved'])
try:query_layouts(pack,slide_size={'width_emu':1,'height_emu':1})
except TaskError as error:check('template_size_mismatch_explicit',error.code=='TEMPLATE_SIZE_MISMATCH')
else:raise AssertionError('template size mismatch accepted')

# Minimal legitimate text PDF, independently parsed through installed Poppler.
stream=b'BT /F1 18 Tf 50 100 Td (Energy 2026 reduced 8 percent) Tj ET'
objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 200] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>',b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
raw=b'%PDF-1.4\n';offsets=[0]
for n,o in enumerate(objects,1):offsets.append(len(raw));raw+=str(n).encode()+b' 0 obj\n'+o+b'\nendobj\n'
xref=len(raw);raw+=b'xref\n0 6\n0000000000 65535 f \n'+b''.join(f'{o:010d} 00000 n \n'.encode() for o in offsets[1:]);raw+=f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
(output/'text-source.pdf').write_bytes(raw);pdf=parse('pdf',raw)
check('pdf_real_poppler_page_text_bbox',len(pdf['document']['pages'])==1 and '2026' in pdf['segments'][0]['raw_text'] and pdf['segments'][0]['data']['bbox'])
check('pdf_table_semantics_explicit',any(w['code']=='PDF_TABLE_STRUCTURE_UNRESOLVED' for w in pdf['warnings']))
reject('pdf_disguised','pdf',b'wrong signature','INPUT_FORMAT_DISGUISED')

j=parse('json',b'{"a":[1,{"b":"8%"}]}');check('json_hierarchy_preserved',j['document']['tree']['a'][1]['b']=='8%')
reject('json_duplicate_key','json',b'{"a":1,"a":2}','JSON_DUPLICATE_KEY');reject('json_nonfinite','json',b'{"a":NaN}','JSON_NONFINITE')
reject('xml_entity_blocked','xml',b'<!DOCTYPE root [<!ENTITY x SYSTEM "file:///etc/passwd">]><root>&x;</root>','XML_ENTITY_BLOCKED')
reject('ooxml_disguised','docx',xlsx,'INPUT_FORMAT_DISGUISED');reject('ooxml_corrupt','docx',b'PK broken','INPUT_ARCHIVE_CORRUPT')
reject('office_encrypted_or_legacy','docx',bytes.fromhex('D0CF11E0A1B11AE1')+b'header','INPUT_ENCRYPTED_OR_LEGACY')
reject('unsupported_format','yaml',b'a: 1','INPUT_UNSUPPORTED')
unsafe=office('word/document.xml','<document/>',{'../escape':'bad'});reject('archive_path_escape','docx',unsafe,'INPUT_UNSAFE_PART')
image_dir=output/'assets';image_dir.mkdir(exist_ok=True);(image_dir/'one.png').write_bytes(png);(image_dir/'same.png').write_bytes(png);(image_dir/'font.ttf').write_bytes(b'not a licensed font')
inv=inventory(image_dir);check('asset_inventory_hash_duplicates_unknown_license',len(inv['entries'])==3 and len(inv['duplicate_groups'])==1 and all(e['license']=='unknown' for e in inv['entries']))

report={'passed':True,'checks':cases,'requirement_ids':['IN-003','IN-004','IN-006','IN-007','IN-008','IN-010','IN-011','IN-012','IN-015','IN-016','TPL-001','TPL-002','TPL-003','TPL-004','TPL-005','TPL-009','TPL-014','AST-001','AST-004','AST-011','SEC-004','SEC-005','SEC-006'],
        'source_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [ROOT/'src/yoloongppt/input_formats.py',ROOT/'src/yoloongppt/templates.py',ROOT/'src/yoloongppt/assets.py']},
        'boundary':'Targeted actual Docker parser/template checks; OCR/PDF table semantics/capacity/visual QA/system AC and multi-project template fixtures remain incomplete.'}
(output/'input-formats-runtime.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':True,'checks':len(cases),'report':str(output/'input-formats-runtime.json'),'fixtures':str(output)},ensure_ascii=False))
