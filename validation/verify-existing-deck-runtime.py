"""Bounded native fixture: import, multiple edits, fidelity, negative cases."""
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path

sys.path.insert(0,'/workspace/src')
from PIL import Image
from pptx import Presentation
from yoloongppt.existing_deck import inspect_deck, apply_patch, select_objects
from yoloongppt.native_objects import add_object, capabilities
from yoloongppt.native_objects import local_asset
from yoloongppt.errors import TaskError
from yoloongppt.schemas import SchemaRegistry

folder=Path('/runtime/native-object-validation');folder.mkdir(exist_ok=True)
checks=[]
def check(name,condition):
    assert condition,name;checks.append(name)
def rejected(name,code,request):
    output=folder/(name+'.pptx')
    try: apply_patch(source,output,request)
    except TaskError as e: check(name,e.code==code and not output.exists())
    else: raise AssertionError(name)

for path in folder.glob('*.pptx'): path.unlink()
image=folder/'image.png';replacement=folder/'replacement.png'
Image.new('RGB',(32,32),'blue').save(image);Image.new('RGB',(32,32),'red').save(replacement)
prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6]);s2=prs.slides.add_slide(prs.slide_layouts[6])
text=add_object(s,{'kind':'text','name':'标题','bounds':[1,1,4,1],'text':'修改前','hyperlink':'https://example.com/','style':{'font_size':24,'color':'112233'}})
shape=add_object(s,{'kind':'shape','bounds':[1,3,1,1],'text':'保留形状','style':{'fill':'ABCDEF'}})
table=add_object(s,{'kind':'table','bounds':[3,3,2,1],'cells':[['A','B'],['C','D']]})
chart=add_object(s,{'kind':'chart','bounds':[5,1,4,3],'data':{'categories':['Q1','Q2'],'series':[{'name':'能耗','values':[8,9]}]}})
picture=add_object(s,{'kind':'image','path':str(image),'bounds':[1,5,1,1],'alt_text':'初始图片'})
connector=add_object(s,{'kind':'connector','bounds':[2,5,1,1],'connector_type':'elbow'})
group=add_object(s,{'kind':'group','bounds':[6,5,1,1],'children':[{'kind':'shape','bounds':[6,5,1,1],'text':'组合子对象'}]})
other=add_object(s2,{'kind':'text','bounds':[1,1,4,1],'text':'其它页原样保留'})
s.notes_slide.notes_text_frame.text='原始备注'
source=folder/'fixture.pptx';prs.save(source)
with zipfile.ZipFile(source,'a') as z: z.writestr('customXml/vendor-unknown.xml',b'<vendor xmlns="urn:vendor"><opaque>KEEP</opaque></vendor>')
model=inspect_deck(source)
check('two_slides_native_inventory',len(model['slides'])==2 and len(model['slides'][0]['objects'])==8)
check('unknown_part_inventoried','customXml/vendor-unknown.xml' in model['parts'])
ids={o['shape_id']:o['object_id'] for o in model['slides'][0]['objects']}
req={'expected_sha256':model['source_sha256'],'patches':[
 {'object_id':ids[text.shape_id],'operation':'text','value':'修改后的标题'},
 {'object_id':ids[text.shape_id],'operation':'format','value':{'font':'Noto Sans CJK SC','font_size':30,'bold':True,'color':'AABBCC','alignment':'center'}},
 {'object_id':ids[text.shape_id],'operation':'geometry','value':[1.1,1.2,4.2,1.3]},
 {'object_id':ids[table.shape_id],'operation':'table_cell','value':{'row':1,'column':1,'text':'已修改'}},
 {'object_id':ids[chart.shape_id],'operation':'chart_data','value':{'categories':['Q1','Q2'],'series':[{'name':'能耗','values':[8,10]}]}},
 {'object_id':ids[picture.shape_id],'operation':'image_replace','value':{'path':str(replacement)}},
 {'object_id':ids[picture.shape_id],'operation':'alt_text','value':'替换后的图片'},
 {'object_id':ids[text.shape_id],'operation':'notes','value':'修改后的演讲备注'},
 {'object_id':ids[group.shapes[0].shape_id],'operation':'text','value':'组合子对象局部修改'}]}
output=folder/'edited.pptx';report=apply_patch(source,output,req)
with zipfile.ZipFile(source) as a,zipfile.ZipFile(output) as b:
 check('all_unchanged_members_byte_identical',all(a.read(n)==b.read(n) for n in report['preserved_parts']))
 check('unknown_xml_byte_identical',a.read('customXml/vendor-unknown.xml')==b.read('customXml/vendor-unknown.xml'))
 check('master_layout_theme_byte_identical',all(a.read(n)==b.read(n) for n in a.namelist() if any(x in n for x in ('slideMasters/','slideLayouts/','theme/'))))
check('second_slide_byte_identical',model['slides'][1]['part'] in report['preserved_parts'])
updated=Presentation(output);check('text_actual_edit',updated.slides[0].shapes[0].text=='修改后的标题')
check('font_actual_edit',updated.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt==30)
check('geometry_actual_edit',updated.slides[0].shapes[0].left==round(1.1*914400))
check('table_actual_edit',updated.slides[0].shapes[2].table.cell(1,1).text=='已修改')
check('chart_cache_actual_edit',list(updated.slides[0].shapes[3].chart.series[0].values)==[8.,10.])
with zipfile.ZipFile(output) as z:
    workbook=next(n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx'))
    with zipfile.ZipFile(io.BytesIO(z.read(workbook))) as x:
        from lxml import etree
        sheet=etree.fromstring(x.read('xl/worksheets/sheet1.xml'))
        check('chart_workbook_matches_cache',sheet.xpath('//*[local-name()="c"][@r="B3"]/*[local-name()="v"]/text()')==['10'])
check('image_actual_edit',updated.slides[0].shapes[4].image.blob==replacement.read_bytes())
check('notes_actual_edit',updated.slides[0].notes_slide.notes_text_frame.text=='修改后的演讲备注')
check('group_child_local_edit',updated.slides[0].shapes[6].shapes[0].text=='组合子对象局部修改')
check('hyperlink_created_preserved',updated.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].hyperlink.address=='https://example.com/')
after=inspect_deck(output);check('object_ids_stable',[o['object_id'] for o in after['slides'][0]['objects']]==[o['object_id'] for o in model['slides'][0]['objects']])
check('exact_selector_name',select_objects(model,{'name':'标题'})['object_id']==ids[text.shape_id])
check('group_child_indexed',any(o['name']==group.shapes[0].name for o in model['slides'][0]['objects']))
check('masters_layouts_themes_imported',all(model['components'][k] for k in ['masters','layouts','themes']))
check('chart_data_imported',model['slides'][0]['objects'][3]['chart']['series'][0]['values']==['8','9'])
rejected('bad_hash','REVISION_SOURCE_CHANGED',{**req,'expected_sha256':'0'*64})
rejected('locked_object','REVISION_OBJECT_LOCKED',{**req,'locked_object_ids':[ids[text.shape_id]]})
rejected('unsupported_operation','REVISION_OPERATION_UNSUPPORTED',{'expected_sha256':model['source_sha256'],'patches':[{'object_id':ids[text.shape_id],'operation':'animation','value':{}}]})
rejected('ambiguous_selector','OBJECT_SELECTION_AMBIGUOUS',{'expected_sha256':model['source_sha256'],'patches':[{'selector':{'slide_index':0},'operation':'text','value':'不执行'}]})
rejected('bad_geometry','TASK_SCHEMA_INVALID',{'expected_sha256':model['source_sha256'],'patches':[{'object_id':ids[text.shape_id],'operation':'geometry','value':[0,0,-1,2]}]})
rejected('bad_table_cell','REVISION_TABLE_INVALID',{'expected_sha256':model['source_sha256'],'patches':[{'object_id':ids[table.shape_id],'operation':'table_cell','value':{'row':9,'column':0,'text':'越界'}}]})
rejected('secret_image_path','NATIVE_ASSET_PATH_FORBIDDEN',{'expected_sha256':model['source_sha256'],'patches':[{'object_id':ids[picture.shape_id],'operation':'image_replace','value':{'path':'/runtime/secrets/deepseek_api_key'}}]})
try: local_asset('/etc/passwd')
except TaskError as e: check('outside_media_path_denied',e.code=='NATIVE_ASSET_PATH_FORBIDDEN')
else: raise AssertionError('outside path permitted')
cap=capabilities();check('capabilities_all30_partial_honest',len(cap['requirements'])==30 and all(x['status']!='supported' for x in cap['requirements']))
check('source_never_changed',hashlib.sha256(source.read_bytes()).hexdigest()==model['source_sha256'])
schemas=SchemaRegistry();schemas.validate('existing-deck-patch.schema.json',req);check('consumed_input_schema',True)
manifest={'ok':True,'checks':checks,'source_sha256':model['source_sha256'],'output_sha256':report['output_sha256'],
 'changed_parts':report['changed_parts'],'capabilities':cap,'qa':'not_executed','acceptance':'partial_component_only',
 'requirement_ids':['REV-001','REV-002','REV-004','REV-005','REV-010','PPT-008','PPT-009','PPT-011','PPT-012','PPT-013','PPT-014','PPT-016','PPT-017','PPT-023','PPT-029','PPT-030'],
 'source_versions':{str(p.relative_to('/workspace')):hashlib.sha256(p.read_bytes()).hexdigest() for p in [Path('/workspace/src/yoloongppt/native_objects.py'),Path('/workspace/src/yoloongppt/existing_deck.py')]}}
(folder/'native-object-runtime.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(folder/'preservation-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'ok':True,'checks':len(checks),'changed_parts':report['changed_parts']},ensure_ascii=False))
