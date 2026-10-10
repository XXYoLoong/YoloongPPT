"""Bounded Docker evidence for DEC-022..029 / SYS-008/009, no model calls.

Fixtures validate actual module functions and independent nodes. Literal arrow
grounding is not semantic proof or complete acceptance of diagram/asset modes.
"""
import copy
import hashlib
import io
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from yoloongppt.artifacts import entity
from yoloongppt.asset_resolver import resolve, download
from yoloongppt.errors import TaskError
from yoloongppt.schemas import SchemaRegistry
from yoloongppt.template_registry import register, get, candidates
from yoloongppt.visual import NODES, prepare, finalize, run_node

output = Path('/runtime/visual-validation'); output.mkdir(exist_ok=True, parents=True)
schemas = SchemaRegistry(); checks = []; samples = []


def check(name, condition, observed=None):
    assert condition, (name, observed)
    checks.append({'id': name, 'passed': True, 'observed': observed})


def reject(name, function, expected):
    try: function()
    except TaskError as error:
        check(name, error.code == expected, error.code)
    else: raise AssertionError(name + ' unexpectedly accepted')


task = json.loads(Path('/workspace/validation/core-task.json').read_text('utf-8'))
task['template_ref'] = None; task['style'] = None; task['sources'] = []
context = {'trace_id': entity('trace'), 'values': {'aspect_ratio': '16:9'}, 'constraints': task['constraints']}
evidence_id = entity('evidence')
text = '需求确认 → 设计 → 制作。2026启动 → 2027扩展 → 2028完成。指标A为8，指标B为9。'
sources = {'evidence': [{'evidence_id': evidence_id, 'raw_text': text}],
           'approved_texts': {evidence_id: text}, 'fact_boundary': {'status': 'ready'}}
refs = [evidence_id]
asset_id = entity('asset')
image_path = output / 'declared-image.png'
buffer = io.BytesIO(); Image.new('RGB', (320, 160), (26, 80, 67)).save(buffer, format='PNG'); image_path.write_bytes(buffer.getvalue())
asset = {'source_id': entity('source'), 'kind': 'image', 'locator': str(image_path),
         'metadata': {'source_role': 'asset', 'asset_id': asset_id, 'license': 'unknown', 'usage_status': 'authorized'}}
task['sources'] = [asset]
base_slide = {'title': '已批准视觉样例', 'goal': '保留来源，选择原生表达', 'kind': 'text', 'body': ['需求确认', '设计', '制作'],
              'notes': '结构关系由原文明示箭头提供，语义复审另行执行。', 'evidence_refs': refs, 'rationale': '明示结构且保留数据', 'table': None, 'chart': None}


def proposal(slide):
    return {'title': '视觉样例', 'main_takeaway': '保留证据', 'storyline': {'pattern': 'linear', 'rationale': '原文顺序'}, 'slides': [slide]}


def diagram(kind, labels, quote):
    return {'type': kind, 'direction': 'horizontal', 'nodes': [{'node_id': f'n{i}', 'label': label, 'evidence_refs': refs} for i, label in enumerate(labels)],
            'edges': [{'from': f'n{i}', 'to': f'n{i+1}', 'relation': 'temporal' if kind == 'timeline' else 'sequence', 'evidence_refs': refs, 'evidence_quote': quote} for i in range(len(labels)-1)]}


images = []
for role in ('information', 'decoration', 'logo'):
    slide = copy.deepcopy(base_slide); slide.update(kind='image', image={'asset_id': asset_id, 'role': role, 'alt_text': '调用者提供的绿色图片', 'fit': 'contain', 'required': True})
    images.append(slide)
tables = []
for columns, rows in [(['指标','值'], [['A','8'],['B','9']]), (['指标'],[['A'],['B']]), (['值'],[['8']])]:
    slide = copy.deepcopy(base_slide);slide.update(kind='table',table={'columns':columns,'rows':rows,'units':'原文单位未指定'});tables.append(slide)
charts = []
for kind in ('bar','column','line'):
    slide = copy.deepcopy(base_slide);slide.update(kind='chart',chart={'type':kind,'categories':['A','B'],'series':[{'name':'指标','values':[8,9]}],'units':'原文单位未指定'});charts.append(slide)
diagrams = []
for kind, labels, quote in [('process',['需求确认','设计','制作'],'需求确认 → 设计 → 制作'),('timeline',['2026启动','2027扩展','2028完成'],'2026启动 → 2027扩展 → 2028完成'),('process',['需求确认','设计'],'需求确认 → 设计')]:
    slide=copy.deepcopy(base_slide);slide.update(kind='diagram',diagram=diagram(kind,labels,quote));diagrams.append(slide)

normal_samples = {'DEC-022': [base_slide,tables[0],images[0]], 'DEC-023': images, 'DEC-024': images,
                  'DEC-025': tables, 'DEC-026': charts, 'DEC-027': diagrams}
for node, slides in normal_samples.items():
    for index, slide in enumerate(slides):
        document = {'task': task, 'sources': sources, 'context': context, 'proposal': proposal(slide)}
        result = run_node(node, document, schemas)
        check(f'{node}_independent_normal_{index+1}', result['node_id'] == node and len(result['decision_trace']['selected']) == 1)
        samples.append({'name':f'{node}_normal_{index+1}', 'input':document, 'result':result})

template = Presentation(); template.slide_width=Inches(13.333333);template.slide_height=Inches(7.5)
slide=template.slides.add_slide(template.slide_layouts[1]);slide.shapes.title.text='登记模板';slide.placeholders[1].text='需求确认'
template_path=output/'wide-template.pptx';template.save(template_path)
record=register(template_path, schemas=schemas, source=str(template_path))
check('template_actual_registered_file_hash', Path(record['path']).is_file() and hashlib.sha256(Path(record['path']).read_bytes()).hexdigest()==record['sha256'])
check('template_persistent_query', get(record['template_id'],schemas=schemas)['template_id']==record['template_id'] and any(candidate['eligible'] for candidate in candidates(record)))
check('template_structural_preview_no_fake_thumbnail', record['preview']['status']=='structural' and bool(record['preview']['layouts']))
for index, family in enumerate(('generic','reference','brand')):
    chosen=copy.deepcopy(task)
    if family!='generic':
        chosen['template_ref']=str(template_path)
        chosen['sources'].append({'source_id':entity('source'),'kind':'pptx','locator':str(template_path),'metadata':{'source_role':'template','brand_locked':family=='brand'}})
    result=run_node('DEC-028',{'task':chosen,'sources':sources,'context':context},schemas)
    check(f'DEC-028_independent_normal_{index+1}',result['output']['family']==family)
    samples.append({'name':f'DEC-028_normal_{index+1}','result':result})
for index, color in enumerate(('147C80','000000','FFFFFF')):
    chosen=copy.deepcopy(task);chosen['style']={'accent':color}
    result=run_node('DEC-029',{'task':chosen,'sources':sources,'context':context},schemas)
    check(f'DEC-029_independent_normal_{index+1}',result['output']['accent']==color)
    samples.append({'name':f'DEC-029_normal_{index+1}','result':result})

packet=prepare(task,sources,context,schemas)
current=proposal(images[0]); before=copy.deepcopy(current)
final,result=finalize(current,packet,sources,schemas)
check('resolved_image_real_file_hash_and_pixels',Path(final['slides'][0]['image']['path']).is_file() and final['slides'][0]['image']['width_px']==320 and final['slides'][0]['image']['sha256']==hashlib.sha256(image_path.read_bytes()).hexdigest())
check('visual_no_rewrite_original_model_proposal',current==before)
check('visual_all_eight_nodes_executed',set(trace['node_id'] for trace in result['decision_traces'])==set(NODES))
check('unknown_license_preserved_with_explicit_authorization',result['resolved_assets']['entries'][0]['license']=='unknown' and any(issue['code']=='ASSET_LICENSE_UNKNOWN' for issue in result['resolved_assets']['issues']))

def bad_slide(node, mutate, code):
    slide=copy.deepcopy({'DEC-022':base_slide,'DEC-023':images[0],'DEC-024':images[0],'DEC-025':tables[0],'DEC-026':charts[0],'DEC-027':diagrams[0]}[node]);mutate(slide)
    reject(node+'_boundary',lambda:run_node(node,{'task':task,'sources':sources,'context':context,'proposal':proposal(slide)},schemas),code)

bad_slide('DEC-022',lambda slide:slide.update(kind='infographic'),'VISUAL_KIND_UNSUPPORTED')
bad_slide('DEC-023',lambda slide:slide['image'].update(role='unknown'),'IMAGE_REQUIREMENT_INVALID')
bad_slide('DEC-024',lambda slide:slide['image'].update(asset_id='unregistered'),'IMAGE_ASSET_UNKNOWN')
bad_slide('DEC-025',lambda slide:slide['table']['rows'].append(['wrong width']),'VISUAL_TABLE_SHAPE')
bad_slide('DEC-026',lambda slide:slide['chart'].update(type='scatter'),'VISUAL_CHART_TYPE_UNSUPPORTED')
bad_slide('DEC-027',lambda slide:slide['diagram']['edges'][0].update(relation='causal'),'DIAGRAM_CAUSAL_PROOF_REQUIRED')
narrow=Presentation();narrow_path=output/'narrow-template.pptx';narrow.save(narrow_path)
chosen=copy.deepcopy(task);chosen['template_ref']=str(narrow_path)
reject('DEC-028_boundary',lambda:run_node('DEC-028',{'task':chosen,'sources':sources,'context':context},schemas),'TEMPLATE_SIZE_MISMATCH')
chosen=copy.deepcopy(task);chosen['style']={'accent':'guess'}
reject('DEC-029_boundary',lambda:run_node('DEC-029',{'task':chosen,'sources':sources,'context':context},schemas),'VISUAL_STYLE_UNSUPPORTED')
bad=copy.deepcopy(diagrams[0]);bad['diagram']['edges'][0]['evidence_quote']='需求确认和设计'
reject('diagram_unquoted_relation_rejected',lambda:finalize(proposal(bad),packet,sources,schemas),'DIAGRAM_EDGE_UNGROUNDED')
bad=copy.deepcopy(diagrams[0]);bad['diagram']['nodes'][0]['label']='臆造因果'
reject('diagram_fabricated_node_rejected',lambda:finalize(proposal(bad),packet,sources,schemas),'DIAGRAM_NODE_LABEL_UNGROUNDED')
bad=copy.deepcopy(sources);bad['approved_texts']={}
reject('missing_fact_projection_never_raw_fallback',lambda:finalize(proposal(diagrams[0]),packet,bad,schemas),'FACT_PROJECTION_INCOMPLETE')
request={'asset_id':asset_id,'source_type':'local','path':str(image_path),'license':'unknown','usage_status':'authorized'}
reject('asset_missing_file_explicit',lambda:resolve({'requests':[{**request,'path':str(output/'missing.png')}]},schemas=schemas),'ASSET_FILE_MISSING')
reject('asset_permission_denied',lambda:resolve({'requests':[{**request,'usage_status':'denied'}]},schemas=schemas),'ASSET_USAGE_UNAUTHORIZED')
reject('asset_low_resolution_no_fake_upscale',lambda:resolve({'requests':[request],'minimum_pixels':[400,200]},schemas=schemas),'ASSET_RESOLUTION_INSUFFICIENT')
reject('asset_hash_mismatch',lambda:resolve({'requests':[{**request,'expected_sha256':'0'*64}]},schemas=schemas),'ASSET_HASH_MISMATCH')
reject('remote_asset_opt_in_required',lambda:download('https://example.invalid/a.png'),'ASSET_REMOTE_DISABLED')
reject('asset_private_address_blocked',lambda:download('https://127.0.0.1/a.png',enabled=True),'ASSET_SSRF_BLOCKED')
duplicate=resolve({'requests':[request,{**request,'asset_id':entity('asset')}]},schemas=schemas)
check('duplicate_without_authorization_explicit',bool(duplicate['duplicate_groups']) and any(issue['code']=='ASSET_DUPLICATE_UNCONFIRMED' for issue in duplicate['issues']))
deliberate=resolve({'requests':[{**request,'deliberate_reuse':True},{**request,'asset_id':entity('asset'),'deliberate_reuse':True}]},schemas=schemas)
check('deliberate_duplicate_use_preserved',deliberate['duplicate_groups'][0]['deliberate_reuse'] is True)
schema_names=[path.name for path in Path('/workspace/contracts').glob('*.schema.json') if path.name.startswith(('visual-','asset-resolver-','template-registry-'))]
report={'checks':checks,'passed':len(checks),'source_hashes':{name:hashlib.sha256(Path('/workspace/src/yoloongppt',name).read_bytes()).hexdigest() for name in ('visual.py','asset_resolver.py','template_registry.py')},
        'schema_hashes':{name:schemas.versions[name] for name in sorted(schema_names)},'node_samples':'node-samples.json',
        'actual_scope':'Independent 3 normal and boundary cases per DEC-022..029; local image bytes/cache/hash/authorization and persistent native template registration/query; no live model, remote public download, image generation or system AC claim.'}
(output/'node-samples.json').write_text(json.dumps(samples,ensure_ascii=False,indent=2),encoding='utf-8')
(output/'visual-runtime.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(output/'visual-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':len(checks),'report':str(output/'visual-runtime.json'),'node_samples':str(output/'node-samples.json')},ensure_ascii=False))
