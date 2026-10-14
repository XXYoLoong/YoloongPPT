"""PPT-015: actual CLI/API/MCP reuse, editable vector output and render."""
import hashlib
import json
import subprocess
import sys
import urllib.request
from pathlib import Path
from uuid import uuid4

from pptx import Presentation
from yoloongppt.rendering import render

out=Path('/runtime/native-svg-entrypoints')/uuid4().hex
out.mkdir(parents=True)
fixture=Path('/runtime/upstream-native-validation-v6/slide-one.svg')
spec={'kind':'svg','name':'object_native_svg_'+uuid4().hex,'path':str(fixture),
      'bounds':[1,1,10,5],'sha256':hashlib.sha256(fixture.read_bytes()).hexdigest(),
      'alt_text':'Upstream native SVG verification, artificial fixture.'}
task={'slides':[{'objects':[spec],'notes':'Software verification; artificial fixture, no business factual assertion.'}]}
(out/'task.json').write_text(json.dumps(task,ensure_ascii=False,indent=2))
checks=[]
def check(name,value,observed=None):
    checks.append({'id':name,'passed':bool(value),'observed':observed})
    if not value:raise AssertionError(name)
def run_cli(args,input=None):
    p=subprocess.run([sys.executable,'-m','yoloongppt',*args],input=input,text=True,capture_output=True,timeout=180)
    check('cli-process-'+args[0],p.returncode==0,p.stderr)
    return p.stdout

results={'cli':json.loads(run_cli(['compose-native',str(out/'task.json')]))}
request=urllib.request.Request('http://127.0.0.1:8000/compose-native',
    data=json.dumps(task).encode(),headers={'Content-Type':'application/json'})
with urllib.request.urlopen(request,timeout=180) as response:results['api']=json.load(response)
messages=[{'jsonrpc':'2.0','id':1,'method':'initialize','params':{'protocolVersion':'2025-06-18','capabilities':{},'clientInfo':{'name':'verification','version':'1'}}},
          {'jsonrpc':'2.0','method':'notifications/initialized'},
          {'jsonrpc':'2.0','id':2,'method':'tools/call','params':{'name':'compose_native','arguments':task}}]
responses=[json.loads(line) for line in run_cli(['mcp-serve'],'\n'.join(json.dumps(x) for x in messages)+'\n').splitlines()]
results['mcp']=next(x for x in responses if x.get('id')==2)['result']['structuredContent']
for entry,result in results.items():
    check(entry+'-success',result.get('ok'),result)
    folder=Path(result['pptx']).parent
    prs=Presentation(result['pptx']);group=prs.slides[0].shapes[0]
    check(entry+'-native-editable-group',hasattr(group,'shapes') and len(group.shapes)>1)
    check(entry+'-stable-object-name',group.name==spec['name'])
    check(entry+'-no-bitmap',not any(s.shape_type==13 for s in group.shapes))
    calls=json.loads((folder/'native-upstream-calls.json').read_text())['calls']
    check(entry+'-original-source-hash',len(calls)==1 and calls[0]['output']['source_sha256']==spec['sha256'])
    native_ids={int(node.get('id')) for node in group._element.iter() if node.tag.rsplit('}',1)[-1]=='cNvPr' and int(node.get('id'))!=group.shape_id}
    mapped_ids={item['shape_id'] for item in calls[0]['output']['shape_ids']}
    check(entry+'-child-object-map',native_ids==mapped_ids)
    check(entry+'-original-upstream',calls[0]['output']['upstream']['commit']=='2d72da616cf9fa40d4dcaf59fd4c980ecf534b7d')
    report=render(result['pptx'],{'slides':[{'slide_id':'slide_'+str(uuid4())}]},folder)
    check(entry+'-actual-render',any(a['type']=='pdf' for a in report['artifacts']) and any(a['type']=='png' for a in report['artifacts']))
    (out/(entry+'-render.json')).write_text(json.dumps(report,indent=2))
(out/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
(out/'checks.json').write_text(json.dumps({'checks':checks,'passed':True,'scope':'Native SVG entrypoint integration; original PPT-015 full CRUD/style/backend and system acceptance remain required.'},ensure_ascii=False,indent=2))
print(json.dumps({'checks':len(checks),'passed':True,'evidence':str(out)}))
