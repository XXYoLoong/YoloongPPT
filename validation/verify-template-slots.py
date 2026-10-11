"""Docker-only bounded native template/brand evidence; no model calls or AC completion claim."""
import copy
import hashlib
import io
import json
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from yoloongppt.errors import TaskError
from yoloongppt.errors import trace_id
from yoloongppt.artifacts import entity
from yoloongppt.atomic import AtomicRegistry
from yoloongppt.planning import compile_deck, DEFAULT_STYLE
from yoloongppt.writer import execute
from yoloongppt.rendering import render
from yoloongppt.schemas import SchemaRegistry
from yoloongppt.template_layout_runtime import choose, create_slide, bind_placeholder, preserve_package, preserved_part_names, serialize_package_xml
from yoloongppt.template_registry import register, get, select

ROOT=Path('/runtime/template-slots-'+str(uuid.uuid4())); ROOT.mkdir(parents=True)
schemas=SchemaRegistry(); checks=[]; records=[]; selections=[]; outputs=[]
style={'font':'Noto Sans CJK SC','foreground':'000000','background':'FFFFFF','accent':'176548'}
policy={'version':'template-policy-1','typography':{'font':'Noto Sans CJK SC','sizes':{'title':28,'body':20}}}
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
rel_ns='http://schemas.openxmlformats.org/package/2006/relationships'
def check(name, condition, observed=None):
    assert condition,(name,observed)
    checks.append({'id':name,'passed':True,'observed':observed})
def reject(name, function, code):
    try:function()
    except TaskError as e:check(name,e.code==code,{'code':e.code,'details':e.details})
    else:raise AssertionError(name+' unexpectedly accepted')
def patch_package(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:parts={n:z.read(n) for n in z.namelist()}
    parts['customXml/brand-extra.xml']=b'<brand-extension version="1">retain unknown semantics</brand-extension>'
    rels=ET.fromstring(parts['_rels/.rels']); ET.SubElement(rels,'{'+rel_ns+'}Relationship',{'Id':'rIdBrandExtension','Type':'urn:yoloong:brand-extension','Target':'customXml/brand-extra.xml'})
    parts['_rels/.rels']=serialize_package_xml(rels)
    ct=ET.fromstring(parts['[Content_Types].xml']);ET.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',{'PartName':'/customXml/brand-extra.xml','ContentType':'application/xml'});parts['[Content_Types].xml']=serialize_package_xml(ct)
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        for n,v in parts.items():z.writestr(n,v)
    return out.getvalue()

p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
exemplar=p.slides.add_slide(p.slide_layouts[6]);master=p.slide_masters[0]
logo=exemplar.shapes.add_textbox(Inches(11.6),Inches(6.9),Inches(1.1),Inches(.35));logo.text='Yoloong';logo.name='LockedLogo'
footer=exemplar.shapes.add_textbox(Inches(.6),Inches(6.9),Inches(4),Inches(.35));footer.text='企业保密 · 原模板页脚';footer.name='LockedFooter'
for shape in (logo,footer):
    for para in shape.text_frame.paragraphs:para.font.name='Noto Sans CJK SC';para.font.size=Pt(16)
    node=copy.deepcopy(shape._element);node.find('.//p:cNvPr',NS).set('id',str(100+len(master.shapes)))
    master._element.find('p:cSld/p:spTree',NS).append(node)
# Native table/chart placeholders, plus distinct naming to test role crosswalk.
for layout_no, ph_role in ((3,'tbl'),(4,'chart')):
    layout=p.slide_layouts[layout_no]
    for ph in layout.placeholders:
        idx=ph.placeholder_format.idx
        if idx==2:ph._element.find('.//p:ph',NS).set('type',ph_role);ph.name='业务数据区域'
        if idx==1:ph._element.find('.//p:ph',NS).set('type','body');ph.name='业务说明区域'
raw=io.BytesIO();p.save(raw);raw=patch_package(raw.getvalue());template=ROOT/'enterprise.pptx';template.write_bytes(raw)
img=ROOT/'source.png';Image.new('RGB',(320,160),'#176548').save(img)
refs=[entity('evidence')]
image={'path':str(img),'sha256':hashlib.sha256(img.read_bytes()).hexdigest(),'width_px':320,'height_px':160,'fit':'contain','crop':[0,0,0,0],'source_refs':refs,'alt_text':'原始图片'}
base={'title':'企业模板实际填充','body':['原生占位符中的可编辑正文。'],'evidence_refs':refs,'kind':'text','goal':'保留真实模板槽位','notes':'原文与品牌结构保留。','rationale':'实际本地原生槽位验证','table':None,'chart':None}
brand={'locked_objects':[{'part':'ppt/slideMasters/slideMaster1.xml','name':'LockedLogo'},{'part':'ppt/slideMasters/slideMaster1.xml','name':'LockedFooter'}],
       'safe_areas':[[0,6.8,13.333333,.7]],'locked_tokens':{'font':'Noto Sans CJK SC','foreground':'000000','background':'FFFFFF','accent':'176548'}}

def registered(extra):
    pol={**copy.deepcopy(policy),**extra}
    result=register(template,root=str(ROOT/'registry'),allowed_roots=['/runtime'],schemas=schemas,policy=pol)
    records.append(result)
    check('registry-'+str(len(records)),get(result['template_id'],root=str(ROOT/'registry'),schemas=schemas)==result)
    return result

def make(result, content, name):
    selected=choose(result,content,style);schemas.validate('template-layout-runtime.schema.json',selected);selections.append(selected)
    check(name+'-actual-font-file',all(Path(e['template_binding']['typography']['font_file']).is_file() for e in selected['elements']))
    pres=Presentation(io.BytesIO(raw))
    for item in list(pres.slides._sldIdLst):pres.part.drop_rel(item.rId);pres.slides._sldIdLst.remove(item)
    slide=create_slide(pres,{'layout':{'template_binding':{'native_layout_part':selected['native_layout_part'],'preserve_master_style':True,'locks':selected['locks']}}})
    ids=[]
    for element in selected['elements']:
        ph=bind_placeholder(slide,element);idx=ph.placeholder_format.idx;ids.append(idx)
        if element['type']=='text':
            ph.text_frame.clear()
            for i,text in enumerate(element['text']):
                para=ph.text_frame.paragraphs[0] if i==0 else ph.text_frame.add_paragraph();para.text=text;para.font.name=selected['page_style']['font'];para.font.size=Pt(element['font_size'])
            shape=ph
        elif element['type']=='image':
            shape=ph.insert_picture(element['data']['path']);shape.crop_left=shape.crop_right=shape.crop_top=shape.crop_bottom=0
        elif element['type']=='table':
            data=element['data'];rows=[data['columns'],*data['rows']];shape=ph.insert_table(len(rows),len(data['columns']))
            for r,row in enumerate(rows):
                for c,value in enumerate(row):shape.table.cell(r,c).text=value
        else:
            data=element['data'];chart=CategoryChartData();chart.categories=data['categories']
            for series in data['series']:chart.add_series(series['name'],series['values'])
            shape=ph.insert_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,chart)
        shape.left,shape.top,shape.width,shape.height=[Inches(v) for v in element['bounds']]
        shape.name=element['object_id']
        check(name+'-native-'+element['role'],shape.is_placeholder and shape.placeholder_format.idx==idx)
    serialized=io.BytesIO();pres.save(serialized);preserved=preserve_package(raw,serialized.getvalue());path=ROOT/(name+'.pptx');path.write_bytes(preserved)
    again=Presentation(path);actual=again.slides[0]
    check(name+'-layout',str(actual.slide_layout.part.partname).lstrip('/')==selected['native_layout_part'])
    check(name+'-no-overlay',len(actual.shapes)==len(slide.shapes),{'shapes':len(actual.shapes),'placeholder_indices':ids})
    with zipfile.ZipFile(io.BytesIO(raw)) as before,zipfile.ZipFile(path) as after:
        parts=preserved_part_names(before.namelist());changed=[n for n in parts if n not in after.namelist() or before.read(n)!=after.read(n)]
        check(name+'-native-parts-byte-preserved',not changed,{'preserved_parts':len(parts),'changed':changed})
    output={'case':name,'pptx':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'native_layout':selected['native_layout_part'],'bindings':ids}
    outputs.append(output);return selected,path

record=registered({'layout_part':'ppt/slideLayouts/slideLayout2.xml','brand':brand})
normal1,path1=make(record,base,'normal-brand-text')
check('brand-explicit-objects',len(normal1['locks']['objects'])==2)
normal2,path2=make(registered({'layout_part':'ppt/slideLayouts/slideLayout9.xml','field_map':{'body':2,'asset':1}}),{**base,'kind':'image','image':image},'normal-native-image')
im=next(x for x in normal2['elements'] if x['type']=='image');check('contain-preserves-aspect',abs(im['bounds'][2]/im['bounds'][3]-2)<1e-8 and im['data']['crop']==[0,0,0,0])
table={'columns':['指标','值'],'rows':[['实际','2']]}
normal3,path3=make(registered({'layout_part':'ppt/slideLayouts/slideLayout4.xml','field_map':{'body':'业务说明区域','data':'业务数据区域'}}),{**base,'kind':'table','table':table},'normal-native-table')
normal4,path4=make(registered({'layout_part':'ppt/slideLayouts/slideLayout5.xml','field_map':{'body':1,'data':2}}),{**base,'kind':'chart','chart':{'type':'column','categories':['A','B'],'series':[{'name':'实际','values':[2,3]}],'units':'个'}},'normal-native-chart')
check('table-chart-native-after-reopen',Presentation(path3).slides[0].placeholders[2].has_table and Presentation(path4).slides[0].placeholders[2].has_chart)

reject('boundary-brand-token',lambda:choose(record,base,{**style,'accent':'FF0000'}),'TEMPLATE_BRAND_TOKEN_CONFLICT')
reject('boundary-brand-safe-area',lambda:choose(registered({'layout_part':'ppt/slideLayouts/slideLayout2.xml','brand':{'safe_areas':[[0,0,13.333333,7.5]]}}),base,style),'TEMPLATE_NO_COMPATIBLE_LAYOUT')
reject('boundary-brand-slot-locked',lambda:choose(registered({'layout_part':'ppt/slideLayouts/slideLayout2.xml','brand':{'locked_objects':[{'part':'ppt/slideLayouts/slideLayout2.xml','name':'Title 1'}]}}),base,style),'TEMPLATE_NO_COMPATIBLE_LAYOUT')
reject('boundary-unmatched-field',lambda:choose(registered({'layout_part':'ppt/slideLayouts/slideLayout2.xml','field_map':{'body':'不存在的字段'}}),base,style),'TEMPLATE_NO_COMPATIBLE_LAYOUT')
reject('boundary-capacity',lambda:choose(record,{**base,'body':['正文'*3000]},style),'TEMPLATE_NO_COMPATIBLE_LAYOUT')
reject('boundary-no-fallback',lambda:choose(registered({'layout_part':'ppt/slideLayouts/slideLayout6.xml'}),base,style),'TEMPLATE_NO_COMPATIBLE_LAYOUT')
fallback=choose(registered({'layout_part':'ppt/slideLayouts/slideLayout6.xml','fallback_layouts':['ppt/slideLayouts/slideLayout2.xml']}),base,style)
check('explicit-authorized-fallback',fallback['selection']['fallback']['used'] and fallback['native_layout_part']=='ppt/slideLayouts/slideLayout2.xml',fallback['selection']['fallback'])
reject('boundary-missing-font',lambda:choose(registered({'typography':{'font':'ThisFontDoesNotExist'}}),base,style),'TEMPLATE_FONT_UNAVAILABLE')
reject('boundary-policy-override',lambda:choose(record,base,style,{**record['policy'],'layout_part':'ppt/slideLayouts/slideLayout1.xml'}),'TEMPLATE_REGISTERED_POLICY_CHANGED')
reject('boundary-policy-invalid',lambda:register(template,root=str(ROOT/'registry'),allowed_roots=['/runtime'],policy={'unknown':'ignored'}),'TEMPLATE_POLICY_INVALID')
reject('boundary-non-native-diagram',lambda:choose(record,{**base,'kind':'diagram'},style),'TEMPLATE_CONTENT_UNSUPPORTED')
slide=create_slide(Presentation(io.BytesIO(raw)),{'layout':{'template_binding':{'native_layout_part':normal1['native_layout_part'],'locks':normal1['locks']}}})
bad=copy.deepcopy(normal1['elements'][0]);bad['bounds'][0]+=1
reject('boundary-bind-geometry',lambda:bind_placeholder(slide,bad),'TEMPLATE_BINDING_GEOMETRY_CHANGED')
bad=copy.deepcopy(normal1['elements'][0]);bad['template_binding']['placeholder_index']=999
reject('boundary-bind-missing-placeholder',lambda:bind_placeholder(slide,bad),'TEMPLATE_PLACEHOLDER_UNAVAILABLE')
check('policy-versioned-identity',records[0]['template_id']!=records[1]['template_id'] and records[0]['sha256']==records[1]['sha256'])

# Formal product compiler/DAG/schema/writer integration on the real nonblank layout.
proposal={'title':'正式企业模板槽位','main_takeaway':'真实槽位原生可编辑','storyline':{'pattern':'结论先行','rationale':'模块限定验证'},'slides':[copy.deepcopy(base)]}
schemas.validate('generation-model.schema.json',proposal)
deck,dag=compile_deck(proposal,{**DEFAULT_STYLE,**style},trace_id(),registry=AtomicRegistry(schemas),visual={'template_selection':{'record':record}})
schemas.validate('deck-execution.schema.json',deck);schemas.validate('deck-execution.schema.json',dag)
compiled_path=ROOT/'formal-compiled.pptx';objects,execution=execute(deck,dag,compiled_path,schemas)
actual=Presentation(compiled_path).slides[0]
check('formal-compiler-nonblank-layout',str(actual.slide_layout.part.partname).lstrip('/')=='ppt/slideLayouts/slideLayout2.xml')
check('formal-writer-native-bound-objectmap',len(objects['objects'])==2 and all(actual.shapes[i].is_placeholder for i in range(2)),objects)
check('formal-writer-no-overlay',len(actual.shapes)==2,len(actual.shapes))
check('formal-writer-editable-source-text',actual.placeholders[0].text==base['title'] and actual.placeholders[1].text==base['body'][0])
recovered,_=compile_deck(proposal,{**DEFAULT_STYLE,**style},trace_id(),previous_deck=deck,registry=AtomicRegistry(schemas),visual={'template_selection':{'record':record}})
check('formal-recovery-object-ids', [e['object_id'] for e in recovered['slides'][0]['elements']]==[e['object_id'] for e in deck['slides'][0]['elements']])
with zipfile.ZipFile(io.BytesIO(raw)) as before,zipfile.ZipFile(compiled_path) as after:
    check('formal-package-preservation',all(after.read(n)==before.read(n) for n in preserved_part_names(before.namelist())))
(ROOT/'formal-deck.json').write_text(json.dumps(deck,ensure_ascii=False,indent=2),'utf-8');(ROOT/'formal-objectmap.json').write_text(json.dumps(objects,ensure_ascii=False,indent=2),'utf-8');(ROOT/'formal-execution.json').write_text(json.dumps(execution,ensure_ascii=False,indent=2),'utf-8')
(ROOT/'pre-render-checks.json').write_text(json.dumps({'checks':checks,'outputs':outputs},ensure_ascii=False,indent=2),'utf-8')

# Render the three ordinary forms and the chart within ordinary data-slot case.
for i,path in enumerate((path1,path2,path3,path4),1):
    dest=ROOT/('render-case-'+str(i));dest.mkdir()
    result=render(path,{'slides':[{'slide_id':'slide_'+str(i)}]},dest)
    check('actual-render-'+str(i),len(result['artifacts'])==2,{'pdf_png':result['artifacts'],'powerpoint_compatibility':result['powerpoint_compatibility']})
    text=(dest/'render/deck.txt').read_text('utf-8')
    check('render-brand-visible-'+str(i),'Yoloong' in text and '企业保密' in text)
    outputs[i-1]['render_root']=str(dest)

source_files=['src/yoloongppt/templates.py','src/yoloongppt/template_registry.py','src/yoloongppt/template_layout_runtime.py','validation/verify-template-slots.py','src/yoloongppt/planning.py','src/yoloongppt/writer.py','src/yoloongppt/layout_decisions.py']
schema_files=['contracts/template-brand-policy.schema.json','contracts/template-layout-runtime.schema.json','contracts/template-registry-record.schema.json']
report={'suite':'native-template-slots-1','passed':len(checks),'checks':checks,'runtime_root':str(ROOT),'normal_cases':['brand-text','native-image','native-table-and-chart'],
        'outputs':outputs,'source_sha256':{n:hashlib.sha256((Path('/workspace')/n).read_bytes()).hexdigest() for n in source_files},'schema_sha256':{n:hashlib.sha256((Path('/workspace')/n).read_bytes()).hexdigest() for n in schema_files},
        'requirements':['AC-006','TPL-002','TPL-003','TPL-004','TPL-006','TPL-008','TPL-015','TPL-016'],'limits':['component Docker verification only; real generation integration is separate','three upstream template families not accepted','all modes/all AC/PowerPoint desktop verification remain incomplete']}
(ROOT/'template-slots.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8');(ROOT/'selections.json').write_text(json.dumps(selections,ensure_ascii=False,indent=2),'utf-8')
print(json.dumps({'passed':len(checks),'report':str(ROOT/'template-slots.json'),'outputs':[x['pptx'] for x in outputs]},ensure_ascii=False))
