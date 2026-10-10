"""Real native objects/render/QA integration using saved approved literal fixtures.

No model is called. This is component evidence, never AI or system AC approval.
"""
import copy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from pptx import Presentation
from yoloongppt.artifacts import entity
from yoloongppt.atomic import AtomicRegistry
from yoloongppt.errors import TaskError
from yoloongppt.planning import DEFAULT_STYLE, compile_deck
from yoloongppt.quality import check as quality_check
from yoloongppt.quality_decisions import decide
from yoloongppt.rendering import render
from yoloongppt.schemas import SchemaRegistry, ROOT
from yoloongppt.template_registry import register
from yoloongppt.visual import prepare, finalize
from yoloongppt.writer import execute

out=Path('/runtime/visual-execution-validation');out.mkdir(parents=True,exist_ok=True)
schemas=SchemaRegistry();registry=AtomicRegistry(schemas);checks=[]
def save(name,value):
    (out/name).write_text(json.dumps(value,ensure_ascii=False,indent=2),encoding='utf-8')
def check(name,condition,observed=None):
    assert condition,(name,observed)
    checks.append({'name':name,'passed':True,'observed':observed})

samples=json.loads(Path('/runtime/visual-validation/node-samples.json').read_text('utf-8'))
fixture=next(item['input'] for item in samples if item.get('name')=='DEC-022_normal_3')
task=copy.deepcopy(fixture['task']);sources=copy.deepcopy(fixture['sources']);context=copy.deepcopy(fixture['context'])
slides=[]
for name in ('DEC-022_normal_3','DEC-027_normal_1','DEC-025_normal_1','DEC-026_normal_1'):
    slides.append(copy.deepcopy(next(item['input']['proposal']['slides'][0] for item in samples if item.get('name')==name)))
proposal=copy.deepcopy(fixture['proposal']);proposal['slides']=slides
schemas.validate('generation-model.schema.json',proposal)
before=copy.deepcopy(proposal)

# Distinguish caller template from default runtime with an explicit fixture marker.
selected_template=out/'selected-template.pptx'
with ZipFile('/runtime/visual-validation/wide-template.pptx') as source, ZipFile(selected_template,'w',ZIP_DEFLATED) as target:
    for item in source.infolist():
        raw=source.read(item.filename)
        if item.filename=='ppt/theme/theme1.xml':
            xml=ET.fromstring(raw)
            xml.find('{http://schemas.openxmlformats.org/drawingml/2006/main}themeElements/{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme').set('name','Caller component template marker')
            raw=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        target.writestr(item,raw)
template_hash=hashlib.sha256(selected_template.read_bytes()).hexdigest()
task['template_ref']=str(selected_template)
task['sources'].append({'source_id':entity('source'),'kind':'pptx','locator':str(selected_template),'metadata':{'source_role':'template','license':'unknown'}})
packet=prepare(task,sources,context,schemas)
final,visual=finalize(proposal,packet,sources,schemas)
check('prepare_finalize_preserve_original_proposal',proposal==before)
check('selected_template_is_distinct_caller_bytes',packet['template_selection']['record']['sha256']==template_hash)
deck,execution=compile_deck(final,DEFAULT_STYLE,context['trace_id'],registry=registry,visual=visual)
check('compile_preserves_resolved_image_declaration','crop' not in final['slides'][0]['image'] and 'fit_result' not in final['slides'][0]['image'])
check('four_visual_kinds_preserved',[slide['content']['kind'] for slide in deck['slides']]==['image','diagram','table','chart'])
check('layout_nine_stages_every_slide',all(len(slide['layout']['selection_trace']['decisions'])==9 for slide in deck['slides']))
check('all_required_native_atomic_bound',{'add_image','add_shape','add_connector','add_table','add_chart'}.issubset({call['capability'] for call in execution['calls']}))
save('source-fixture.json',sources);save('proposal.json',proposal);save('final-proposal.json',final)
save('visual-packet.json',packet);save('visual-result.json',visual);save('deck-spec.json',deck);save('execution-plan.json',execution)
object_map,result=execute(deck,execution,out/'deck.pptx',schemas)
save('object-map.json',object_map);save('execution-result.json',result)
native=Presentation(out/'deck.pptx')
check('native_real_four_pages',len(native.slides)==4)
check('all_native_objects_mapped',len(object_map['objects'])==sum(len(slide['elements']) for slide in deck['slides']))
image=next(shape for shape in native.slides[0].shapes if shape.shape_type==13)
check('embedded_image_byte_identity',hashlib.sha256(image.image.blob).hexdigest()==final['slides'][0]['image']['sha256'])
check('contain_no_crop',[image.crop_left,image.crop_top,image.crop_right,image.crop_bottom]==[0,0,0,0])
process=deck['slides'][1];nodes=[element for element in process['elements'] if element['type']=='shape'];edges=[element for element in process['elements'] if element['type']=='connector']
check('process_three_native_shapes_two_native_edges',len(nodes)==3 and len(edges)==2 and sum(shape.shape_type==1 for shape in native.slides[1].shapes)==3 and sum(shape.shape_type==9 for shape in native.slides[1].shapes)==2)
check('process_labels_exact_original',all(next(shape for shape in native.slides[1].shapes if shape.name==element['object_id']).text==element['data']['text'] for element in nodes))
table=next(shape.table for shape in native.slides[2].shapes if shape.has_table)
check('native_table_original_values',[[cell.text for cell in row.cells] for row in table.rows]==[slides[2]['table']['columns'],*slides[2]['table']['rows']])
chart=next(shape.chart for shape in native.slides[3].shapes if shape.has_chart)
check('native_chart_original_values',[list(series.values) for series in chart.series]==[series['values'] for series in slides[3]['chart']['series']])
with ZipFile(selected_template) as original,ZipFile(out/'deck.pptx') as written:
    protected=[name for name in original.namelist() if name.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/'))]
    mismatches=[name for name in protected if name not in written.namelist() or original.read(name)!=written.read(name)]
    check('caller_master_layout_theme_byte_preserved',bool(protected) and not mismatches,{'protected_parts':len(protected),'mismatches':mismatches})
    check('chart_embedded_editable_workbook',any(name.startswith('ppt/embeddings/') and name.endswith('.xlsx') for name in written.namelist()))
check('caller_template_unchanged',hashlib.sha256(selected_template.read_bytes()).hexdigest()==template_hash)

# Explicitly unsupported caller layout: no content-free layout. Never fall back.
unsupported=Presentation(selected_template)
blank=unsupported.slide_layouts[6]
placeholder=copy.deepcopy(unsupported.slide_layouts[0].placeholders[0]._element)
placeholder.xpath('.//p:cNvPr')[0].set('id','90');placeholder.xpath('.//p:cNvPr')[0].set('name','Caller required template title')
blank.shapes._spTree.insert_element_before(placeholder,'p:extLst')
unsupported_path=out/'unsupported-template.pptx';unsupported.save(unsupported_path)
unsupported_deck=copy.deepcopy(deck);unsupported_deck['template']=register(unsupported_path,schemas=schemas)
try:execute(unsupported_deck,execution,out/'must-not-exist.pptx',schemas)
except TaskError as error:check('unsupported_template_explicit_prewrite_failure',error.code=='EXECUTION_TEMPLATE_LAYOUT_AMBIGUOUS' and not (out/'must-not-exist.pptx').exists(),error.code)
else:raise AssertionError('unsupported template silently replaced')

renders=render(out/'deck.pptx',deck,out);save('render-report.json',renders)
check('actual_pdf_and_four_png',sum(artifact['type']=='pdf' for artifact in renders['artifacts'])==1 and sum(artifact['type']=='png' for artifact in renders['artifacts'])==4)
check('every_render_hash_matches_bytes',all(hashlib.sha256((out/artifact['path']).read_bytes()).hexdigest()==artifact['hash'] for artifact in renders['artifacts']))
qa=quality_check(deck,out/'deck.pptx',object_map,sources,renders,out);save('quality-report.json',qa)
decisions=decide({'quality_report':qa,'deck':deck,'object_map':object_map},schemas,context['trace_id']);save('quality-decisions.json',decisions)
check('actual_component_qa_no_p0',qa['p0_issue_count']==0,[{'code':item['code'],'severity':item['severity']} for item in qa['issues']])
check('DEC039040_actual_report_consumed',[item['node_id'] for item in decisions['decision_traces']]==['DEC-039','DEC-040'])
check('incomplete_semantic_AC_cannot_finalize',decisions['completion']['state']=='warning' and decisions['completion']['final_allowed'] is False and decisions['completion']['missing_acceptance']==['AC-001','GOV-008'])
hashes={str(path.relative_to(ROOT)):hashlib.sha256(path.read_bytes()).hexdigest() for path in (ROOT/'src/yoloongppt').glob('*.py')}
report={'ok':True,'check_count':len(checks),'checks':checks,'source_hashes':hashes,'schema_hashes':schemas.versions,
    'p0_issue_count':qa['p0_issue_count'],'warning_issue_codes':[item['code'] for item in qa['issues'] if item['severity']!='P0'],
    'completion':decisions['completion'],'actual_scope':'Saved approved literal test fixtures → prepare/finalize → native compiler/atomic execution → real four-page PPTX with selected template preserved → LibreOffice/PDF/PNG → deterministic QA → DEC039/040. No model call, no semantic/AI/system AC claim.'}
save('visual-execution.json',report)
print(json.dumps({'ok':True,'checks':len(checks),'p0':qa['p0_issue_count'],'completion':decisions['completion']['state'],'report':str(out/'visual-execution.json')}))
