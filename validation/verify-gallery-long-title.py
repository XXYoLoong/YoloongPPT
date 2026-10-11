"""One incremental native 8-image long-title fixture; no original-suite replay."""
import copy,hashlib,json
from pathlib import Path
from uuid import uuid4
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from yoloongppt.visual import finalize
from yoloongppt.gallery_layout import evaluate
from yoloongppt.planning import compile_deck,DEFAULT_STYLE
from yoloongppt.writer import execute
from yoloongppt.rendering import render
from yoloongppt.quality_rules import audit,package_audit
from yoloongppt.schemas import SchemaRegistry
from yoloongppt.errors import trace_id
from yoloongppt.artifacts import entity

parent=Path('/runtime/gallery-validation/b81e5a47-0d59-4963-87a5-daa257d654ef')
root=Path('/runtime/gallery-long-title-validation')/str(uuid4());root.mkdir(parents=True)
load=lambda name:json.loads((parent/name).read_text('utf-8'))
def save(name,value):(root/name).write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n','utf-8')
checks=[]
def test(name,ok,details):
    checks.append({'name':name,'passed':bool(ok),'observation':details});save('incremental-checks.json',checks)
    if not ok:save('failure.json',{'failed_check':name,'checks':checks})
    assert ok,name
schema=SchemaRegistry();content=load('declared-gallery-fixture.json')
long_title='完整保留混合长标题并逐图显示不同宽高比真实素材与原据图注。'*3
content.update(title=long_title,goal='人工长标题八图容量验证',rationale='人工fixture，仅验证真实几何和可编辑图片',table=None,chart=None)
sources=load('approved-caption-fixture.json')
for eid in content['evidence_refs']:sources['approved_texts'][eid]+='\n'+long_title
packet=json.loads(Path('/runtime/runs/39a2dbf3-c19b-4da2-99a6-27fe7b82ef7a/visual-packet.json').read_text('utf-8'));packet['resolved_assets']=load('resolved-assets.json')
proposal={'title':'长标题八图增量验证','main_takeaway':'全部人工素材按原据显示','storyline':{'pattern':'linear','rationale':'原图顺序保持'},'slides':[content]}
proposal,visual=finalize(proposal,packet,sources,schemas=schema)
schema.validate('generation-model.schema.json',proposal)
save('source-fixture.json',sources);save('model-independent-proposal.json',proposal);save('visual-result.json',visual)
deck,execution=compile_deck(proposal,copy.deepcopy(DEFAULT_STYLE),trace_id())
save('deck-spec.json',deck);save('execution-plan.json',execution)
selection=deck['slides'][0]['layout']['selection_trace']['gallery_selection'];chosen=next(c for c in selection['candidates'] if c['candidate']==selection['selected'])
title=next(e for e in deck['slides'][0]['elements'] if e['slot_key']=='title')
test('long-title-native-box-expanded-no-font-reduction',title['text']==[long_title] and title['font_size']==30 and title['bounds'][3]>.95 and chosen['adaptive_title']['expanded'],{'characters':len(long_title),'title_bounds':title['bounds'],'trace':chosen['adaptive_title']})
test('eight-native-images-captions-preserved-at-actual-dpi',len([e for e in deck['slides'][0]['elements'] if e['type']=='image'])==8 and len([e for e in deck['slides'][0]['elements'] if e['slot_key'].startswith('caption.')])==8 and all(i['effective_dpi_unrounded']>=i['minimum_dpi'] for i in chosen['images']),{'selected':selection['selected'],'images':chosen['images']})
mapping,calls=execute(deck,execution,root/'long-title-gallery.pptx',schema);save('object-map.json',mapping);save('execution-trace.json',calls)
presentation=Presentation(root/'long-title-gallery.pptx')
test('shared-compiler-writer-eight-editable-pictures-executed',sum(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in presentation.slides[0].shapes)==8 and any(s.has_text_frame and s.text==long_title for s in presentation.slides[0].shapes),{'call_count':len(calls['calls']),'mapped_object_count':len(mapping['objects'])})
render_report=render(root/'long-title-gallery.pptx',deck,root);save('render-report.json',render_report)
test('one-real-page-render',len([a for a in render_report['artifacts'] if a['type']=='png'])==1,render_report['artifacts'])
issues,qa_checks,coverage=audit(deck,presentation,mapping);package_issues,package_checks=package_audit(root/'long-title-gallery.pptx')
save('native-quality.json',{'issues':issues+package_issues,'checks':qa_checks+package_checks,'coverage':coverage})
test('long-title-gallery-native-quality-no-issues',not issues and not package_issues,{'issues':issues+package_issues,'coverage':coverage})
hash_paths=['src/yoloongppt/gallery_layout.py','src/yoloongppt/visual.py','src/yoloongppt/planning.py','src/yoloongppt/layouts.py','src/yoloongppt/writer.py','src/yoloongppt/adaptive_layout.py','contracts/generation-model.schema.json','contracts/deck-execution.schema.json','validation/verify-gallery-long-title.py']
report={'passed':all(c['passed'] for c in checks),'checks':checks,'source_hashes':{p:hashlib.sha256((Path('/workspace')/p).read_bytes()).hexdigest() for p in hash_paths},'parent_fixture_id':parent.name,'reuse_input_assets':[{'asset_id':a['asset_id'],'sha256':a['sha256'],'path':a['path']} for a in packet['resolved_assets']['entries']],'fixture_identity':'synthetic long-title variant; reuses original 8 artificial raster inputs; no AI/model generation','model_calls':0,'render_calls':1,'system_acceptance':'not_passed','output_path':str(root)}
save('gallery-long-title-runtime.json',report)
print(json.dumps({'passed':report['passed'],'checks':len(checks),'output_path':str(root)},ensure_ascii=False))
