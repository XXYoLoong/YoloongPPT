"""RES-032/PPT-015/SYS-012/013: direct reuse of the pinned PPT Master engine.

The upstream exporter, importer, OPC audit and canvas/relationship guards remain
unmodified under vendor/ppt_master. An isolated Python worker prevents its flat
module names and console bootstrap from contaminating the application's imports.
No raster exporter or silent nonconforming/compatibility fallback is enabled.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET

from .errors import TaskError

VENDOR = Path(__file__).resolve().parents[2] / 'vendor' / 'ppt_master'
UPSTREAM_COMMIT = '2d72da616cf9fa40d4dcaf59fd4c980ecf534b7d'
IDS = ['RES-032', 'PPT-015', 'PPT-030', 'SYS-012', 'SYS-013']
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def _fail(code, message, details=None):
    raise TaskError(code, message, 'PPTMasterAdapter', IDS, details or [])


def _hash(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def upstream_version():
    manifest = json.loads((VENDOR / 'SOURCE.json').read_text(encoding='utf-8'))
    return {'project': 'PPT Master', 'repository': manifest['repository'],
            'commit': manifest['commit'], 'license': 'MIT',
            'third_party_data': ['Apache POI Apache-2.0', 'Open XML SDK MIT'],
            'manifest_sha256': _hash(VENDOR / 'SOURCE.json')}


def _activate():
    manifest = json.loads((VENDOR / 'SOURCE.json').read_text(encoding='utf-8'))
    if manifest['commit'] != UPSTREAM_COMMIT:
        _fail('UPSTREAM_VERSION_MISMATCH', 'PPT Master 固定版本不一致。')
    for row in manifest['files']:
        if _hash(VENDOR / row['path']) != row['sha256']:
            _fail('UPSTREAM_SOURCE_CHANGED', '上游模块或许可通知已变更。', [row])
    sys.path.insert(0, str(VENDOR / 'scripts'))


def _svg_guard(path, resource_root):
    from .input_formats import MAX_RAW, xml
    path = Path(path).resolve()
    raw = path.read_bytes()
    if not raw or len(raw) > MAX_RAW:
        _fail('SVG_INPUT_LIMIT', 'SVG 字节数超过已配置输入限制。')
    root = xml(raw)
    for node in root.iter():
        tag = node.tag.rsplit('}', 1)[-1]
        if tag in {'script', 'foreignObject'} or any(k.lower().startswith('on') for k in node.attrib):
            _fail('SVG_ACTIVE_CONTENT_BLOCKED', 'SVG 含活动内容或未支持的 HTML 对象。')
        if tag=='a':
            link=node.get('href') or node.get('{http://www.w3.org/1999/xlink}href') or ''
            if link and not link.startswith('#'):
                parsed_link=urlsplit(link)
                if parsed_link.scheme not in {'https','http','mailto'} or parsed_link.username or parsed_link.password or (parsed_link.scheme in {'https','http'} and not parsed_link.hostname):
                    _fail('SVG_LINK_UNSAFE','SVG链接不允许活动脚本、文件或含凭据的URL。')
        if tag not in {'image', 'use'}:
            continue
        value = node.get('href') or node.get('{http://www.w3.org/1999/xlink}href') or ''
        if value.startswith('data:') and not value.lower().startswith(('data:image/png;','data:image/jpeg;')):
            _fail('SVG_INLINE_ASSET_UNSUPPORTED','内联SVG/活动文档需独立安全解析，当前只允许内联PNG/JPEG；未隐式渲染。')
        if not value or value.startswith(('#', 'data:')):
            continue
        parsed = urlsplit(value)
        if parsed.scheme or parsed.netloc:
            _fail('SVG_EXTERNAL_ASSET_BLOCKED', 'SVG 外部资产须先安全取得并登记为本地资源。')
        asset = (path.parent / unquote(parsed.path)).resolve()
        try:
            asset.relative_to(resource_root)
        except ValueError:
            _fail('SVG_ASSET_PATH_ESCAPE', 'SVG 资产越过本次资源根目录。')
        if not asset.is_file():
            _fail('SVG_ASSET_MISSING', 'SVG 引用的本地资产缺失。', [{'path': str(asset)}])
        from .native_objects import local_asset
        local_asset(asset)
    return root


def _package_guard(path):
    from .input_formats import MAX_RAW, package, xml
    raw = Path(path).read_bytes()
    if len(raw) > MAX_RAW:
        _fail('PPTX_INPUT_LIMIT', 'PPTX 超过已配置输入限制。')
    with package(raw, 'ppt/presentation.xml') as archive:
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                xml(archive.read(name))


def _objects(path):
    import zipfile
    rows = []
    with zipfile.ZipFile(path) as archive:
        slide_parts = sorted((name for name in archive.namelist()
                              if name.startswith('ppt/slides/slide') and name.endswith('.xml')),
                             key=lambda name: int(Path(name).stem[5:]))
        for part in slide_parts:
            root = ET.fromstring(archive.read(part))
            for node in root.findall('.//p:spTree/*', NS):
                if node.tag.rsplit('}', 1)[-1] in {'nvGrpSpPr', 'grpSpPr'}:
                    continue
                identity = node.find('.//p:cNvPr', NS)
                rows.append({'slide_part': part, 'kind': node.tag.rsplit('}', 1)[-1],
                             'id': identity.get('id') if identity is not None else None,
                             'name': identity.get('name') if identity is not None else None,
                             'text': ''.join(t.text or '' for t in node.findall('.//a:t', NS)),
                             'custom_geometry': node.find('.//a:custGeom', NS) is not None})
    return rows


def _audit(path):
    from pptx_delivery_check import audit_pptx_delivery
    _package_guard(path)
    return audit_pptx_delivery(path)


def _run(operation, payload, *, workspace, timeout=120):
    workspace = Path(workspace).resolve()
    workspace.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='ppt-master-', dir=workspace) as folder:
        request, response = Path(folder) / 'request.json', Path(folder) / 'response.json'
        request.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
        env = dict(os.environ, PYTHONDONTWRITEBYTECODE='1', PYTHONUTF8='1',
                   TMPDIR=folder, TEMP=folder, TMP=folder)
        env['PYTHONPATH'] = os.pathsep.join(filter(None, [str(Path(__file__).resolve().parents[1]),
                                                         env.get('PYTHONPATH', '')]))
        try:
            result = subprocess.run([sys.executable, '-B', '-m', 'yoloongppt.upstream_native',
                                     '_worker', operation, str(request), str(response)],
                                    capture_output=True, text=True, encoding='utf-8',
                                    errors='replace', timeout=timeout, env=env)
        except subprocess.TimeoutExpired:
            _fail('UPSTREAM_EXECUTION_TIMEOUT', 'PPT Master 操作超时，未返回成功产物。')
        if not response.is_file():
            _fail('UPSTREAM_EXECUTION_FAILED', 'PPT Master 未返回结构化执行结果。',
                  [{'returncode': result.returncode, 'stderr': result.stderr[-8000:]}])
        value = json.loads(response.read_text(encoding='utf-8'))
        if not value['ok']:
            error = value['error']
            _fail(error['code'], error['message'], error.get('details'))
        return value['result']


def export_svg_deck(svg_paths, output, *, title, notes=None, resource_root=None):
    """Export authored SVG in the explicit flat native layout route.

    This entry has no supplied template; it does not substitute a default template
    for a template request. Template-preserving edits use the roundtrip entry.
    """
    files = [Path(path).resolve() for path in svg_paths]
    if not files or len({path.stem for path in files}) != len(files):
        _fail('SVG_ROSTER_INVALID', 'SVG 页序须非空且页文件名不可重复。')
    if not isinstance(title, str) or not title.strip():
        _fail('SVG_TITLE_REQUIRED', '原生导出须提供明确标题。')
    resource_root = Path(resource_root or files[0].parent).resolve()
    output = Path(output).resolve()
    if output.suffix.lower() != '.pptx':
        _fail('PPTX_OUTPUT_REQUIRED', '输出路径须为 PPTX。')
    if notes is not None and (not isinstance(notes, dict) or
                             set(notes) - {p.stem for p in files} or
                             any(not isinstance(v, str) for v in notes.values())):
        _fail('SVG_NOTES_ROSTER_INVALID', '备注须逐页文件名绑定，不能绑定缺失页面。')
    return _run('export', {'files': [str(p) for p in files], 'output': str(output),
                          'resource_root': str(resource_root), 'title': title, 'notes': notes},
                workspace=output.parent)


def audit_native_deck(path):
    """Return upstream read-only delivery audit; not full product/semantic QA."""
    path = Path(path).resolve()
    return _run('audit', {'path': str(path)}, workspace=Path(tempfile.gettempdir()))


def import_native_deck(source, workspace):
    """Create the upstream source-preserving workspace for editing existing PPTX."""
    return _run('import', {'source': str(Path(source).resolve()),
                           'workspace': str(Path(workspace).resolve())}, workspace=workspace)


def export_native_roundtrip(workspace, output):
    """Export that workspace with unchanged source objects restored natively."""
    output = Path(output).resolve()
    return _run('roundtrip', {'workspace': str(Path(workspace).resolve()), 'output': str(output)},
                workspace=output.parent)


def canvas_guard(svg_path):
    path = Path(svg_path).resolve()
    return _run('canvas', {'path': str(path)}, workspace=Path(tempfile.gettempdir()))


def relationship_guard(xml_fragment):
    """Use upstream QName guard; presence needs explicit relationship rebinding."""
    return _run('relationships', {'xml': _xml_text(xml_fragment)}, workspace=Path(tempfile.gettempdir()))


def chart_edit_guard(chart_xml):
    """Return upstream classification before category-data edits; never mutate."""
    return _run('chart_guard', {'xml': _xml_text(chart_xml)}, workspace=Path(tempfile.gettempdir()))


def table_merge_guard(table_xml):
    """Return upstream merge topology used to protect unsupported cell edits."""
    return _run('table_guard', {'xml': _xml_text(table_xml)}, workspace=Path(tempfile.gettempdir()))


def _xml_text(value):
    if isinstance(value, bytes):
        try:
            return value.decode('utf-8')
        except UnicodeDecodeError:
            _fail('XML_ENCODING_UNSUPPORTED', '关系/编辑 guard 须接收 UTF-8 XML。')
    if not isinstance(value, str):
        _fail('XML_FRAGMENT_REQUIRED', '关系/编辑 guard 须接收 XML 文本或字节。')
    return value


def insert_native_svg(slide, svg_path, object_id, bounds, resource_root=None):
    """Insert upstream-converted editable vectors in an identified native group.

    Bounds use the writer's inches with explicit contain placement. Font/line
    units scale with the vectors because OOXML group transforms do not scale
    those absolute units consistently in LibreOffice. The target slide/master/theme and
    relationships are retained. SVG requiring image/chart/hyperlink relationships
    uses export_svg_deck/roundtrip instead; it fails before target mutation here.
    """
    import math
    from pptx.oxml import parse_xml
    from pptx.util import Inches
    path = Path(svg_path).resolve()
    if not isinstance(object_id, str) or not object_id:
        _fail('SVG_OBJECT_ID_REQUIRED', '原生 SVG 必须绑定稳定 object_id。')
    if len(bounds) != 4 or any(not isinstance(v, (int, float)) or isinstance(v, bool)
                               or not math.isfinite(v) for v in bounds) or min(bounds[2:]) <= 0:
        _fail('SVG_BOUNDS_INVALID', '原生 SVG bounds 须为有效英寸坐标与正尺寸。')
    if any(node.get('name') == object_id for node in slide._element.iter()
           if node.tag.rsplit('}', 1)[-1] == 'cNvPr'):
        _fail('SVG_OBJECT_ID_DUPLICATE', '目标 slide 已有该 object_id。')
    value = _run('shapes', {'path': str(path),
                            'resource_root': str(Path(resource_root or path.parent).resolve())},
                 workspace=Path(tempfile.gettempdir()))
    source_tree = parse_xml(value['slide_xml'].encode('utf-8')).find('.//p:spTree', NS)
    nodes = [node for node in source_tree if node.tag.rsplit('}', 1)[-1]
             not in {'nvGrpSpPr', 'grpSpPr'}]
    if not nodes:
        _fail('SVG_NATIVE_OBJECT_EMPTY', 'SVG 未产生可插入的原生对象。')
    highest = max((int(node.get('id')) for node in slide._element.iter()
                   if node.tag.rsplit('}', 1)[-1] == 'cNvPr'), default=1)
    # Group receives the next existing-library ID; nested objects follow it.
    id_map, identities = {}, []
    for node in nodes:
        for identity in node.iter():
            if identity.tag.rsplit('}', 1)[-1] == 'cNvPr':
                old_id = identity.get('id')
                new_id = highest + 2 + len(id_map)
                if old_id in id_map:
                    _fail('SVG_NATIVE_ID_DUPLICATE', '上游转换产生重复原生对象 ID。')
                id_map[old_id] = new_id
                identities.append({'source_id': old_id, 'shape_id': new_id,
                                   'source_name': identity.get('name')})
                identity.set('id', str(new_id))
    for node in nodes:
        for child in node.iter():
            if child.tag.rsplit('}', 1)[-1] in {'stCxn', 'endCxn'}:
                old_id = child.get('id')
                if old_id not in id_map:
                    _fail('SVG_CONNECTOR_SCOPE_INVALID', 'SVG connector 引用本对象范围外的 shape。')
                child.set('id', str(id_map[old_id]))
    canvas_w, canvas_h = value['size_emu']
    box_w, box_h = Inches(bounds[2]), Inches(bounds[3])
    scale = min(box_w / canvas_w, box_h / canvas_h)
    fitted_w, fitted_h = round(canvas_w * scale), round(canvas_h * scale)
    offset_x = Inches(bounds[0]) + round((box_w - fitted_w) / 2)
    offset_y = Inches(bounds[1]) + round((box_h - fitted_h) / 2)
    absolute_units = {'rPr': {'sz', 'kern', 'spc'}, 'defRPr': {'sz', 'kern', 'spc'},
                      'endParaRPr': {'sz', 'kern', 'spc'}, 'ln': {'w'},
                      'bodyPr': {'lIns', 'tIns', 'rIns', 'bIns', 'spcCol'},
                      'pPr': {'marL', 'marR', 'indent', 'defTabSz'},
                      'spcPts': {'val'}, 'buSzPts': {'val'}, 'tab': {'pos'},
                      'outerShdw': {'blurRad', 'dist'}, 'innerShdw': {'blurRad', 'dist'},
                      'glow': {'rad'}, 'softEdge': {'rad'}}
    scaled_units = []
    for node in nodes:
        for child in node.iter():
            tag = child.tag.rsplit('}', 1)[-1]
            for attribute in absolute_units.get(tag, ()):
                raw = child.get(attribute)
                if raw is None:
                    continue
                before = int(raw)
                after = round(before * scale)
                if attribute == 'sz' and after <= 0:
                    _fail('SVG_TEXT_SCALE_INVALID', '原生 SVG 缩放产生不可表示的字号。')
                child.set(attribute, str(after))
                scaled_units.append({'tag': tag, 'attribute': attribute,
                                     'before': before, 'after': after})
    group = slide.shapes.add_group_shape()
    group.name = object_id
    for node in nodes:
        group.shapes._spTree.insert_element_before(node, 'p:extLst')
    transform = group._element.grpSpPr.get_or_add_xfrm()
    transform.off.x, transform.off.y = offset_x, offset_y
    transform.ext.cx, transform.ext.cy = fitted_w, fitted_h
    transform.chOff.x, transform.chOff.y = 0, 0
    transform.chExt.cx, transform.chExt.cy = value['size_emu']
    return {'shape_id': group.shape_id, 'shape_name': group.name, 'object_id': object_id,
            'native_svg': True, 'shape_ids': identities, 'source_sha256': _hash(path),
            'upstream': value['upstream'], 'conversion': value['conversion'],
            'placement': {'fit': 'contain', 'scale': scale, 'bounds_emu':
                          [offset_x, offset_y, fitted_w, fitted_h],
                          'scaled_absolute_units': scaled_units},
            'relationship_policy': 'no relationships added or overwritten'}


def _perform(operation, payload):
    _activate()
    if operation == 'audit':
        return {'upstream': upstream_version(), 'audit': _audit(payload['path']),
                'native_objects': _objects(payload['path'])}
    if operation == 'relationships':
        from pptx_shapes.xml_safety import has_relationship_attributes
        from .input_formats import xml
        return {'has_relationship_attributes': has_relationship_attributes(xml(payload['xml'].encode('utf-8')))}
    if operation in {'chart_guard', 'table_guard'}:
        from pptx_ooxml.edit_safety import _chart_edit_capability, _table_merge_topology
        from .input_formats import xml
        node = xml(payload['xml'].encode('utf-8'))
        return _chart_edit_capability(node) if operation == 'chart_guard' else _table_merge_topology(node)
    if operation == 'canvas':
        from dataclasses import asdict
        from svg_to_pptx.canvas_contract import parse_project_svg_root, require_powerpoint_slide_size
        viewbox = parse_project_svg_root(_svg_guard(payload['path'], Path(payload['path']).parent))
        return {'viewbox': {k: str(v) for k, v in asdict(viewbox).items()},
                'size_emu': list(require_powerpoint_slide_size(viewbox))}
    if operation == 'shapes':
        from svg_to_pptx.canvas_contract import parse_project_svg_root, require_powerpoint_slide_size
        from svg_to_pptx.drawingml.converter import convert_svg_to_slide_shapes
        from pptx_shapes.xml_safety import has_relationship_attributes
        path, resource_root = Path(payload['path']), Path(payload['resource_root'])
        viewbox = parse_project_svg_root(_svg_guard(path, resource_root))
        traces = []
        slide_xml, media, rels, animation, parts, content_types = convert_svg_to_slide_shapes(
            path, resource_root=resource_root, verbose=False, image_optimize=False,
            native_objects=True, promote_background=False, trace_out=traces)
        if media or rels or parts or has_relationship_attributes(ET.fromstring(slide_xml)):
            _fail('SVG_INSERT_RELATIONSHIPS_UNSUPPORTED',
                  '此 SVG 含图像/图表/链接等关系对象；须采用上游整页或保真 roundtrip 路线。')
        if any(trace.get('summary', {}).get('skipped', 0) for trace in traces):
            _fail('SVG_NATIVE_OBJECT_SKIPPED', 'SVG 含未转换对象，拒绝静默丢失。', traces)
        return {'slide_xml': slide_xml, 'size_emu': list(require_powerpoint_slide_size(viewbox)),
                'upstream': upstream_version(), 'conversion': traces}
    if operation == 'import':
        from pptx_to_svg.converter import ConvertOptions, convert_pptx_to_svg
        _package_guard(payload['source'])
        workspace = Path(payload['workspace'])
        if (workspace / 'sources' / 'source.pptx').exists():
            _fail('NATIVE_WORKSPACE_EXISTS', '原生编辑工作区已有源文件；拒绝覆盖。')
        result = convert_pptx_to_svg(Path(payload['source']), workspace,
                                     ConvertOptions(inheritance_mode='both', roundtrip=True))
        return {'upstream': upstream_version(), 'workspace': str(workspace),
                'source_sha256': _hash(payload['source']), 'slide_count': len(result.slides),
                'manifest': str(workspace / 'analysis' / 'roundtrip_manifest.json'),
                'authoring_directory': str(workspace / 'authoring-svg-flat')}
    output = Path(payload['output'])
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='native-output-', dir=output.parent) as folder:
        staged = Path(folder) / 'deck.pptx'
        trace = Path(folder) / 'conversion-trace.json'
        if operation == 'export':
            from svg_to_pptx.pptx_package.builder import create_pptx_with_native_svg
            files, resource_root = [Path(p) for p in payload['files']], Path(payload['resource_root'])
            for path in files:
                _svg_guard(path, resource_root)
            ok = create_pptx_with_native_svg(files, staged, resource_root=resource_root,
                                            verbose=False, transition=None, use_compat_mode=False,
                                            notes=payload['notes'], enable_notes=payload['notes'] is not None,
                                            use_native_shapes=True, native_objects=True,
                                            pptx_structure='flat',
                                            conversion_trace_path=trace, image_optimize=False,
                                            workers=1, doc_metadata={'title': payload['title']})
            if not ok:
                _fail('SVG_NATIVE_EXPORT_INCOMPLETE', '原生 SVG 导出未完成全部页面。')
        elif operation == 'roundtrip':
            from svg_to_pptx.pptx_package.cli import main
            code = main([payload['workspace'], '--roundtrip', '-o', str(staged),
                         '--native-charts-and-tables', '--conversion-trace', str(trace), '-q'])
            if code:
                _fail('NATIVE_ROUNDTRIP_FAILED', '上游保真编辑导出失败。', [{'exit_code': code}])
        else:
            _fail('UPSTREAM_OPERATION_UNKNOWN', '未知上游操作。')
        audit = _audit(staged)
        if audit['errors']:
            _fail('NATIVE_DELIVERY_AUDIT_FAILED', 'PPTX 未通过上游结构交付检查。', audit['errors'])
        conversion = json.loads(trace.read_text(encoding='utf-8'))
        skipped = [entry for entry in conversion['slides'] if entry.get('summary', {}).get('skipped', 0)]
        if skipped:
            _fail('SVG_NATIVE_OBJECT_SKIPPED', 'SVG 含未转换对象，拒绝静默丢失。', skipped)
        objects = _objects(staged)
        os.replace(staged, output)
        conversion['output'] = str(output)
        trace_path = output.with_suffix('.native-trace.json')
        trace_path.write_text(json.dumps(conversion, ensure_ascii=False, indent=2), encoding='utf-8')
        audit['file']['path'] = str(output)
        return {'output': str(output), 'sha256': _hash(output), 'upstream': upstream_version(),
                'inputs': [{'path': path, 'sha256': _hash(path)} for path in payload.get('files', [])],
                'native_objects': objects, 'delivery_audit': audit,
                'conversion_trace': str(trace_path), 'conversion': conversion,
                'limitations': ['Upstream structure audit is not visual/fact/system acceptance.',
                                'Only supported native objects; unsupported paths fail explicitly.']}


if __name__ == '__main__':
    if len(sys.argv) != 5 or sys.argv[1] != '_worker':
        raise SystemExit('PPT Master adapter worker requires operation, request and response paths.')
    try:
        result = _perform(sys.argv[2], json.loads(Path(sys.argv[3]).read_text(encoding='utf-8')))
        response = {'ok': True, 'result': result}
    except TaskError as error:
        response = {'ok': False, 'error': error.result('upstream-worker')['error']}
    except ModuleNotFoundError as error:
        response = {'ok': False, 'error': {'code': 'UPSTREAM_DEPENDENCY_MISSING',
                    'message': '上游模块依赖缺失。', 'details': [{'module': error.name}]}}
    except Exception as error:
        response = {'ok': False, 'error': {'code': 'UPSTREAM_NATIVE_FAILED',
                    'message': '上游原生操作失败，未降级为图片。',
                    'details': [{'type': type(error).__name__, 'reason': str(error)}]}}
    Path(sys.argv[4]).write_text(json.dumps(response, ensure_ascii=False, indent=2), encoding='utf-8')
