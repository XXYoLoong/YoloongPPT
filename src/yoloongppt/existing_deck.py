"""REV-001/002/004/005/010: preserve-first, explicitly targeted OOXML edits."""
import copy
import hashlib
import io
import json
import math
import os
import posixpath
import re
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation

from .errors import TaskError
from .native_objects import chart_data, local_asset

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'


def fail(code, message, details=None):
    raise TaskError(code, message, 'ExistingDeck', ['REV-001', 'REV-002', 'REV-004', 'REV-005'], details or [])


def _xml(blob):
    if b'<!DOCTYPE' in blob or b'<!ENTITY' in blob: fail('PPTX_XML_UNSAFE', '不允许外部实体或 DTD。')
    return etree.fromstring(blob, etree.XMLParser(resolve_entities=False, no_network=True))


def _package(path):
    path = Path(path)
    if path.stat().st_size > 128*1024*1024: fail('PPTX_SIZE_LIMIT', 'PPTX 超过 128 MiB 显式读取上限。')
    with zipfile.ZipFile(path) as archive:
        info = archive.infolist(); names = [x.filename for x in info]
        if len(info) > 20000 or sum(x.file_size for x in info) > 256*1024*1024: fail('PPTX_SIZE_LIMIT', 'PPTX 解压内容超过显式上限。')
        if len(names) != len(set(names)) or any(n.startswith('/') or '\\' in n or '..' in n.split('/') for n in names): fail('PPTX_PACKAGE_INVALID', '部件路径重复或不安全。')
        parts = {x.filename: archive.read(x) for x in info}
    if 'ppt/presentation.xml' not in parts: fail('PPTX_PACKAGE_INVALID', '缺少演示文稿主部件。')
    for name,blob in parts.items():
        if name.endswith(('.xml','.rels')): _xml(blob)
    return info, parts


def _rels(parts, name):
    folder, base = posixpath.split(name); relname = folder+'/_rels/'+base+'.rels'
    if relname not in parts: return {}
    return {x.get('Id'): {'type': x.get('Type'), 'external': x.get('TargetMode') == 'External',
                        'target': x.get('Target') if x.get('TargetMode') == 'External' else posixpath.normpath(posixpath.join(folder, x.get('Target').lstrip('/'))) if not x.get('Target').startswith('/') else x.get('Target').lstrip('/')}
            for x in _xml(parts[relname])}


def _slides(parts):
    main = _xml(parts['ppt/presentation.xml']); rels = _rels(parts, 'ppt/presentation.xml')
    return [(int(x.get('id')), rels[x.get('{'+NS['r']+'}id')]['target']) for x in main.xpath('./p:sldIdLst/p:sldId', namespaces=NS)]


def _objects(root):
    result = []
    for node in root.xpath('./p:cSld/p:spTree//*[self::p:sp or self::p:pic or self::p:graphicFrame or self::p:cxnSp or self::p:grpSp]', namespaces=NS):
        own = node.xpath('./p:nvSpPr/p:cNvPr|./p:nvPicPr/p:cNvPr|./p:nvGraphicFramePr/p:cNvPr|./p:nvCxnSpPr/p:cNvPr|./p:nvGrpSpPr/p:cNvPr', namespaces=NS)
        if own: result.append((node, own[0]))
    return result


def inspect_deck(path):
    """No native object execution; stable identity survives text/style patches."""
    _,parts = _package(path); slides=[]
    for index,(slide_id,part) in enumerate(_slides(parts)):
        root = _xml(parts[part]); objects=[]
        for order,(node,identity) in enumerate(_objects(root)):
            shape_id = int(identity.get('id'))
            geometry = node.xpath('./p:spPr/a:xfrm|./p:xfrm|./p:grpSpPr/a:xfrm', namespaces=NS)
            g = geometry[0] if geometry else None
            off = g.find('a:off', NS) if g is not None else None; ext = g.find('a:ext', NS) if g is not None else None
            ph = node.xpath('./*/p:nvPr/p:ph', namespaces=NS)
            item = {'object_id': f'pptx:s{slide_id}:shape{shape_id}', 'shape_id': shape_id, 'name': identity.get('name',''),
                    'kind': etree.QName(node).localname, 'order': order, 'part': part,
                    'text': '\n'.join(''.join(p.itertext()) for p in node.xpath('./p:txBody/a:p', namespaces=NS)),
                    'alt_text': identity.get('descr',''), 'placeholder': dict(ph[0].attrib) if ph else None,
                    'bounds_emu': [int(off.get('x','0')),int(off.get('y','0')),int(ext.get('cx','0')),int(ext.get('cy','0'))] if off is not None and ext is not None else None,
                    'format_xml': [etree.tostring(x).decode('utf-8') for x in node.xpath('./p:spPr|./p:txBody/a:bodyPr|./p:txBody/a:p/a:pPr|./p:txBody/a:p/a:r/a:rPr', namespaces=NS)]}
            table = node.xpath('.//a:tbl', namespaces=NS)
            if table: item['table'] = [[ '\n'.join(''.join(p.itertext()) for p in c.findall('a:txBody/a:p',NS)) for c in r.findall('a:tc',NS)] for r in table[0].findall('a:tr',NS)]
            chart = node.xpath('.//c:chart', namespaces=NS)
            if chart:
                item['chart_part'] = _rels(parts,part)[chart[0].get('{'+NS['r']+'}id')]['target']
                chart_xml=_xml(parts[item['chart_part']])
                item['chart']={'types':[etree.QName(x).localname for x in chart_xml.xpath('//c:plotArea/*[c:ser]',namespaces=NS)],
                               'series':[{'name':ser.xpath('./c:tx//c:v/text()',namespaces=NS),
                                          'categories':ser.xpath('./c:cat//c:pt/c:v/text()',namespaces=NS),
                                          'values':ser.xpath('./c:val//c:pt/c:v/text()|./c:yVal//c:pt/c:v/text()',namespaces=NS)}
                                         for ser in chart_xml.xpath('//c:ser',namespaces=NS)]}
            blip=node.xpath('./p:blipFill/a:blip',namespaces=NS)
            if blip:
                item['image_relationship']=_rels(parts,part).get(blip[0].get('{'+NS['r']+'}embed') or blip[0].get('{'+NS['r']+'}link'))
                rect=node.xpath('./p:blipFill/a:srcRect',namespaces=NS)
                item['crop']=dict(rect[0].attrib) if rect else {}
            objects.append(item)
        related = _rels(parts,part)
        notes_parts = [v['target'] for v in related.values() if v['type'].endswith('/notesSlide') and not v['external']]
        slides.append({'slide_id':slide_id,'slide_index':index,'part':part,'hidden':root.get('show') == '0',
                       'objects':objects,'relationships':related,
                       'notes': [{'part':n,'text':'\n'.join(_xml(parts[n]).xpath('//a:t/text()',namespaces=NS))} for n in notes_parts]})
    components={'masters':[],'layouts':[],'themes':[],'comments':[]}
    for name,blob in parts.items():
        if not name.endswith('.xml'): continue
        key=next((key for prefix,key in [('ppt/slideMasters/','masters'),('ppt/slideLayouts/','layouts'),('ppt/theme/','themes'),('ppt/comments/','comments')] if name.startswith(prefix)),None)
        if key:
            xml=_xml(blob)
            entry={'part':name,'attributes':dict(xml.attrib),'relationships':_rels(parts,name)}
            if key in {'masters','layouts'}: entry['placeholders']=[dict(x.attrib) for x in xml.xpath('//p:ph',namespaces=NS)]
            if key=='themes':
                entry['fonts']=[dict(x.attrib) for x in xml.xpath('//a:fontScheme//*[self::a:latin or self::a:ea or self::a:cs or self::a:font]',namespaces=NS)]
                entry['colors']=[{'name':etree.QName(x.getparent()).localname,**dict(x.attrib)} for x in xml.xpath('//a:clrScheme/*/*',namespaces=NS)]
            if key=='comments': entry['text']='\n'.join(xml.xpath('//*[local-name()="text"]/text()'))
            components[key].append(entry)
    properties={n:{etree.QName(x).localname:x.text for x in _xml(parts[n])} for n in ['docProps/core.xml','docProps/app.xml'] if n in parts}
    return {'source_sha256': hashlib.sha256(Path(path).read_bytes()).hexdigest(), 'slides':slides,'components':components,'properties':properties,
            'slide_size': dict(_xml(parts['ppt/presentation.xml']).find('p:sldSz',NS).attrib),
            'parts': {n:hashlib.sha256(b).hexdigest() for n,b in parts.items()},
            'preserve_first':True, 'unsupported_features': [n for n in parts if any(k in n for k in ('diagrams/','comments','vbaProject','customXml/','embeddings/'))],
            'limitations':['Exact structured selectors only; natural-language semantic selection is not implemented.', 'ZIP container byte identity differs; unchanged member payload bytes are preserved.']}


def select_objects(model, selector):
    allowed={'object_id','slide_index','shape_id','name','text','placeholder_idx'}
    if not selector or set(selector)-allowed: fail('OBJECT_SELECTION_INVALID', '须指定支持的精确对象选择器。')
    found=[]
    for slide in model['slides']:
        if 'slide_index' in selector and slide['slide_index'] != selector['slide_index']: continue
        for item in slide['objects']:
            if any(item.get(k) != selector[k] for k in ('object_id','shape_id','name','text') if k in selector): continue
            if 'placeholder_idx' in selector and (item['placeholder'] or {}).get('idx','0') != str(selector['placeholder_idx']): continue
            found.append(item)
    if len(found) != 1: fail('OBJECT_SELECTION_AMBIGUOUS' if found else 'OBJECT_NOT_FOUND', '对象选择必须唯一；未扩大修改范围。', [{'matches':len(found)}])
    return found[0]


def _set_text(body, value):
    if not isinstance(value,str): fail('REVISION_VALUE_INVALID','文本必须为字符串。')
    first=body.find('a:p',NS)
    template=first.find('a:r/a:rPr',NS) if first is not None else None
    ppr=first.find('a:pPr',NS) if first is not None else None
    for p in list(body.findall('a:p',NS)): body.remove(p)
    for line in value.split('\n'):
        p=etree.SubElement(body,'{'+NS['a']+'}p')
        if ppr is not None: p.append(copy.deepcopy(ppr))
        run=etree.SubElement(p,'{'+NS['a']+'}r')
        if template is not None: run.append(copy.deepcopy(template))
        text=etree.SubElement(run,'{'+NS['a']+'}t');text.text=line


def _write_xml(root):
    return etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)


def _text_property_child(prop, tag):
    """Respect CT_TextCharacterProperties child order when adding properties."""
    order={'ln':0,'noFill':1,'solidFill':1,'gradFill':1,'blipFill':1,'pattFill':1,'grpFill':1,
           'effectLst':2,'effectDag':2,'highlight':3,'uLnTx':4,'uLn':4,'uFillTx':5,'uFill':5,
           'latin':6,'ea':7,'cs':8,'sym':9,'hlinkClick':10,'hlinkMouseOver':11,'rtl':12,'extLst':13}
    existing=prop.find('a:'+tag,NS)
    if existing is not None: return existing
    child=etree.Element('{'+NS['a']+'}'+tag)
    index=next((i for i,c in enumerate(prop) if order.get(etree.QName(c).localname,99)>order[tag]),len(prop))
    prop.insert(index,child);return child


def apply_patch(source_path, output_path, request):
    """Atomic save-as; audit all changed and untouched payload hashes."""
    source_path=Path(source_path); output_path=Path(output_path)
    if source_path.resolve() == output_path.resolve() or output_path.exists(): fail('REVISION_OUTPUT_EXISTS','须另存为不存在的路径，禁止覆盖。')
    before=inspect_deck(source_path)
    if request.get('expected_sha256') != before['source_sha256']: fail('REVISION_SOURCE_CHANGED','源文件 SHA-256 不匹配。')
    if set(request)-{'expected_sha256','patches','locked_object_ids'} or not isinstance(request.get('patches'),list) or not request['patches'] or len(request['patches']) > 200: fail('REVISION_PATCH_INVALID','补丁须为 1–200 条显式操作。')
    operations={'text','format','geometry','table_cell','chart_data','image_replace','alt_text','notes'}
    if any(not isinstance(p,dict) or p.get('operation') not in operations for p in request['patches']):
        fail('REVISION_OPERATION_UNSUPPORTED','修改操作未实现，未执行降级。')
    from .schemas import SchemaRegistry
    SchemaRegistry().validate('existing-deck-patch.schema.json',request)
    infos,parts=_package(source_path); original=dict(parts); roots={}; history=[]; locked=set(request.get('locked_object_ids',[]))
    if any(n.startswith('_xmlsignatures/') for n in parts): fail('SIGNED_DECK_UNSUPPORTED','签名演示文稿不能在未处理签名策略时修改。')
    for patch in request['patches']:
        if set(patch)-{'object_id','selector','operation','value'} or ('object_id' in patch) == ('selector' in patch): fail('REVISION_PATCH_INVALID','补丁须提供一种对象选择方式。')
        item=select_objects(before, {'object_id':patch['object_id']} if 'object_id' in patch else patch['selector'])
        if item['object_id'] in locked: fail('REVISION_OBJECT_LOCKED','对象已锁定，未执行修改。')
        part=item['part']; root=roots.setdefault(part,_xml(parts[part]))
        node,identity=next((n,i) for n,i in _objects(root) if int(i.get('id'))==item['shape_id'])
        operation=patch.get('operation'); value=patch.get('value')
        if operation == 'text':
            body=node.find('p:txBody',NS)
            if body is None: fail('OBJECT_TYPE_UNSUPPORTED','目标不是文本框。')
            _set_text(body,value)
        elif operation == 'alt_text':
            if not isinstance(value,str): fail('REVISION_VALUE_INVALID','替代文字必须为字符串。')
            identity.set('descr',value)
        elif operation == 'geometry':
            if not isinstance(value,list) or len(value)!=4 or any(isinstance(x,bool) or not isinstance(x,(int,float)) or not math.isfinite(x) for x in value) or min(value[:2])<0 or min(value[2:])<=0: fail('REVISION_GEOMETRY_INVALID','几何须为有限非负位置、正尺寸英寸值。')
            if node.getparent().tag != '{'+NS['p']+'}spTree': fail('GROUP_GEOMETRY_UNSUPPORTED','组合子对象坐标变换尚未实现。')
            xfrm=node.xpath('./p:spPr/a:xfrm|./p:xfrm',namespaces=NS)
            if not xfrm: fail('OBJECT_GEOMETRY_UNSUPPORTED','目标没有本地几何；不覆盖继承属性。')
            off=xfrm[0].find('a:off',NS);ext=xfrm[0].find('a:ext',NS)
            for name,v in zip(('x','y'),value[:2]): off.set(name,str(round(v*914400)))
            for name,v in zip(('cx','cy'),value[2:]): ext.set(name,str(round(v*914400)))
        elif operation == 'format':
            allowed={'font','font_size','bold','italic','underline','color','alignment'}
            if not isinstance(value,dict) or not value or set(value)-allowed: fail('REVISION_FORMAT_UNSUPPORTED','格式字段未实现。')
            body=node.find('p:txBody',NS)
            if body is None: fail('OBJECT_TYPE_UNSUPPORTED','目标没有文本格式。')
            if 'color' in value and not re.fullmatch('[0-9A-Fa-f]{6}',str(value['color'])): fail('REVISION_FORMAT_INVALID','颜色须为六位 RGB。')
            if 'font_size' in value and (isinstance(value['font_size'],bool) or not isinstance(value['font_size'],(int,float)) or not math.isfinite(value['font_size']) or not 1<=value['font_size']<=400): fail('REVISION_FORMAT_INVALID','字号须在 1–400 pt。')
            if any(not isinstance(value[k],bool) for k in ('bold','italic','underline') if k in value): fail('REVISION_FORMAT_INVALID','布尔格式须为布尔值。')
            if 'alignment' in value and value['alignment'] not in {'left','center','right','justify'}: fail('REVISION_FORMAT_INVALID','段落对齐无效。')
            for p in body.findall('a:p',NS):
                if 'alignment' in value:
                    ppr=p.find('a:pPr',NS)
                    if ppr is None: ppr=etree.Element('{'+NS['a']+'}pPr');p.insert(0,ppr)
                    ppr.set('algn',{'left':'l','center':'ctr','right':'r','justify':'just'}[value['alignment']])
                for run in p.findall('a:r',NS):
                    prop=run.find('a:rPr',NS)
                    if prop is None: prop=etree.Element('{'+NS['a']+'}rPr');run.insert(0,prop)
                    for k,a in [('bold','b'),('italic','i')]:
                        if k in value: prop.set(a,'1' if value[k] else '0')
                    if 'underline' in value: prop.set('u','sng' if value['underline'] else 'none')
                    if 'font_size' in value: prop.set('sz',str(round(value['font_size']*100)))
                    if 'font' in value:
                        if not isinstance(value['font'],str) or not value['font']: fail('REVISION_FORMAT_INVALID','字体名称无效。')
                        for tag in ('latin','ea','cs'):
                            el=_text_property_child(prop,tag)
                            el.set('typeface',value['font'])
                    if 'color' in value:
                        for el in list(prop):
                            if etree.QName(el).localname in {'solidFill','gradFill','noFill','pattFill','blipFill','grpFill'}: prop.remove(el)
                        fill=_text_property_child(prop,'solidFill');etree.SubElement(fill,'{'+NS['a']+'}srgbClr',val=value['color'].upper())
        elif operation == 'table_cell':
            if not isinstance(value,dict) or set(value)!={'row','column','text'} or any(isinstance(value[k],bool) or not isinstance(value[k],int) or value[k]<0 for k in ('row','column')): fail('REVISION_TABLE_INVALID','单元格补丁需要非负行列和文本。')
            rows=node.xpath('.//a:tbl/a:tr',namespaces=NS)
            if value['row']>=len(rows): fail('REVISION_TABLE_INVALID','表格行越界。')
            cells=rows[value['row']].findall('a:tc',NS)
            if value['column']>=len(cells): fail('REVISION_TABLE_INVALID','表格列越界。')
            cell=cells[value['column']]
            if cell.get('hMerge')=='1' or cell.get('vMerge')=='1': fail('REVISION_TABLE_MERGED','合并从属单元格不能独立修改。')
            _set_text(cell.find('a:txBody',NS),value['text'])
        elif operation == 'chart_data':
            if 'chart_part' not in item: fail('OBJECT_TYPE_UNSUPPORTED','目标不是原生图表。')
            chart_part=item['chart_part']; incoming=sum(v['target']==chart_part for n in parts if n.endswith('.xml') for v in _rels(parts,n).values() if not v['external'])
            if incoming!=1: fail('SHARED_CHART_UNSUPPORTED','共享图表部件不能局部替换。')
            # Load current edited package so multiple chart patches cannot reset earlier ones.
            current=dict(parts)
            for n,r in roots.items(): current[n]=_write_xml(r)
            buffer=io.BytesIO()
            with zipfile.ZipFile(buffer,'w') as z:
                for n,b in current.items(): z.writestr(n,b)
            prs=Presentation(io.BytesIO(buffer.getvalue())); slide=next(s for s in prs.slides if s.slide_id==int(item['object_id'].split(':s')[1].split(':')[0]))
            def recursive(shapes):
                for shape in shapes:
                    yield shape
                    if hasattr(shape,'shapes'): yield from recursive(shape.shapes)
            shape=next(s for s in recursive(slide.shapes) if s.shape_id==item['shape_id'])
            allowed_parts={chart_part}; related=_rels(parts,chart_part)
            for v in related.values():
                if v['external']: fail('LINKED_CHART_UNSUPPORTED','外部链接图表需要单独策略。')
                if v['type'].endswith('/package'): allowed_parts.add(v['target'])
            if not any(n.endswith('.xlsx') for n in allowed_parts): fail('CHART_WORKBOOK_UNSUPPORTED','图表缺少已存在的内嵌工作簿。')
            shape.chart.replace_data(chart_data(value)); saved=io.BytesIO();prs.save(saved)
            with zipfile.ZipFile(saved) as z:
                for n in allowed_parts: parts[n]=z.read(n)
        elif operation == 'image_replace':
            blips=node.xpath('./p:blipFill/a:blip',namespaces=NS)
            if not blips or not isinstance(value,dict) or set(value)!={'path'}: fail('OBJECT_TYPE_UNSUPPORTED','图片替换需要图片对象和文件路径。')
            rel=_rels(parts,part).get(blips[0].get('{'+NS['r']+'}embed'))
            if not rel or rel['external']: fail('LINKED_IMAGE_UNSUPPORTED','链接图片不能隐式改为嵌入图片。')
            target=rel['target']; refs=sum(v['target']==target for n in parts if n.endswith('.xml') for v in _rels(parts,n).values() if not v['external'])
            occurrences=len(root.xpath('//a:blip[@r:embed=$rid]',namespaces=NS,rid=blips[0].get('{'+NS['r']+'}embed')))
            if refs!=1 or occurrences>1: fail('SHARED_IMAGE_UNSUPPORTED','共享图片替换会扩大范围，已阻断。')
            from PIL import Image
            replacement=local_asset(value['path']);blob=replacement.read_bytes()
            if len(blob)>32*1024*1024: fail('IMAGE_SIZE_LIMIT','替换图片超过 32 MiB。')
            with Image.open(io.BytesIO(blob)) as image:
                ext={'.png':'PNG','.jpeg':'JPEG','.jpg':'JPEG'}.get(Path(target).suffix.lower())
                if ext is None or image.format!=ext: fail('IMAGE_FORMAT_CHANGED','须保持图片原部件格式；未静默转换。')
                if image.width*image.height>40_000_000: fail('IMAGE_SIZE_LIMIT','替换图片超过 4000 万像素。')
                image.verify()
            parts[target]=blob
        elif operation == 'notes':
            notes=next((v['target'] for v in _rels(parts,part).values() if v['type'].endswith('/notesSlide') and not v['external']),None)
            if not notes: fail('NOTES_CREATION_UNSUPPORTED','没有既有备注部件，未隐式创建。')
            note=roots.setdefault(notes,_xml(parts[notes])); candidates=note.xpath('//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]/p:txBody',namespaces=NS)
            if len(candidates)!=1: fail('NOTES_BODY_UNSUPPORTED','备注正文占位符不唯一。')
            _set_text(candidates[0],value)
        else: fail('REVISION_OPERATION_UNSUPPORTED','修改操作未实现，未执行降级。')
        effects = ['Text replacement retains the first paragraph/run style; original run segmentation is intentionally replaced.'] if operation in {'text','table_cell','notes'} else []
        history.append({'object_id':item['object_id'],'operation':operation,'before':item,'requested_value':value,'effects':effects})
    for part,root in roots.items():
        # Serialization only for XML whose semantic tree actually changed.
        if etree.tostring(root)!=etree.tostring(_xml(original[part])): parts[part]=_write_xml(root)
    changed=[n for n in parts if parts[n]!=original[n]]
    if not changed: fail('REVISION_NO_CHANGE','补丁没有产生修改。')
    output_path.parent.mkdir(parents=True,exist_ok=True); temporary=output_path.with_name(output_path.name+'.pending')
    if temporary.exists(): fail('REVISION_OUTPUT_EXISTS','暂存输出已存在。')
    try:
        with zipfile.ZipFile(temporary,'x') as archive:
            for info in infos: archive.writestr(info,parts[info.filename])
        after=inspect_deck(temporary)
        # Assert untouched shape trees within modified slide parts remain identical.
        targeted={h['object_id'] for h in history}
        for sid,part in _slides(original):
            if part not in changed: continue
            original_nodes={int(i.get('id')):n for n,i in _objects(_xml(original[part]))}
            old={i:etree.tostring(n) for i,n in original_nodes.items()}
            new={int(i.get('id')):etree.tostring(n) for n,i in _objects(_xml(parts[part]))}
            for shape_id,blob in old.items():
                descendants={f'pptx:s{sid}:shape{i.get("id")}' for i in original_nodes[shape_id].xpath('.//p:cNvPr',namespaces=NS)}
                if f'pptx:s{sid}:shape{shape_id}' not in targeted and not descendants.intersection(targeted) and blob!=new.get(shape_id): fail('PRESERVATION_OBJECT_CHANGED','非目标对象发生变化。')
        # Hard-link atomically refuses an output created by another writer.
        os.link(temporary,output_path); temporary.unlink()
    except Exception:
        if temporary.exists(): temporary.unlink()
        raise
    return {'source_sha256':before['source_sha256'],'output_sha256':after['source_sha256'],'changed_parts':changed,
            'preserved_parts':[n for n in parts if n not in changed], 'part_hashes':after['parts'], 'history':history,
            'after':after,'preservation':{'unchanged_member_payloads':'byte_identical','untargeted_shapes':'XML_identical',
            'system_acceptance':'not_passed','qa':'not_executed','limitations':before['limitations']}}
