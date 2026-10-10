"""IN-007, TPL-001..009/014/016: native PPTX template inspection and instantiation."""
import hashlib
import io
import re
from pathlib import Path
from xml.etree import ElementTree as ET

from .errors import TaskError


def parse_template(raw, *, source=None, license=None):
    from .input_formats import NS, package, xml, relationships, part_inventory
    layouts = []; masters = []; themes = []; issues = []
    with package(raw, 'ppt/presentation.xml') as archive:
        names = archive.namelist(); presentation = xml(archive.read('ppt/presentation.xml')); size = presentation.find('p:sldSz', NS)
        if size is None:
            raise TaskError('TEMPLATE_SLIDE_SIZE_MISSING', '模板缺少幻灯片尺寸。', 'TemplateRegistry', ['IN-007','TPL-014'])
        dimensions = {'width_emu': int(size.get('cx')), 'height_emu': int(size.get('cy')), 'type': size.get('type')}
        def bounds(shape):
            xfrm = shape.find('.//a:xfrm', NS)
            if xfrm is None:return None
            off = xfrm.find('a:off', NS); ext = xfrm.find('a:ext', NS)
            return {'x_emu': int(off.get('x')), 'y_emu': int(off.get('y')), 'width_emu': int(ext.get('cx')), 'height_emu': int(ext.get('cy'))} if off is not None and ext is not None else None
        def shapes(node):
            result = []
            for shape in node.findall('p:cSld/p:spTree/*', NS):
                if shape.tag.rsplit('}',1)[-1] in {'nvGrpSpPr','grpSpPr'}:continue
                prop = shape.find('.//p:cNvPr', NS); ph = shape.find('.//p:ph', NS)
                role = ph.get('type','obj') if ph is not None else None
                result.append({'shape_id': prop.get('id') if prop is not None else None, 'name': prop.get('name') if prop is not None else None,
                               'kind': shape.tag.rsplit('}',1)[-1], 'placeholder': ph.attrib if ph is not None else None,
                               'placeholder_index': int(ph.get('idx',0)) if ph is not None else None, 'role': role,
                               'bounds': bounds(shape), 'text': '\n'.join(''.join(t.text or '' for t in p.findall('.//a:t',NS)) for p in shape.findall('.//a:p',NS)),
                               'raw_xml': ET.tostring(shape, encoding='unicode')})
            return result
        for part in names:
            if re.fullmatch(r'ppt/slideMasters/slideMaster\d+\.xml', part):
                node = xml(archive.read(part)); background = node.find('p:cSld/p:bg',NS); textstyles=node.find('p:txStyles',NS); cmap=node.find('p:clrMap',NS)
                masters.append({'part':part,'shapes':shapes(node), 'relationships':relationships(archive,part),
                                'background_xml':ET.tostring(background,encoding='unicode') if background is not None else None,
                                'text_styles_xml':ET.tostring(textstyles,encoding='unicode') if textstyles is not None else None,
                                'color_map':cmap.attrib if cmap is not None else None,
                                'transition_xml':[ET.tostring(x,encoding='unicode') for x in node.findall('p:transition',NS)],
                                'timing_xml':[ET.tostring(x,encoding='unicode') for x in node.findall('p:timing',NS)],
                                'raw_xml':ET.tostring(node,encoding='unicode')})
            elif re.fullmatch(r'ppt/theme/theme\d+\.xml',part):
                node=xml(archive.read(part)); colors=[]; fonts={}
                for scheme in node.findall('a:themeElements/a:clrScheme',NS):
                    for color in scheme:
                        leaf=next(iter(color),None)
                        colors.append({'token':color.tag.rsplit('}',1)[-1], 'kind':leaf.tag.rsplit('}',1)[-1] if leaf is not None else None,
                                       'value':leaf.get('val') if leaf is not None else None,'last_color':leaf.get('lastClr') if leaf is not None else None,
                                       'raw_xml':ET.tostring(color,encoding='unicode')})
                for role in ['majorFont','minorFont']:
                    font=node.find('a:themeElements/a:fontScheme/a:'+role,NS)
                    fonts[role]=[{'kind':x.tag.rsplit('}',1)[-1],**x.attrib} for x in font] if font is not None else []
                themes.append({'part':part,'name':node.get('name'),'colors':colors,'fonts':fonts,'raw_xml':ET.tostring(node,encoding='unicode')})
        master_by_part={m['part']:m for m in masters}
        for part in names:
            if not re.fullmatch(r'ppt/slideLayouts/slideLayout\d+\.xml',part):continue
            node=xml(archive.read(part)); rels=relationships(archive,part)
            master_rel=next((r for r in rels if (r['type'] or '').endswith('/slideMaster') and not r['external']),None)
            master=master_by_part.get(master_rel['target_part']) if master_rel else None
            local=shapes(node); slots=[]
            for item in local:
                if item['placeholder'] is None:continue
                inherited=None
                if master:
                    inherited=next((s for s in master['shapes'] if s['placeholder'] is not None and s['placeholder_index']==item['placeholder_index'] and s['role']==item['role']),None)
                    if inherited is None:inherited=next((s for s in master['shapes'] if s['placeholder'] is not None and s['role']==item['role']),None)
                geometry=item['bounds'] or (inherited['bounds'] if inherited else None)
                role=item['role']; types={'pic':['image'],'chart':['chart'],'tbl':['table'],'media':['media'],'obj':['text','image','table','chart']}.get(role,['text'])
                slots.append({'slot_id':f'{part}#placeholder={item["placeholder_index"]}', 'placeholder_index':item['placeholder_index'],
                              'name':item['name'],'role':role,'content_types':types,'bounds':geometry,
                              'bounds_origin':'layout' if item['bounds'] else 'master' if inherited and inherited['bounds'] else 'unresolved',
                              'required':None,'repeatable':False,'capacity':{'status':'unmeasured','characters':None,'lines':None,'minimum_font_pt':None},
                              'overflow_policy':'fail_until_capacity_measured','style_xml':item['raw_xml'],
                              'inherited_shape_ref':f'{master["part"]}#shape={inherited["shape_id"]}' if master and inherited else None,
                              'brand_lock':'unspecified'})
                if geometry is None:issues.append({'code':'TEMPLATE_SLOT_BOUNDS_UNRESOLVED','part':part,'slot_id':slots[-1]['slot_id'],'message':'占位符缺少可解析布局/母版几何，不能假定容量。'})
            cslide=node.find('p:cSld',NS)
            layouts.append({'part':part,'name':cslide.get('name') if cslide is not None else None,'type':node.get('type'),'preserve':node.get('preserve'),
                            'show_master_shapes':node.get('showMasterSp','1') not in {'0','false'},'master_part':master['part'] if master else None,
                            'slots':slots,'shapes':local,'relationships':rels,'raw_xml':ET.tostring(node,encoding='unicode')})
        issues.append({'code':'TEMPLATE_CAPACITY_UNMEASURED','message':'原生母版/布局/字体/颜色/占位符已抽取；容量、品牌授权、安全区和非OOXML设计tokens不能凭XML猜测。'})
        return {'record_type':'native_template_pack','status':'partial','package_sha256':hashlib.sha256(raw).hexdigest(),
                'source':source,'license':license or 'unknown','version':'sha256:'+hashlib.sha256(raw).hexdigest(),
                'slide_size':dimensions,'masters':masters,'layouts':layouts,'themes':themes,
                'assets':[p for p in part_inventory(archive) if p['part'].startswith('ppt/media/')],
                'coverage':{'master_layout_theme':'parsed','placeholder_geometry':'parsed_with_explicit_unresolved','slot_capacity':'unmeasured',
                            'brand_lock':'user_policy_required','preview':'not_rendered','editable_instantiation':'available_for_text_placeholders'},
                'issues':issues}


def query_layouts(pack, *, content_type='text', minimum_slots=1, slide_size=None):
    """Explain deterministic structural candidates; no claim that unmeasured content fits."""
    if slide_size and any(slide_size.get(k)!=pack['slide_size'][k] for k in ['width_emu','height_emu']):
        raise TaskError('TEMPLATE_SIZE_MISMATCH','模板尺寸与输出不符，须显式转换策略。','TemplateRegistry',['TPL-014'])
    candidates=[]
    for layout in pack['layouts']:
        matching=[s for s in layout['slots'] if content_type in s['content_types'] and s['bounds'] is not None]
        candidates.append({'layout_part':layout['part'],'eligible':len(matching)>=minimum_slots,
                           'matching_slot_ids':[s['slot_id'] for s in matching], 'score':len(matching),
                           'reasons':['仅按内容类型和已定位槽位筛选；内容容量尚未测量。'], 'capacity_status':'unmeasured'})
    return sorted(candidates,key=lambda c:(not c['eligible'],-c['score'],c['layout_part']))


def instantiate_template(raw, layout_part, slot_content):
    """Append an editable native slide. Keep existing template slides and master/layout/theme.

    Values map integer placeholder indices to strings. This operation intentionally
    accepts text slots only; callers must measure capacity/render/QA before delivery.
    """
    from pptx import Presentation
    from .input_formats import package, xml
    pack=parse_template(raw); match=next((x for x in pack['layouts'] if x['part']==layout_part),None)
    if match is None:raise TaskError('TEMPLATE_LAYOUT_UNKNOWN','模板布局不存在。','TemplateRegistry',['TPL-009'])
    contents={}
    for key,value in slot_content.items():
        try:index=int(key)
        except (TypeError,ValueError):raise TaskError('TEMPLATE_SLOT_INVALID','占位符索引必须为整数。','TemplateRegistry',['TPL-004']) from None
        if str(index)!=str(key) or not isinstance(value,str):raise TaskError('TEMPLATE_SLOT_INVALID','槽位索引/文本不合法。','TemplateRegistry',['TPL-004'])
        slot=next((s for s in match['slots'] if s['placeholder_index']==index),None)
        if slot is None or 'text' not in slot['content_types']:raise TaskError('TEMPLATE_SLOT_UNSUPPORTED','占位符不存在或不是文本槽位。','TemplateRegistry',['TPL-004','TPL-009'])
        contents[index]=value
    presentation=Presentation(io.BytesIO(raw)); selected=None
    for master in presentation.slide_masters:
        for layout in master.slide_layouts:
            if str(layout.part.partname).lstrip('/')==layout_part:selected=layout
    if selected is None:raise TaskError('TEMPLATE_LAYOUT_UNAVAILABLE','原生布局不可实例化。','TemplateRegistry',['TPL-009'])
    slide=presentation.slides.add_slide(selected)
    for index,text in contents.items():
        try:placeholder=slide.placeholders[index]
        except KeyError:raise TaskError('TEMPLATE_PLACEHOLDER_UNAVAILABLE','该占位符未继承到实际新建页面。','TemplateRegistry',['TPL-009']) from None
        if not placeholder.has_text_frame:raise TaskError('TEMPLATE_SLOT_UNSUPPORTED','原生占位符不是文本框。','TemplateRegistry',['TPL-009'])
        placeholder.text=text
    output=io.BytesIO(); presentation.save(output); result=output.getvalue()
    with package(raw,'ppt/presentation.xml') as before, package(result,'ppt/presentation.xml') as after:
        protected=[p for p in before.namelist() if p.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/')) and not p.endswith('.rels')]
        mismatches=[p for p in protected if p not in after.namelist() or ET.tostring(xml(before.read(p)))!=ET.tostring(xml(after.read(p)))]
    if mismatches:raise TaskError('TEMPLATE_INHERITANCE_CHANGED','实例化改变了母版/布局/主题结构，拒绝结果。','TemplateRegistry',['TPL-009'],[{'parts':mismatches}])
    return result, {'status':'partial','layout_part':layout_part,'appended_slide_index':len(presentation.slides)-1,
                    'native_text_placeholders':sorted(contents),'master_layout_theme_preserved':True,
                    'sha256':hashlib.sha256(result).hexdigest(),'issues':[{'code':'TEMPLATE_RENDER_QA_PENDING','message':'已实际新建原生可编辑页面；渲染、容量和视觉QA待执行。'}]}
