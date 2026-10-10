"""Shared operation dispatcher for CLI, HTTP, durable jobs and MCP (SYS-019–022)."""
import json
from pathlib import Path
from uuid import UUID

from .artifacts import sha256
from .errors import TaskError
from .schemas import SchemaRegistry
from .evidence import EvidenceStore


def run_folder(run_id):
    try:
        if str(UUID(run_id)) != run_id:
            raise ValueError()
    except (ValueError, TypeError, AttributeError):
        raise TaskError('RUN_ID_INVALID', '运行ID必须为标准UUID。', 'Operations', ['SYS-017']) from None
    folder = Path('/runtime/runs') / run_id
    if not folder.is_dir():
        raise TaskError('RUN_NOT_FOUND', '运行不存在。', 'Operations', ['SYS-017'], status=404)
    return folder


def artifacts(run_id, trace):
    folder = run_folder(run_id)
    path = folder/'manifest.json'
    if not path.is_file():
        raise TaskError('RUN_MANIFEST_MISSING', '运行没有完成产物清单。', 'Operations', ['OBS-006'])
    manifest = json.loads(path.read_text('utf-8'))
    verified = []
    for item in manifest['artifacts']:
        p = (folder/item['path']).resolve()
        if not p.is_relative_to(folder.resolve()) or not p.is_file() or sha256(p) != item['hash']:
            raise TaskError('ARTIFACT_CHANGED', '产物清单与文件不符。', 'Operations', ['OBS-006'])
        verified.append({**item, 'absolute_path': str(p)})
    return {'ok': True, 'trace_id': trace, 'run_id': run_id, 'artifacts': verified}


def dispatch(operation, document, trace, schemas=None, store=None):
    schemas = schemas or SchemaRegistry()
    store = store or EvidenceStore()
    if operation == 'validate':
        from .tasks import validate_task
        return validate_task(document, schemas, trace)
    if operation == 'inspect':
        from .sources import inspect_sources
        return inspect_sources(document, schemas, store, trace)
    if operation in {'create_deck','generate'}:
        from .generation import generate
        return generate(document, schemas, store, trace)
    if operation in {'revise_deck','revise'}:
        from .revision import revise
        return revise(document['run_id'], document['request'], schemas, trace)
    if operation == 'resume':
        from .generation import resume
        return resume(document['run_id'], schemas, store, trace,document.get('task_patch'))
    if operation == 'recheck':
        from .revision import recheck
        return recheck(document['run_id'], schemas, trace)
    if operation == 'capabilities':
        from .capabilities import CapabilityRegistry
        result = CapabilityRegistry(schemas).list(trace)
        try:
            from .native_objects import capabilities
            result['native_objects'] = capabilities()
        except ImportError:
            pass
        return result
    if operation == 'artifacts':
        return artifacts(document['run_id'], trace)
    if operation == 'compare':
        from .observability import compare_runs
        return compare_runs(document['left_run_id'],document['right_run_id'],trace)
    if operation == 'inventory_assets':
        from .assets import inventory
        return {'ok':True,'trace_id':trace,'inventory':inventory(document['directory'])}
    if operation == 'doctor':
        import shutil
        from .observability import version_manifest
        return {'ok':True,'trace_id':trace,'versions':version_manifest(schemas),
                'tools':{name:bool(shutil.which(name)) for name in ['libreoffice','pdftotext','pdftoppm']},
                'storage':{'root':'/runtime','writable':__import__('os').access('/runtime',__import__('os').W_OK)},
                'system_acceptance':'not_passed'}
    if operation == 'render':
        from .rendering import render
        from .artifacts import RunArtifacts
        artifacts(document['run_id'],trace)
        old=run_folder(document['run_id']);output=RunArtifacts()
        if not (old/'deck-spec.json').is_file():
            raise TaskError('RENDER_SPEC_MISSING','该运行缺少渲染用deck-spec。','Operations',['SYS-014'])
        deck=json.loads((old/'deck-spec.json').read_text('utf-8'))
        import shutil
        shutil.copyfile(old/'deck.pptx',output.path/'deck.pptx')
        report=render(output.path/'deck.pptx',deck,output.path)
        output.json('render-report.json',report)
        output.json('manifest.json',{'run_id':output.run_id,'artifacts':output.manifest()})
        return {'ok':True,'trace_id':trace,'run_id':output.run_id,'render':report,'source_run_id':document['run_id']}
    if operation == 'compose_native':
        from pptx import Presentation
        from pptx.util import Inches
        from .native_objects import add_object
        from .existing_deck import inspect_deck
        from .artifacts import RunArtifacts
        slides=document.get('slides')
        if not isinstance(slides,list) or not 1<=len(slides)<=30:
            raise TaskError('NATIVE_DECK_INVALID','原生执行需1–30页。','Operations',['PPT-001'])
        for slide in slides:
            if not isinstance(slide,dict) or set(slide)-{'objects','notes'} or not isinstance(slide.get('objects'),list):
                raise TaskError('NATIVE_DECK_INVALID','页面需objects，可提供notes。','Operations',['PPT-001'])
            for spec in slide['objects']:schemas.validate('native-object-spec.schema.json',spec)
        output=RunArtifacts();prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
        for content in slides:
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            for spec in content['objects']:add_object(slide,spec)
            if 'notes' in content:slide.notes_slide.notes_text_frame.text=content['notes']
        prs.save(output.path/'deck.pptx');model=inspect_deck(output.path/'deck.pptx')
        result={'ok':True,'trace_id':trace,'run_id':output.run_id,'pptx':str(output.path/'deck.pptx'),
                'object_map':model,'scope':'caller supplied native object execution; not AI generation or system acceptance','qa':'not_executed'}
        output.json('result.json',result);output.json('manifest.json',{'run_id':output.run_id,'artifacts':output.manifest()})
        return result
    if operation in {'inspect_deck','edit_deck','parse_template','instantiate_template'}:
        from .schemas import ROOT
        from .artifacts import RunArtifacts
        path=Path(document.get('path',''))
        path=(ROOT/path).resolve() if not path.is_absolute() else path.resolve()
        if not path.is_file() or 'secrets' in path.parts or not (path.is_relative_to(ROOT.resolve()) or path.is_relative_to(Path('/runtime'))):
            raise TaskError('SOURCE_PATH_OUTSIDE_ROOT','文件须在工作区或运行目录。','Operations',['SYS-003'])
        if operation=='parse_template':
            from .templates import parse_template
            result=parse_template(path.read_bytes())
            schemas.validate('template-native-pack.schema.json',result)
        elif operation=='instantiate_template':
            from .templates import instantiate_template
            output=RunArtifacts()
            blob,result=instantiate_template(path.read_bytes(),document['layout_part'],document['slot_content'])
            (output.path/'deck.pptx').write_bytes(blob)
            result.update(run_id=output.run_id,pptx=str(output.path/'deck.pptx'))
            output.json('template-report.json',result)
            output.json('manifest.json',{'run_id':output.run_id,'artifacts':output.manifest()})
        elif operation=='inspect_deck':
            from .existing_deck import inspect_deck
            result=inspect_deck(path)
        else:
            from .existing_deck import apply_patch
            output=RunArtifacts()
            result=apply_patch(path,output.path/'deck.pptx',document['request'])
            result.update(run_id=output.run_id,pptx=str(output.path/'deck.pptx'),artifact_root=str(output.path))
            output.json('revision-report.json',result)
            output.json('manifest.json',{'run_id':output.run_id,'artifacts':output.manifest()})
        return {'ok':True,'trace_id':trace,'operation':operation,'result':result}
    if operation == 'debug':
        node = document.get('node_id')
        payload = document.get('input', {})
        if node == 'DEC-001':
            from .routing import run_route
            return run_route(payload, schemas, trace)
        if node == 'DEC-002':
            from .context import run_context
            return run_context(payload, schemas, trace)
        if node == 'DEC-003':
            from .source_roles import run_roles
            return run_roles(payload, schemas, trace)
        if node in {'DEC-004','DEC-005'}:
            from .facts import run_node
            return run_node(payload, schemas, node, trace)
        if isinstance(node,str) and node in {f'DEC-{i:03d}' for i in range(6,22)}:
            from .narrative import run_node
            from .artifacts import RunArtifacts
            output=RunArtifacts()
            result=run_node(node,payload,schemas,output)
            output.json('result.json',result)
            output.json('manifest.json',{'run_id':output.run_id,'artifacts':output.manifest()})
            return {'ok':True,'trace_id':trace,'run_id':output.run_id,'artifact_root':str(output.path),**result}
        raise TaskError('DEBUG_NODE_UNSUPPORTED', '该调试节点尚未接入；未伪造结果。', 'DebugRunner', ['SYS-022'])
    raise TaskError('OPERATION_UNSUPPORTED', '未登记该操作。', 'Operations', ['SYS-019','SYS-020','SYS-021'])
