"""QA-001..020 deterministic measurements; partial semantic/Office coverage explicit."""
import hashlib
import math
import posixpath
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pptx.enum.shapes import MSO_SHAPE_TYPE
from .artifacts import entity

EMU=914400
A='{http://schemas.openxmlformats.org/drawingml/2006/main}'
P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
REL='{http://schemas.openxmlformats.org/package/2006/relationships}'
CT='{http://schemas.openxmlformats.org/package/2006/content-types}'

def issue(code,slide_id,detail,requirement,severity='P0',object_id=None,measurement=None,threshold=None,action=None):
 return {'issue_id':entity('issue'),'code':code,'slide_id':slide_id,'object_id':object_id,'severity':severity,
         'detail':detail,'requirement_id':requirement,'measurement':measurement or {},'threshold':threshold or {},
         'action':action or {'node':'QA','mode':'fail','reason':'No safe automatic repair registered'}}

def contrast(left,right):
 def luminance(value):
  channels=[int(value[i:i+2],16)/255 for i in (0,2,4)]
  channels=[c/12.92 if c<=0.04045 else ((c+0.055)/1.055)**2.4 for c in channels]
  return sum(c*w for c,w in zip(channels,(0.2126,0.7152,0.0722)))
 a,b=sorted((luminance(left),luminance(right)))
 return (b+0.05)/(a+0.05)

def package_audit(path):
 issues=[];checks=[]
 with ZipFile(path) as z:
  names=z.namelist();existing=set(names)
  if len(existing)!=len(names):issues.append(issue('ZIP_DUPLICATE_PART',None,'Duplicate ZIP members','QA-001'))
  if '[Content_Types].xml' not in existing:issues.append(issue('CONTENT_TYPES_MISSING',None,'Missing content types','QA-001'));return issues,checks
  try:ct=ET.fromstring(z.read('[Content_Types].xml'))
  except ET.ParseError: return [issue('CONTENT_TYPES_INVALID',None,'Invalid content types XML','QA-001')],checks
  defaults={n.get('Extension') for n in ct.findall(CT+'Default')};overrides={n.get('PartName','').lstrip('/') for n in ct.findall(CT+'Override')}
  for n in overrides-existing:issues.append(issue('CONTENT_TYPE_PART_MISSING',None,n,'QA-001'))
  for name in names:
   if name.endswith('/') or name=='[Content_Types].xml':continue
   if name not in overrides and name.rsplit('.',1)[-1] not in defaults:issues.append(issue('CONTENT_TYPE_UNDECLARED',None,name,'QA-001'))
   if not name.endswith('.rels'):continue
   try:rels=ET.fromstring(z.read(name))
   except ET.ParseError:continue # primary QA records invalid XML
   owner='' if name=='_rels/.rels' else posixpath.dirname(posixpath.dirname(name))
   seen=set()
   for r in rels.findall(REL+'Relationship'):
    rid=r.get('Id');target=r.get('Target','')
    if rid in seen:issues.append(issue('RELATIONSHIP_DUPLICATE_ID',None,{'part':name,'id':rid},'QA-001'))
    seen.add(rid)
    if r.get('TargetMode')=='External':continue
    resolved=posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(owner,target)).split('#',1)[0]
    if resolved not in existing:issues.append(issue('RELATIONSHIP_TARGET_MISSING',None,{'part':name,'target':target,'resolved':resolved},'QA-001'))
  checks.append({'kind':'structure','check':'Content_Types declarations and all internal relationship targets/IDs','status':'executed'})
 return issues,checks

def _flatten(shapes):
 for s in shapes:
  yield s
  if s.shape_type==MSO_SHAPE_TYPE.GROUP:yield from _flatten(s.shapes)

def _rgb(color):
 try:return str(color.rgb) if color.rgb is not None else None
 except (AttributeError,TypeError):return None

def audit(deck,presentation,object_map):
 issues=[];checks=[];known=set();mapped={o['logical_object_id']:o for o in object_map['objects']}
 fonts=set();footers=[];images={}
 globalstyle=deck.get('style',{})
 def add(code,spec,obj,detail,req,severity='P0',measurement=None,threshold=None,node='layout'):
  issues.append(issue(code,spec['slide_id'],detail,req,severity,obj,measurement,threshold,{'node':node,'mode':'local_revision','reason':code}))
 for index,spec in enumerate(deck['slides']):
  if index>=len(presentation.slides):break
  slide=presentation.slides[index];style=spec.get('style',{});elements=spec['elements'];actual={s.name:s for s in _flatten(slide.shapes)}
  expected={e['object_id']:e for e in elements};known.update(expected)
  if not len(slide.shapes):add('SLIDE_EMPTY',spec,None,'No shapes','QA-003')
  if style.get('font')!=globalstyle.get('font'):add('DECK_FONT_INCONSISTENT',spec,None,'Page font differs from deck font','QA-015','P1',node='style')
  actual_background=None
  try:actual_background=_rgb(slide.background.fill.fore_color)
  except (AttributeError,TypeError):pass
  layout=spec.get('layout',{});selection=layout.get('selection_trace',{});chosen=selection.get('selected');chosen_ids=chosen if isinstance(chosen,list) else [chosen]
  selected_candidates=[c for c in selection.get('candidates',[]) if c.get('candidate_id') in chosen_ids]
  design_name=layout.get('selected_layout')
  explicit_variation=bool(design_name in {'cover','closing'} and selection.get('selected_layout')==design_name and len(selected_candidates)==1 and selected_candidates[0].get('layout')==design_name and selected_candidates[0].get('eligible') is True and actual_background==style.get('background'))
  if actual_background is not None and actual_background!=style.get('background'):add('SLIDE_BACKGROUND_CHANGED',spec,None,'Actual native page background differs from executed SlideSpec','QA-015','P1',{'actual':actual_background,'expected':style.get('background')},{'rule':'actual must match executed SlideSpec'},'style')
  if style.get('background')!=globalstyle.get('background') and not explicit_variation:add('DECK_BACKGROUND_INCONSISTENT',spec,None,'Page background differs without matching native/selected cover or closing design evidence','QA-015','P1',node='style')
  checks.append({'slide_id':spec['slide_id'],'kind':'global_consistency','check':'native/background/design_variation','status':'executed','explicit_variation':explicit_variation,'actual_background':actual_background,'spec_background':style.get('background'),'deck_background':globalstyle.get('background'),'selected_layout':design_name,'selected_candidate_ids':[c.get('candidate_id') for c in selected_candidates]})
  boxes=[];title=[];body=[];footer=[]
  for e in elements:
   oid=e['object_id'];shape=actual.get(oid)
   if shape is None:continue
   x,y,w,h=[float(v) for v in e['bounds']]
   if not all(math.isfinite(v) for v in (x,y,w,h)) or (e['type']!='connector' and (w<=0 or h<=0)) or (e['type']=='connector' and w==0 and h==0):add('OBJECT_ZERO_OR_INVALID_SIZE',spec,oid,{'bounds':e['bounds']},'QA-003')
   real=[v/EMU for v in (shape.left,shape.top,shape.width,shape.height)]
   if any(abs(a-b)>0.002 for a,b in zip(real,(x,y,w,h))):add('OBJECT_GEOMETRY_MISMATCH',spec,oid,'Native coordinates differ from spec','QA-004',measurement={'actual':real,'expected':e['bounds']},threshold={'tolerance_inches':0.002},node='geometry')
   if e['role'] not in {'background','decoration','connector'} and e['type']!='connector':boxes.append((e,(x,y,w,h)))
   if e['type']=='text':
    role=e['role'];text=''.join(e.get('text',[]))
    if not text.strip():add('TEXT_EMPTY',spec,oid,'Empty text object','QA-003','P1',node='content')
    size=e['font_size'];(title if role=='title' else footer if role=='footer' else body).append(size)
    minimum=12 if role=='footer' else 16
    if size<minimum:add('FONT_TOO_SMALL',spec,oid,'Font below visible-content policy','QA-006','P0',{'size_pt':size},{'minimum_pt':minimum},'style')
    foreground=e.get('format',{}).get('color',style.get('accent') if role=='title' else style.get('foreground','000000'))
    background=e.get('format',{}).get('fill',style.get('background','FFFFFF'))
    if foreground and background:
     ratio=contrast(foreground,background);minimum_ratio=3.0 if size>=24 or (size>=18 and e.get('format',{}).get('bold',role=='title')) else 4.5
     if ratio<minimum_ratio:add('TEXT_CONTRAST_LOW',spec,oid,'Text foreground/background readability ratio below policy','QA-009','P0',{'ratio':ratio,'foreground':foreground,'background':background},{'minimum_ratio':minimum_ratio},'style')
    actual_background=background
    try:
     actual_background=_rgb(shape.fill.fore_color) or _rgb(slide.background.fill.fore_color) or background
    except (AttributeError,TypeError):pass
    for p in shape.text_frame.paragraphs:
     actual_colors=[_rgb(p.font.color),*[_rgb(run.font.color) for run in p.runs]]
     for actual_color in set(c for c in actual_colors if c):
      actual_ratio=contrast(actual_color,actual_background)
      if actual_ratio<minimum_ratio:add('ACTUAL_TEXT_CONTRAST_LOW',spec,oid,'Native paragraph/run overrides have insufficient contrast','QA-009','P0',{'ratio':actual_ratio,'foreground':actual_color,'background':actual_background},{'minimum_ratio':minimum_ratio},'style')
     if p.font.name:fonts.add(p.font.name)
     for run in p.runs:
      if run.font.name:fonts.add(run.font.name)
      if run.font.size is not None and run.font.size.pt<minimum:add('ACTUAL_FONT_TOO_SMALL',spec,oid,'Run overrides font below policy','QA-018','P1',{'size_pt':run.font.size.pt},{'minimum_pt':minimum},'style')
    if role=='footer':
     footers.append({'slide_id':spec['slide_id'],'text':text})
     if str(index+1) not in text:add('PAGE_NUMBER_MISSING',spec,oid,'Footer does not contain the actual page number','QA-015','P1',node='content')
   elif e['type']=='image':
    if shape.shape_type!=MSO_SHAPE_TYPE.PICTURE:add('IMAGE_NOT_NATIVE',spec,oid,'Image spec is not a native picture','QA-016',node='adapter');continue
    c=shape._element.xpath('.//p:cNvPr');alt=c[0].get('descr','') if c else ''
    if not alt.strip():add('ALT_TEXT_MISSING',spec,oid,'Native picture alt text absent','QA-018','P1',node='accessibility')
    pixels=shape.image.size;dpi=min(pixels[0]/max(w,0.001),pixels[1]/max(h,0.001))
    if dpi<96:add('IMAGE_RESOLUTION_LOW',spec,oid,'Image below rendering policy resolution','QA-010','P1',{'pixels':list(pixels),'effective_dpi':dpi},{'minimum_dpi':96},'asset')
    crop=shape.crop_left+shape.crop_right+shape.crop_top+shape.crop_bottom
    if crop==0 and abs((pixels[0]/pixels[1])/(w/h)-1)>0.05:add('IMAGE_STRETCHED',spec,oid,'Uncropped image aspect differs from native bounds','QA-010','P1',{'source_ratio':pixels[0]/pixels[1],'display_ratio':w/h},{'relative_tolerance':0.05},'geometry')
    digest=hashlib.sha256(shape.image.blob).hexdigest()
    if e.get('data',{}).get('sha256') and digest!=e['data']['sha256']:add('IMAGE_PAYLOAD_CHANGED',spec,oid,'Embedded native image bytes differ from resolved source','QA-010',measurement={'actual_sha256':digest,'expected_sha256':e['data']['sha256']},node='asset')
    images.setdefault(digest,[]).append((spec['slide_id'],oid))
   elif e['type']=='shape':
    expected_text=e.get('data',{}).get('text','')
    if not shape.has_text_frame or shape.text!=expected_text:add('DIAGRAM_LABEL_CHANGED',spec,oid,{'actual':shape.text if shape.has_text_frame else None,'expected':expected_text},'QA-013',node='adapter')
    if shape.shape_type!=MSO_SHAPE_TYPE.AUTO_SHAPE:add('DIAGRAM_NODE_NOT_NATIVE',spec,oid,'Expected native editable auto-shape','QA-016',node='adapter')
   elif e['type']=='connector':
    if shape.shape_type!=MSO_SHAPE_TYPE.LINE:add('DIAGRAM_CONNECTOR_NOT_NATIVE',spec,oid,'Expected native editable connector','QA-016',node='adapter')
    data=e.get('data',{});start=data.get('start');end=data.get('end')
    if start and end:
     if any(v<0 for point in (start,end) for v in point) or any(point[0]>deck['width_inches'] or point[1]>deck['height_inches'] for point in (start,end)):add('DIAGRAM_ENDPOINT_OUT_OF_BOUNDS',spec,oid,{'start':start,'end':end},'QA-004',node='geometry')
    props=shape._element.xpath('.//p:cNvPr');actual_label=props[0].get('descr','') if props else ''
    if actual_label!=(data.get('label') or data.get('relation','')):add('DIAGRAM_CONNECTOR_LABEL_CHANGED',spec,oid,{'actual':actual_label,'expected':data.get('label') or data.get('relation','')},'QA-013',node='adapter')
   elif e['type']=='diagram':
    if shape.shape_type!=MSO_SHAPE_TYPE.GROUP:add('DIAGRAM_NOT_NATIVE_GROUP',spec,oid,'Diagram expected editable group','QA-016',node='adapter')
    elif not len(shape.shapes):add('DIAGRAM_EMPTY',spec,oid,'Native diagram contains no child shapes','QA-003',node='diagram')
    else:
     data=e.get('data',{});children=list(_flatten(shape.shapes));labels=[c.text for c in children if c.has_text_frame and c.text.strip()]
     expected_labels=[n.get('label',n.get('text','')) for n in data.get('nodes',[])]+[edge['label'] for edge in data.get('edges',[]) if edge.get('label')]
     if sorted(labels)!=sorted(expected_labels):add('DIAGRAM_LABEL_CHANGED',spec,oid,{'actual':labels,'expected':expected_labels},'QA-013',node='adapter')
     connectors=[c for c in children if c.shape_type==MSO_SHAPE_TYPE.LINE]
     if len(connectors)!=len(data.get('edges',[])):add('DIAGRAM_EDGE_COUNT_CHANGED',spec,oid,{'actual':len(connectors),'expected':len(data.get('edges',[]))},'QA-016',node='adapter')
   elif e['type'] not in {'table','chart','shape','connector','group'}:
    add('QA_OBJECT_TYPE_UNSUPPORTED',spec,oid,{'type':e['type'],'reason':'No editability rule registered; not silently ignored'},'QA-016',node='adapter')
   if shape.has_table:
    if not shape.table.first_row:add('TABLE_HEADER_FLAG_MISSING',spec,oid,'Native table does not mark first row as header','QA-018','P1',node='accessibility')
   # Never infer semantic reading order from geometry. Report a measurable native order violation.
  for i,(left,lb) in enumerate(boxes):
   for right,rb in boxes[i+1:]:
    if left.get('overlap_allowed') or right.get('overlap_allowed'):continue
    x=max(lb[0],rb[0]);y=max(lb[1],rb[1]);w=min(lb[0]+lb[2],rb[0]+rb[2])-x;h=min(lb[1]+lb[3],rb[1]+rb[3])-y
    if w>0.02 and h>0.02:add('OBJECT_OVERLAP',spec,left['object_id'],{'other_object_id':right['object_id']},'QA-005','P0',{'intersection_inches':[x,y,w,h]},{'whitelist':'background/decoration/connector or explicit overlap_allowed'},'geometry')
  if title and body and min(title)<=max(body):add('TITLE_HIERARCHY_WEAK',spec,None,'Title does not exceed body font size','QA-008','P1',{'title_sizes':title,'body_sizes':body},{'rule':'title > body'},'style')
  ids=[s.name for s in slide.shapes];expected_titles=[e['object_id'] for e in elements if e['role']=='title'];others=[e['object_id'] for e in elements if e['role'] not in {'title','footer','background','decoration'}]
  if expected_titles and others and all(n in ids for n in expected_titles+others) and min(ids.index(n) for n in expected_titles)>min(ids.index(n) for n in others):add('READING_ORDER_TITLE_AFTER_BODY',spec,expected_titles[0],'Native z-order puts body before title','QA-018','P1',node='accessibility')
  checks.append({'slide_id':spec['slide_id'],'kind':'empty_object/geometry/collision/font/hierarchy/contrast/image/native_editability/table_headers/native_reading_order/deck_style','status':'executed'})
 stale=set(mapped)-known
 for oid in stale:issues.append(issue('OBJECT_MAP_STALE',None,'Object mapping has no specification','QA-016',object_id=oid))
 if len(fonts)>1:issues.append(issue('DECK_ACTUAL_FONT_INCONSISTENT',None,{'actual_fonts':sorted(fonts)},'QA-015','P1',action={'node':'style','mode':'local_revision','reason':'Native font families vary'}))
 return issues,checks,{'accessibility':'partial_native_alt_font_contrast_table_header_reading_order','global_consistency':'partial_style_font_footer_hierarchy','visual_geometry':'partial_collision_geometry_contrast_image_dpi','image_semantics':'not_executed','powerpoint_compatibility':'not_executed','full_semantic_reading_order':'not_executed'}
