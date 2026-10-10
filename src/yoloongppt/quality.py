"""SYS-015: independent structure/editability/source-number/render checks."""
import json
import hashlib
import re
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from PIL import Image
from pptx import Presentation

from .artifacts import entity
from .facts import approved_text
from .quality_rules import audit, package_audit, issue as measured_issue

NUMBER = re.compile(r'(?<![0-9A-Za-z_.])-?\d+(?:\.\d+)?%?')


def numbers(text):
    # Typed provenance IDs are metadata, not asserted quantities.
    text=re.sub(r'\b[a-z_]+_[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b','',text)
    return set(NUMBER.findall(text))



DIMENSION_NOTE = re.compile(r'(?<![0-9A-Za-z_.])(?P<width>\d+)\s*[×xX]\s*(?P<height>\d+)\s*(?P<unit>像素|pixels?|px|PNG|JPEG|JPG)(?![0-9A-Za-z_])',re.IGNORECASE)


def verified_asset_notes(notes, spec, native_slide):
    """Only verified same-page native image dimensions enter a metadata namespace.

    Mask the exact dimensions+unit span for numeric fact comparison; never add
    its values to the allowed-number set. Business quantities elsewhere remain.
    The original notes are kept byte-for-byte by caller/writer.
    """
    verified=[]
    images={}
    for element in spec['elements']:
        if element['type']!='image':continue
        data=element.get('data',{});asset_id=data.get('asset_id')
        if not asset_id or not data.get('sha256'):continue
        found=[shape for shape in native_slide.shapes if shape.name==element['object_id']]
        if len(found)!=1 or not hasattr(found[0],'image'):continue
        native=found[0].image;dimensions=tuple(native.size);digest=hashlib.sha256(native.blob).hexdigest()
        if digest!=data['sha256'] or dimensions!=(data.get('width_px'),data.get('height_px')):continue
        images.setdefault(asset_id,[]).append((element,digest,dimensions,native.ext.lower()))
    spans=[]
    for clause in re.finditer(r'[^。；;\n]+',notes):
        text=clause.group();ids=[aid for aid in images if re.search(r'(?<![A-Za-z0-9_])'+re.escape(aid)+r'(?![A-Za-z0-9_])',text)]
        # An ambiguous clause may refer to several assets; dimensions cannot be
        # assigned by numeric coincidence. Require one same-page registered ID.
        if len(ids)!=1 or len(images[ids[0]])!=1:continue
        element,digest,dimensions,ext=images[ids[0]][0]
        for match in DIMENSION_NOTE.finditer(text):
            width,height=int(match['width']),int(match['height']);unit=match['unit'].lower()
            if (width,height)!=dimensions:continue
            if unit in {'png','jpeg','jpg'} and unit not in ({'jpg','jpeg'} if ext in {'jpg','jpeg'} else {ext}):continue
            start,end=clause.start()+match.start(),clause.start()+match.end();spans.append((start,end))
            verified.append({'slide_id':spec['slide_id'],'object_id':element['object_id'],'kind':'asset_metadata_namespace',
                'check':'same-page native image bytes/sha256 and pixel dimensions','status':'executed','namespace':'native_image_dimensions',
                'asset_id':ids[0],'sha256':digest,'actual_dimensions_px':list(dimensions),'native_format':ext,
                'matched_span':notes[start:end],'span_start':start,'span_end':end,
                'reason':'Only this explicit verified dimension token is metadata, not a business fact allowance'})
    masked=list(notes)
    for start,end in spans:masked[start:end]=' '* (end-start)
    return ''.join(masked),verified

def check(deck, pptx, object_map, sources, renders, root):
    issues = []; checks = []
    def issue(code, slide_id, detail, severity='P0'):
        issues.append({'issue_id': entity('issue'), 'code': code, 'slide_id': slide_id, 'severity': severity, 'detail': detail})
    presentation = Presentation(pptx)
    if len(presentation.slides) != len(deck['slides']):issue('SLIDE_COUNT', None, '实际PPTX页数不同。')
    with ZipFile(pptx) as archive:
        if archive.testzip() is not None:issue('ZIP_CORRUPT', None, '压缩包CRC失败。')
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                try:ET.fromstring(archive.read(name))
                except ET.ParseError:issue('XML_INVALID', None, name)
        checks.append({'kind': 'structure', 'check': 'OOXML ZIP CRC and XML parsing', 'status': 'executed'})
        workbooks = [n for n in archive.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')]
    package_issues,package_checks=package_audit(pptx)
    issues.extend(package_issues);checks.extend(package_checks)
    rule_issues,rule_checks,rule_coverage=audit(deck,presentation,object_map)
    issues.extend(rule_issues);checks.extend(rule_checks)
    lookup = {e['evidence_id']: e for e in sources['evidence']}
    shape_lookup = {o['logical_object_id']: o for o in object_map['objects']}
    pages_text = (Path(root)/renders['render_text_path']).read_text(encoding='utf-8').split('\f')
    bbox_pages=None
    if renders.get('render_bbox_path'):
        bbox_pages=ET.parse(Path(root)/renders['render_bbox_path']).findall('.//{http://www.w3.org/1999/xhtml}page')
    for index, spec in enumerate(deck['slides']):
        if index >= len(presentation.slides):break
        slide = presentation.slides[index]
        refs = spec['source_refs']
        raw = '\n'.join(approved_text(sources,lookup[r]) for r in refs)
        content = spec['content']
        assumption_values=sources.get('fact_boundary',{}).get('assumption_values',{})
        aids=content.get('assumption_refs',[])
        if set(aids)-set(assumption_values):issue('ASSUMPTION_REFERENCE_UNKNOWN',spec['slide_id'],{'assumption_refs':aids})
        if aids and not any('【假设】' in t or '【推算】' in t or '【来源明确假设】' in t for t in content['body']):
            issue('ASSUMPTION_NOT_VISIBLE',spec['slide_id'],'假设/推算未在可见正文标明。')
        raw+='\n'+'\n'.join(assumption_values[a]['raw_text'] for a in aids if a in assumption_values)
        displayed = '\n'.join([content['title'], *content['body']])
        if content.get('table'):
            displayed += '\n' + '\n'.join(' '.join(row) for row in [content['table']['columns'], *content['table']['rows']])
        # New native visuals remain within the exact selected-fact boundary.
        for visual in spec['elements']:
            if visual['type']=='diagram':
                data=visual.get('data',{})
                displayed+='\n'+'\n'.join(str(n.get('text',n.get('label',''))) for n in data.get('nodes',[]))
            elif visual['type'] in {'shape','connector'}:
                displayed+='\n'+str(visual.get('data',{}).get('text',visual.get('data',{}).get('label','')))
            elif visual['type']=='image':
                displayed+='\n'+str(visual.get('data',{}).get('alt_text',''))
        new_numbers = numbers(displayed) - numbers(raw)
        if new_numbers:issue('FACT_NUMBER_UNSUPPORTED', spec['slide_id'], {'numbers_not_found_in_cited_evidence': sorted(new_numbers)})
        numeric_notes,asset_note_checks=verified_asset_notes(content['notes'],spec,slide)
        checks.extend(asset_note_checks)
        note_numbers=numbers(numeric_notes)-numbers(raw)
        # Auditing an explicitly rejected value in a negative clause is valid
        # provenance, never permission to put that value in presentation copy.
        rejected=set()
        for group in sources.get('evidence_resolution',{}).get('groups',[]):
            if group['conflicting']:
                chosen=next(c for c in group['claims'] if c['claim_id']==group['selected_claim_id'])
                for c in group['claims']:
                    if c['value']!=chosen['value']:rejected|=numbers(c['value']['raw_text'])
        audited=set()
        for clause in re.split('[。；\n]',numeric_notes):
            if re.search(r'不采用(?:冲突值)?|未采用(?:冲突值)?|已否决|已排除',clause):audited|=numbers(clause)&rejected
        unsupported_notes=note_numbers-audited
        if unsupported_notes:issue('FACT_NOTE_NUMBER_UNSUPPORTED',spec['slide_id'],{'numbers':sorted(unsupported_notes)})
        if content.get('chart'):
            for series in content['chart']['series']:
                for value in series['values']:
                    # Numeric identity only. It does not prove matching entity/unit.
                    candidates = {str(value), format(value, 'g')}
                    if not candidates & {n.rstrip('%') for n in numbers(raw)}:
                        issue('CHART_VALUE_UNSUPPORTED', spec['slide_id'], {'value': value})
        for element in spec['elements']:
            mapping = shape_lookup.get(element['object_id'])
            found = [s for s in slide.shapes if s.name == element['object_id']]
            if len(found) != 1 or mapping is None:
                issue('OBJECT_MAP_MISSING', spec['slide_id'], {'object_id': element['object_id']});continue
            shape = found[0]
            if shape.shape_id != mapping['shape_id'] or mapping['part'] != str(slide.part.partname):issue('OBJECT_MAP_MISMATCH', spec['slide_id'], element['object_id'])
            if shape.left < 0 or shape.top < 0 or shape.left+shape.width > presentation.slide_width+2 or shape.top+shape.height > presentation.slide_height+2:
                issue('OBJECT_OUT_OF_BOUNDS', spec['slide_id'], element['object_id'])
            if element['type'] == 'text':
                if not shape.has_text_frame or shape.text != '\n'.join(element['text']):issue('TEXT_NOT_EDITABLE_OR_CHANGED', spec['slide_id'], element['object_id'])
                if element['role'] != 'footer' and element['text']:
                    clean = lambda s: re.sub(r'\s+', '', s)
                    rendered = clean(pages_text[index]) if index < len(pages_text) else ''
                    if bbox_pages is not None:
                        # Filter actual PDF words to the owning native object's
                        # box before checking sequence; columns otherwise interleave.
                        x,y,w,h=element['bounds']
                        if index>=len(bbox_pages):issue('RENDER_BBOX_PAGE_MISSING',spec['slide_id'],element['object_id']);continue
                        page=bbox_pages[index]
                        words=[]
                        for word in page.findall('{http://www.w3.org/1999/xhtml}word'):
                            left=float(word.attrib['xMin']);right=float(word.attrib['xMax']);top=float(word.attrib['yMin']);bottom=float(word.attrib['yMax'])
                            cx=(left+right)/2;cy=(top+bottom)/2
                            if x*72-1<=cx<=(x+w)*72+1 and y*72-1<=cy<=(y+h)*72+1:words.append(word.text or '')
                        rendered=clean(''.join(words))
                    if any(clean(t) not in rendered for t in element['text']):issue('RENDER_TEXT_MISSING', spec['slide_id'], element['object_id'])
            elif element['type'] == 'table':
                rows = [element['data']['columns'], *element['data']['rows']]
                if not shape.has_table:issue('TABLE_NOT_NATIVE', spec['slide_id'], element['object_id'])
                elif [[c.text for c in row.cells] for row in shape.table.rows] != rows:issue('TABLE_DATA_CHANGED', spec['slide_id'], element['object_id'])
            elif element['type'] == 'chart':
                if not shape.has_chart:issue('CHART_NOT_NATIVE', spec['slide_id'], element['object_id'])
                else:
                    if [list(s.values) for s in shape.chart.series] != [s['values'] for s in element['data']['series']]:issue('CHART_DATA_CHANGED', spec['slide_id'], element['object_id'])
                    if not workbooks:issue('CHART_WORKBOOK_MISSING', spec['slide_id'], element['object_id'])
                    categories=[c.label for c in shape.chart.plots[0].categories]
                    if categories!=element['data']['categories']:issue('CHART_CATEGORIES_CHANGED',spec['slide_id'],element['object_id'])
                    if [s.name for s in shape.chart.series]!=[s['name'] for s in element['data']['series']]:issue('CHART_SERIES_NAMES_CHANGED',spec['slide_id'],element['object_id'])
                    units=element['data'].get('units','')
                    if units and (not shape.chart.value_axis.has_title or shape.chart.value_axis.axis_title.text_frame.text!=units):issue('CHART_UNITS_CHANGED',spec['slide_id'],element['object_id'])
            elif element['type'] in {'image','diagram','shape','connector'}:
                if element['type'] in {'image','connector'}:
                    props=shape._element.xpath('.//p:cNvPr');native_text=props[0].get('descr','') if props else ''
                elif element['type']=='shape':native_text=shape.text if shape.has_text_frame else ''
                else:
                    def child_text(shapes):
                        for native in shapes:
                            if native.has_text_frame:yield native.text
                            if hasattr(native,'shapes'):yield from child_text(native.shapes)
                    native_text='\n'.join(child_text(shape.shapes)) if hasattr(shape,'shapes') else ''
                unsupported=numbers(native_text)-numbers(raw)
                if unsupported:issue('VISUAL_FACT_NUMBER_UNSUPPORTED',spec['slide_id'],{'object_id':element['object_id'],'numbers':sorted(unsupported)})
        notes = slide.notes_slide.notes_text_frame.text
        if content['notes'] not in notes or any(r not in notes for r in refs):issue('NOTES_OR_REFERENCES_MISSING', spec['slide_id'], '备注/证据引用缺失。')
        checks.append({'slide_id': spec['slide_id'], 'kind': 'structure/editability/geometry/selected_source_numeric/visible_assumption/render_text', 'status': 'executed'})
    for artifact in renders['artifacts']:
        if artifact['type'] == 'png':
            with Image.open(Path(root)/artifact['path']) as image:
                colors = image.convert('RGB').resize((160, 90)).getcolors(14400)
                if colors is not None and len(colors) < 4:issue('RENDER_BLANK', artifact['slide_id'], '实际渲染缺少可见内容。')
    visible='\n'.join(t for slide in deck['slides'] for t in slide['content']['body'])
    for action in sources.get('fact_boundary',{}).get('actions',[]):
        if action['mode']=='placeholder' and action['label'] not in visible:
            issue('PLACEHOLDER_NOT_VISIBLE',None,{'fact_key':action['fact_key'],'label':action['label']})
    # Every legacy issue receives explicit measurement/threshold/action metadata.
    mapping={'OBJECT_OUT_OF_BOUNDS':('QA-004','geometry'),'RENDER_TEXT_MISSING':('QA-006','layout'),'FACT_NUMBER_UNSUPPORTED':('QA-014','fact'),'FACT_NOTE_NUMBER_UNSUPPORTED':('QA-014','fact'),'CHART_VALUE_UNSUPPORTED':('QA-011','fact'),'VISUAL_FACT_NUMBER_UNSUPPORTED':('QA-014','fact'),'TABLE_DATA_CHANGED':('QA-012','adapter')}
    for entry in issues:
        req,node=mapping.get(entry['code'],('QA-020','QA'))
        entry.setdefault('object_id',entry['detail'] if isinstance(entry['detail'],str) and entry['detail'] in shape_lookup else entry['detail'].get('object_id') if isinstance(entry['detail'],dict) else None)
        entry.setdefault('requirement_id',req);entry.setdefault('measurement',{'observed':entry['detail']});entry.setdefault('threshold',{'rule':entry['code']})
        entry.setdefault('action',{'node':node,'mode':'requires_user_input' if node=='fact' else 'local_revision' if node in {'geometry','layout'} else 'fail','reason':'Explicit mapped repair or fail; do not guess facts'})
    return {'qa_id': str(entity('artifact')), 'status': 'partial', 'issues': issues, 'checks': checks,
            'p0_issue_count': sum(i['severity'] == 'P0' for i in issues),
            'coverage': {'structure': 'partial', 'editability': 'partial', 'geometry': 'partial', 'fact_numeric': 'partial',
                         'visual_semantic': 'not_executed', 'fact_semantic': 'not_executed', 'accessibility': 'not_executed',
                         'powerpoint_compatibility': 'not_executed',**rule_coverage},
            'acceptance': {'AC-001': 'not_passed', 'GOV-008': 'not_passed'},
            'reason': 'Numeric membership is not semantic truth; PNG/text presence is not visual approval. Full DEC/QA/revision acceptance remains required.'}
