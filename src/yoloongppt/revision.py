"""SYS-016: selected text edit; unchanged OOXML parts are byte-preserved."""
import copy
import json
from pathlib import Path
from uuid import UUID
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from .artifacts import RunArtifacts, sha256
from .errors import TaskError
from .quality import check,numbers
from .facts import approved_text
from .rendering import render
from .review import review
from .writer import capacity

DECISION_FILES = ['fact-source-result.json','task-effective.json','decision-route.json','decision-context.json',
                  'presentation-context.json','decision-source-roles.json','source-role-map.json','fact-interpretation.json',
                  'decision-evidence.json','evidence-resolution.json','decision-fact-boundary.json','fact-boundary.json','approved-source-result.json']


def revision_sources(prior):
    """Revisions keep the parent's factual boundary, including QA-only recovery."""
    if (prior/'source-role-map.json').is_file() and not (prior/'fact-source-result.json').is_file():
        raise TaskError('FACT_SOURCE_SNAPSHOT_MISSING', '已判定角色的原版本缺少事实快照，不能退回全部来源。', 'RevisionEngine', ['DEC-003','SYS-016'])
    name='fact-source-result.json' if (prior/'fact-source-result.json').is_file() else 'source-result.json'
    result=json.loads((prior/name).read_text(encoding='utf-8'))
    if (prior/'source-role-map.json').is_file():
        roles=json.loads((prior/'source-role-map.json').read_text(encoding='utf-8'))
        if roles['status']!='assigned' or any(e['source_id'] not in roles['fact_source_ids'] for e in result['evidence']):
            raise TaskError('FACT_SOURCE_ROLE_MISMATCH', '事实快照与原角色判定不符，停止复审。', 'RevisionEngine', ['DEC-003','SYS-016'])
    if (prior/'decision-evidence.json').is_file():
        if not (prior/'approved-source-result.json').is_file():
            raise TaskError('FACT_BOUNDARY_SNAPSHOT_MISSING','原版本缺少已选事实快照，不能退回全部事实。','RevisionEngine',['DEC-004','DEC-005','SYS-016'])
        selected=json.loads((prior/'approved-source-result.json').read_text(encoding='utf-8'))
        if selected['evidence']!=result['evidence'] or selected['fact_boundary']['status']!='ready':
            raise TaskError('FACT_BOUNDARY_SNAPSHOT_INVALID','事实快照原始证据不一致或边界未决。','RevisionEngine',['DEC-004','DEC-005','SYS-016'])
        manifest=json.loads((prior/'manifest.json').read_text(encoding='utf-8'))
        for name in ['approved-source-result.json','fact-interpretation.json','evidence-resolution.json','fact-boundary.json']:
            entries=[a for a in manifest['artifacts'] if a['path']==name]
            if len(entries)!=1 or not (prior/name).is_file() or sha256(prior/name)!=entries[0]['hash']:
                raise TaskError('FACT_BOUNDARY_SNAPSHOT_CHANGED','已选事实快照哈希改变或缺失，停止修订/复审。','RevisionEngine',['DEC-004','DEC-005','SYS-016'])
        return selected
    return result


def revise(run_id, request, schemas, trace):
    schemas.validate('revision-request.schema.json', request)
    try:
        if str(UUID(run_id)) != run_id:raise ValueError()
    except ValueError:
        raise TaskError('REVISION_RUN_ID_INVALID', '原版本ID必须为标准UUID。', 'RevisionEngine', ['SYS-016']) from None
    prior = Path('/runtime/runs') / run_id
    names = ['task.json', 'deck-spec.json', 'source-result.json', 'object-map.json', 'quality-report.json', 'deck.pptx']
    if not all((prior/n).is_file() for n in names):
        raise TaskError('REVISION_BASE_INCOMPLETE', '原版本缺少实际产物、规格、对象映射或QA。', 'RevisionEngine', ['SYS-016'])
    load = lambda n: json.loads((prior/n).read_text(encoding='utf-8'))
    deck = load('deck-spec.json'); object_map = load('object-map.json'); sources = revision_sources(prior); task = load('task-effective.json') if (prior/'task-effective.json').is_file() else load('task.json')
    schemas.validate('deck-execution.schema.json', deck); schemas.validate('deck-execution.schema.json', object_map)
    candidates = [(s, e) for s in deck['slides'] for e in s['elements'] if e['object_id'] == request['object_id']]
    if len(candidates) != 1:
        raise TaskError('REVISION_OBJECT_NOT_FOUND', '没有唯一匹配的逻辑对象。', 'RevisionEngine', ['SYS-016'])
    spec, element = candidates[0]
    if element['type'] != 'text' or element['role'] not in {'title', 'body'}:
        raise TaskError('REVISION_OBJECT_UNSUPPORTED', '当前只支持标题/正文原生文本对象修订，其他对象未替换。', 'RevisionEngine', ['SYS-016'])
    if element['role'] == 'title' and (len(request['replacement_text']) != 1 or len(request['replacement_text'][0]) > 60):
        raise TaskError('REVISION_TITLE_INVALID', '标题需一个不超过60字的文本项。', 'RevisionEngine', ['SYS-016'])
    formatting=element.get('format',{});page_style=spec['style']
    capacity(request['replacement_text'], element['bounds'], element['font_size'],padding=formatting.get('margin',0.12),paragraph_spacing=formatting.get('paragraph_spacing',10))
    old_text = copy.deepcopy(element['text'])
    element['text'] = request['replacement_text']
    if element['role'] == 'title':spec['content']['title'] = request['replacement_text'][0]
    else:spec['content']['body'] = request['replacement_text']
    if 'fact_boundary' in sources:
        known={e['evidence_id']:e for e in sources['evidence']}
        raw='\n'.join(approved_text(sources,known[r]) for r in spec['source_refs'])
        aids=spec['content'].get('assumption_refs',[]);values=sources['fact_boundary']['assumption_values']
        if set(aids)-set(values):raise TaskError('REVISION_ASSUMPTION_UNKNOWN','修订引用未知假设。','RevisionEngine',['DEC-005','SYS-016'])
        raw+='\n'+'\n'.join(values[a]['raw_text'] for a in aids)
        unsupported=numbers('\n'.join([spec['content']['title'],*spec['content']['body']]))-numbers(raw)
        if unsupported:raise TaskError('REVISION_FACT_UNSUPPORTED','修订试图使用未选/无来源数值，未写入PPT。','RevisionEngine',['DEC-004','DEC-005','SYS-016'],[{'numbers':sorted(unsupported)}])
    schemas.validate('deck-execution.schema.json', deck)
    mapping = next(o for o in object_map['objects'] if o['logical_object_id'] == request['object_id'])
    presentation = Presentation(prior/'deck.pptx')
    slide = presentation.slides[spec['order']]
    matches = [s for s in slide.shapes if s.shape_id == mapping['shape_id'] and s.name == mapping['name']]
    if len(matches) != 1:
        raise TaskError('REVISION_OBJECT_MAP_STALE', '实际对象与映射不一致，停止修订。', 'RevisionEngine', ['SYS-016'])
    shape = matches[0]; shape.text_frame.clear()
    for n, text in enumerate(request['replacement_text']):
        p = shape.text_frame.paragraphs[0] if n == 0 else shape.text_frame.add_paragraph()
        p.text = text; p.space_after = Pt(formatting.get('paragraph_spacing',10)); p.line_spacing = 1.15
        p.font.name = page_style['font']; p.font.size = Pt(element['font_size']); p.font.bold = formatting.get('bold',element['role'] == 'title')
        p.font.color.rgb = RGBColor.from_string(formatting.get('color',page_style['accent'] if element['role'] == 'title' else page_style['foreground']))
        default = p._p.get_or_add_pPr().find('{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr')
        if default is not None:
            ea = OxmlElement('a:ea'); ea.set('typeface', page_style['font']); default.append(ea)
    artifacts = RunArtifacts()
    artifacts.json('revision-request.json', request)
    expanded = artifacts.path/'revision-expanded.pptx'
    presentation.save(expanded)
    selected_part = mapping['part'].lstrip('/')
    with ZipFile(prior/'deck.pptx') as original, ZipFile(expanded) as edited, ZipFile(artifacts.path/'deck.pptx', 'w') as output:
        for info in original.infolist():
            output.writestr(info, edited.read(info.filename) if info.filename == selected_part else original.read(info.filename))
    expanded.unlink()
    with ZipFile(prior/'deck.pptx') as original, ZipFile(artifacts.path/'deck.pptx') as edited:
        changed = [n for n in original.namelist() if original.read(n) != edited.read(n)]
        if changed != [selected_part]:
            raise TaskError('REVISION_SCOPE_CHANGED', '修订没有保持指定part边界。', 'RevisionEngine', ['SYS-016'])
    for name, value in [('task.json', task), ('deck-spec.json', deck), ('source-result.json', load('source-result.json')), ('object-map.json', object_map)]:artifacts.json(name, value)
    for name in DECISION_FILES:
        if (prior/name).is_file():artifacts.json(name,load(name))
    renders = render(artifacts.path/'deck.pptx', deck, artifacts.path)
    artifacts.json('render-report.json', renders)
    quality = check(deck, artifacts.path/'deck.pptx', object_map, sources, renders, artifacts.path)
    fresh_review = review(task, deck, sources, renders, artifacts, [spec['order']])
    previous_quality = load('quality-report.json')
    old_review = previous_quality.get('model_review')
    if old_review:
        carried = [i for i in old_review['issues'] if i['slide_order'] != spec['order']]
        combined = {'checked_slide_orders': old_review['checked_slide_orders'], 'issues': [*carried, *fresh_review['issues']],
                    'summary': '修订页实际重审；其余页复用未改内容及原审查。',
                    'fresh_checked_slide_orders': fresh_review['checked_slide_orders'], 'carried_from_run': run_id}
    else:
        combined = fresh_review
    quality['model_review'] = combined
    quality['p0_issue_count'] += sum(i['severity'] == 'P0' for i in combined['issues'])
    quality['coverage']['visual_semantic'] = 'model_review_executed_with_unchanged_page_evidence_reuse'
    quality['coverage']['fact_semantic'] = 'model_review_executed_with_unchanged_page_evidence_reuse'
    artifacts.json('quality-report.json', quality)
    history = {'parent_run_id': run_id, 'trace_id': trace, 'object_id': request['object_id'], 'slide_id': spec['slide_id'],
               'before_text': old_text, 'after_text': request['replacement_text'], 'reason': request['reason'],
               'before_pptx_hash': sha256(prior/'deck.pptx'), 'after_pptx_hash': sha256(artifacts.path/'deck.pptx'),
               'changed_parts': changed, 'unchanged_parts_byte_preserved': True,
               'entity_ids_preserved': True, 'rerun_nodes': ['selected_text_object', 'render', 'QA'],
               'scope': 'single title/body text; complete RevisionPlan/DEC rerun, other objects, locks and external decks still pending'}
    artifacts.json('revision-history.json', history)
    result = {'ok': True, 'trace_id': trace, 'operation': 'revise', 'run_id': artifacts.run_id, 'parent_run_id': run_id,
              'state': 'draft_revised', 'artifact_root': str(artifacts.path), 'pptx': str(artifacts.path/'deck.pptx'),
              'changed_parts': changed, 'qa': {'p0_issue_count': quality['p0_issue_count'], 'status': 'partial'}, 'system_acceptance': 'not_passed'}
    artifacts.json('result.json', result)
    artifacts.json('manifest.json', {'run_id': artifacts.run_id, 'artifacts': artifacts.manifest()})
    return result


def recheck(run_id, schemas, trace):
    """QA-only recovery: reuse the exact deck/render bytes, preserve failed review."""
    import shutil
    try:
        if str(UUID(run_id)) != run_id:raise ValueError()
    except ValueError:
        raise TaskError('REVISION_RUN_ID_INVALID', '原版本ID必须为标准UUID。', 'RevisionEngine', ['SYS-016']) from None
    prior = Path('/runtime/runs')/run_id
    required = ['task.json', 'deck-spec.json', 'source-result.json', 'object-map.json', 'quality-report.json', 'render-report.json', 'deck.pptx']
    if not all((prior/n).is_file() for n in required):
        raise TaskError('QA_CHECKPOINT_INCOMPLETE', 'QA断点缺少实际文件。', 'QAEngine', ['SYS-015'])
    artifacts = RunArtifacts()
    for name in required:shutil.copyfile(prior/name, artifacts.path/name)
    shutil.copytree(prior/'render', artifacts.path/'render')
    for name in ['revision-history.json', 'revision-request.json',*DECISION_FILES]:
        if (prior/name).is_file():shutil.copyfile(prior/name, artifacts.path/name)
    load = lambda n: json.loads((prior/n).read_text(encoding='utf-8'))
    deck = load('deck-spec.json'); task = load('task-effective.json') if (prior/'task-effective.json').is_file() else load('task.json'); sources = revision_sources(prior); renders = load('render-report.json')
    schemas.validate('deck-execution.schema.json', deck)
    for a in renders['artifacts']:
        if sha256(artifacts.path/a['path']) != a['hash']:
            raise TaskError('QA_RENDER_HASH_CHANGED', '复用渲染哈希不符，不能直接复审。', 'QAEngine', ['SYS-015'])
    previous_quality = load('quality-report.json')
    artifacts.json('quality-before-recheck.json', previous_quality)
    orders = previous_quality.get('model_review', {}).get('fresh_checked_slide_orders')
    fresh = review(task, deck, sources, renders, artifacts, orders)
    quality = check(deck, artifacts.path/'deck.pptx', load('object-map.json'), sources, renders, artifacts.path)
    carried = [i for i in previous_quality.get('model_review', {}).get('issues', []) if orders is not None and i['slide_order'] not in orders]
    quality['model_review'] = {'checked_slide_orders': list(range(len(deck['slides']))), 'fresh_checked_slide_orders': fresh['checked_slide_orders'],
                               'issues': [*carried, *fresh['issues']], 'summary': fresh['summary'], 'carried_from_run': run_id}
    quality['p0_issue_count'] += sum(i['severity'] == 'P0' for i in quality['model_review']['issues'])
    quality['coverage']['visual_semantic'] = quality['coverage']['fact_semantic'] = 'actual_model_review_with_unchanged_evidence_reuse'
    artifacts.json('quality-report.json', quality)
    result = {'ok': True, 'operation': 'recheck', 'trace_id': trace, 'run_id': artifacts.run_id, 'parent_run_id': run_id,
              'artifact_root': str(artifacts.path), 'pptx': str(artifacts.path/'deck.pptx'), 'deck_and_render_bytes_reused': True,
              'qa': {'p0_issue_count': quality['p0_issue_count'], 'status': 'partial'}, 'system_acceptance': 'not_passed'}
    artifacts.json('result.json', result)
    artifacts.json('manifest.json', {'run_id': artifacts.run_id, 'artifacts': artifacts.manifest()})
    return result
