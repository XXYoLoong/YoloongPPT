"""PP-05 / SYS-012/013: execute native OOXML objects and record real backend IDs."""
import math
import time
from pathlib import Path
from importlib.metadata import version

from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .errors import TaskError
from .atomic import AtomicRegistry

FONT_FILE = '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'


def capacity(texts, bounds, size, padding=0.12, paragraph_spacing=10):
    """Conservative CJK glyph measurement; actual rendered QA remains necessary."""
    font = ImageFont.truetype(FONT_FILE, int(size * 4 / 3))
    width = (bounds[2] - padding * 2) * 96
    height = (bounds[3] - padding * 2) * 96
    lines = 0
    for text in texts:
        current = 0
        lines += 1
        for char in text:
            if char == '\n':
                current = 0; lines += 1
                continue
            advance = font.getlength(char)
            if current + advance > width:
                lines += 1; current = advance
            else:
                current += advance
    used = lines * size * 4 / 3 * 1.25 + max(0, len(texts)-1) * paragraph_spacing * 4/3
    if used > height:
        raise TaskError('TEXT_CAPACITY_EXCEEDED', '文字超出已测量槽位容量；未缩小字号、截断或改为图片。',
                        'PageExecutor', ['SYS-012'], [{'estimated_height_px': round(used), 'available_height_px': round(height)}])
    return {'estimated_lines': lines, 'estimated_height_px': round(used), 'available_height_px': round(height), 'measurement': 'Pillow Noto glyph advances; conservative, not visual QA'}


def text_shape(slide, element, style):
    formatting=element.get('format',{})
    measurement = capacity(element['text'], element['bounds'], element['font_size'],padding=formatting.get('margin',0.12),paragraph_spacing=formatting.get('paragraph_spacing',10))
    shape = slide.shapes.add_textbox(*(Inches(x) for x in element['bounds']))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Inches(formatting.get('margin',0.12))
    frame.margin_top = frame.margin_bottom = Inches(formatting.get('margin',0.12))
    if 'fill' in formatting:shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(formatting['fill'])
    if 'border' in formatting:shape.line.color.rgb=RGBColor.from_string(formatting['border']);shape.line.width=Pt(0.7)
    else:shape.line.fill.background()
    for n, text in enumerate(element['text']):
        p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
        p.text = text
        p.space_after = Pt(formatting.get('paragraph_spacing',10))
        p.line_spacing = 1.15
        p.font.name = style['font']; p.font.size = Pt(element['font_size'])
        p.font.bold = formatting.get('bold',element['role'] == 'title')
        p.font.color.rgb = RGBColor.from_string(formatting.get('color',style['accent'] if element['role'] == 'title' else style['foreground']))
        # Explicit East Asian typeface; do not rely on the host's theme fallback.
        props = p._p.get_or_add_pPr()
        default = props.find('{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr')
        if default is not None:
            ea = OxmlElement('a:ea'); ea.set('typeface', style['font']); default.append(ea)
    return shape, measurement


def table_shape(slide, element, style):
    data = element['data']; rows = [data['columns'], *data['rows']]
    count = len(rows); cols = len(data['columns'])
    for row in rows:
        for text in row:
            capacity([text], [0, 0, element['bounds'][2]/cols, element['bounds'][3]/count], 16, padding=0.07, paragraph_spacing=0)
    shape = slide.shapes.add_table(count, cols, *(Inches(v) for v in element['bounds']))
    table = shape.table
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c); cell.text = text
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.07)
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(style['accent'] if r == 0 else ('FFFFFF' if r % 2 else 'E7ECE9'))
            for p in cell.text_frame.paragraphs:
                p.font.name = style['font']; p.font.size = Pt(16); p.font.bold = r == 0
                p.font.color.rgb = RGBColor.from_string('FFFFFF' if r == 0 else style['foreground'])
    return shape, {'rows': count, 'columns': cols, 'native_table': True}


def chart_shape(slide, element, style):
    spec = element['data']; data = CategoryChartData()
    data.categories = spec['categories']
    for series in spec['series']:
        if any(not math.isfinite(v) for v in series['values']):
            raise TaskError('CHART_NONFINITE_DATA', '图表包含非有限数值。', 'PageExecutor', ['SYS-012'])
        data.add_series(series['name'], series['values'])
    kind = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'line': XL_CHART_TYPE.LINE}[spec['type']]
    shape = slide.shapes.add_chart(kind, *(Inches(v) for v in element['bounds']), data)
    chart = shape.chart
    chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False; chart.legend.font.size = Pt(16); chart.legend.font.name = style['font']
    chart.category_axis.tick_labels.font.size = Pt(16); chart.category_axis.tick_labels.font.name = style['font']
    chart.value_axis.tick_labels.font.size = Pt(16)
    chart.value_axis.has_title = bool(spec['units'])
    if spec['units']:
        frame = chart.value_axis.axis_title.text_frame; frame.text = spec['units']
        frame.paragraphs[0].font.size = Pt(16); frame.paragraphs[0].font.name = style['font']
    palette = [style['accent'], 'D17D32', '5669A4', 'A34567']
    for index, series in enumerate(chart.series):
        series.format.line.color.rgb = RGBColor.from_string(palette[index])
        if spec['type'] != 'line':
            series.format.fill.solid(); series.format.fill.fore_color.rgb = RGBColor.from_string(palette[index])
    return shape, {'native_chart': True, 'embedded_workbook': True, 'series': len(spec['series']), 'categories': len(spec['categories']),
                   'series_colors': palette[:len(spec['series'])]}


def execute(deck, execution, destination, schemas):
    schemas.validate('deck-execution.schema.json', deck)
    schemas.validate('deck-execution.schema.json', execution)
    registry = AtomicRegistry(schemas)
    # Validate every binding/input before beginning writes, including later calls.
    for call in execution['calls']:
        registry.check_call(call)
    presentation = Presentation()
    presentation.slide_width = Inches(deck['width_inches']); presentation.slide_height = Inches(deck['height_inches'])
    presentation.core_properties.title = deck['title']
    presentation.core_properties.author = 'YoloongPPT'
    slides = {}; completed = set(); outputs = []; object_map = {'objects': []}
    specs = {s['slide_id']: s for s in deck['slides']}
    for call in execution['calls']:
        if set(call['deps']) - completed or call['call_id'] in completed:
            raise TaskError('EXECUTION_DAG_INVALID', '执行依赖未满足或调用ID重复。', 'PageExecutor', ['SYS-011', 'SYS-012'])
        begin = time.monotonic(); spec = specs[call['slide_id']]
        if call['capability'] == 'create_slide':
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(spec['style']['background'])
            slides[call['slide_id']] = slide
            observed = {'native_slide_id': slide.slide_id, 'part': str(slide.part.partname), 'order': spec['order']}
        elif call['capability'] == 'add_notes':
            slide = slides[call['slide_id']]
            slide.notes_slide.notes_text_frame.text = call['inputs']['text'] + '\n\n证据：\n' + '\n'.join(call['inputs']['evidence_refs'])
            observed = {'part': str(slide.notes_slide.part.partname), 'notes_saved': True}
        else:
            element = call['inputs']; slide = slides[call['slide_id']]
            function = {'add_text': text_shape, 'add_table': table_shape, 'add_chart': chart_shape}.get(call['capability'])
            if function is None:
                raise TaskError('EXECUTION_CAPABILITY_UNSUPPORTED', '执行能力未实现。', 'PageExecutor', ['SYS-012'])
            shape, observed = function(slide, element, spec['style'])
            shape.name = element['object_id']
            object_map['objects'].append({'logical_object_id': element['object_id'], 'slide_id': call['slide_id'],
                'backend_object_id': f'{slide.part.partname}#{shape.shape_id}', 'shape_id': shape.shape_id, 'name': shape.name,
                'type': element['type'], 'source_ref': element['source_refs'], 'execution_call_id': call['call_id'], 'part': str(slide.part.partname)})
            observed.update(shape_id=shape.shape_id, name=shape.name)
        completed.add(call['call_id'])
        registry.check_call(call, observed)
        outputs.append({'call_id': call['call_id'], 'capability_id': call['capability_id'], 'implementation_id': call['implementation_id'], 'backend': 'PP-05',
                        'input': call['inputs'], 'output': observed, 'warnings': [], 'duration_seconds': round(time.monotonic()-begin, 4), 'status': 'executed'})
    schemas.validate('deck-execution.schema.json', object_map)
    temporary = Path(destination).with_suffix('.tmp.pptx')
    presentation.save(temporary)
    temporary.replace(destination)
    return object_map, {'library': 'python-pptx', 'version': version('python-pptx'), 'calls': outputs,
                        'limitations': ['new-deck text/table/category-chart/notes only; no imported-template preservation, advanced objects or live editing']}
