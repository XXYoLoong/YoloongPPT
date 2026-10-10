"""PP-05 native primitives and an honest PPT-001–030 capability matrix."""
import math
from pathlib import Path
from urllib.parse import urlparse

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.util import Inches, Pt

from .errors import TaskError


def capabilities():
    implemented = {
        1: 'create/open/save-as; properties and dimensions; merge not implemented',
        2: 'add/read slides; hide/duplicate/move/delete not implemented',
        7: 'read placeholder type/idx/bounds; fill existing text placeholders',
        8: 'create/read/update text frames; preserve unrelated XML',
        9: 'font/size/bold/italic/underline/color/alignment; advanced bullets/tabs not implemented',
        11: 'rectangle/ellipse/rounded rectangle/arrows; fill/stroke/rotation',
        12: 'straight/elbow/curved connectors; connection points/reroute/freeform not implemented',
        13: 'native group creation and child indexing; ungroup/locks not implemented',
        14: 'embedded PNG/JPEG create/read/replace; crop/effects not implemented; shared image replacement rejected',
        16: 'native rows/columns/cells and cell text edit; pagination not implemented',
        17: 'category charts and replace data with embedded workbook; combo/error bars not implemented',
        22: 'safe https/http/mailto text hyperlink; executable file/action links rejected',
        23: 'read/update existing notes; notes master preserved',
        29: 'alt text and shape order indexing; complete accessibility audit not implemented',
        30: 'slide-id + shape-id stable mapping; extension/custom XML preserved',
    }
    return {'backend': 'PP-05', 'library': 'python-pptx', 'version': '1.0.2',
            'system_acceptance': 'not_passed', 'requirements': [
                {'requirement_id': f'PPT-{i:03}', 'status': 'partial' if i in implemented else 'unsupported',
                 'operations': {'create': 'partial' if i in {1,2,8,9,11,12,13,14,16,17,22} else 'unsupported',
                                'read': 'partial', 'update': 'partial' if i in implemented else 'unsupported',
                                'delete': 'unsupported', 'preserve': 'supported', 'render': 'untested'},
                 'reason': implemented.get(i, 'Existing OOXML parts are inventoried and preserved; native editing is unavailable.'),
                 'preservation_scope': 'unchanged ZIP member payloads in existing_deck.apply_patch'}
                for i in range(1, 31)]}


def _bounds(spec):
    b = spec.get('bounds')
    if not isinstance(b, list) or len(b) != 4 or any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in b) or min(b[:2]) < 0 or min(b[2:]) <= 0:
        raise TaskError('NATIVE_GEOMETRY_INVALID', '对象几何须为四个有限英寸值和正尺寸。', 'NativeObjects', ['PPT-011'])
    return [Inches(v) for v in b]


def format_shape(shape, style):
    """Apply explicitly supplied styles; no guessed substitution or autofit."""
    allowed = {'font', 'font_size', 'bold', 'italic', 'underline', 'color', 'fill', 'stroke', 'stroke_width', 'rotation'}
    if set(style) - allowed:
        raise TaskError('NATIVE_FORMAT_UNSUPPORTED', '包含未实现的格式字段。', 'NativeObjects', ['PPT-009', 'PPT-028'])
    for key in ('color', 'fill', 'stroke'):
        if key in style and (not isinstance(style[key], str) or len(style[key]) != 6 or any(c not in '0123456789abcdefABCDEF' for c in style[key])):
            raise TaskError('NATIVE_COLOR_INVALID', '颜色必须为六位 RGB。', 'NativeObjects', ['PPT-009'])
    if 'fill' in style:
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(style['fill'])
    if 'stroke' in style:
        shape.line.color.rgb = RGBColor.from_string(style['stroke'])
    if 'stroke_width' in style:
        shape.line.width = Pt(style['stroke_width'])
    if 'rotation' in style:
        shape.rotation = style['rotation']
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                for key, target in [('font','name'),('bold','bold'),('italic','italic'),('underline','underline')]:
                    if key in style: setattr(run.font, target, style[key])
                if 'font_size' in style: run.font.size = Pt(style['font_size'])
                if 'color' in style: run.font.color.rgb = RGBColor.from_string(style['color'])


def chart_data(value):
    if not isinstance(value, dict) or set(value) != {'categories', 'series'} or not value['categories'] or not value['series']:
        raise TaskError('NATIVE_CHART_INVALID', '图表需要类别与系列。', 'NativeObjects', ['PPT-017'])
    data = CategoryChartData(); data.categories = value['categories']
    for series in value['series']:
        if set(series) != {'name', 'values'} or len(series['values']) != len(value['categories']) or any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in series['values']):
            raise TaskError('NATIVE_CHART_INVALID', '系列数值须有限且与类别长度一致。', 'NativeObjects', ['PPT-017'])
        data.add_series(series['name'], series['values'])
    return data


def add_object(slide, spec, _validated=False):
    """Consume a native spec into an existing slide; callers own artifact storage."""
    if not _validated:
        from .schemas import SchemaRegistry
        SchemaRegistry().validate('native-object-spec.schema.json', spec)
    kind = spec.get('kind'); bounds = _bounds(spec)
    if kind == 'text':
        shape = slide.shapes.add_textbox(*bounds); shape.text = spec.get('text', '')
    elif kind == 'shape':
        types = {'rectangle': MSO_AUTO_SHAPE_TYPE.RECTANGLE, 'ellipse': MSO_AUTO_SHAPE_TYPE.OVAL,
                 'rounded_rectangle': MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 'right_arrow': MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW}
        if spec.get('shape_type', 'rectangle') not in types:
            raise TaskError('NATIVE_SHAPE_UNSUPPORTED', '形状未登记；没有替换。', 'NativeObjects', ['PPT-011'])
        shape = slide.shapes.add_shape(types[spec.get('shape_type', 'rectangle')], *bounds)
        shape.text = spec.get('text', '')
    elif kind == 'connector':
        types = {'straight': MSO_CONNECTOR.STRAIGHT, 'elbow': MSO_CONNECTOR.ELBOW, 'curve': MSO_CONNECTOR.CURVE}
        if spec.get('connector_type', 'straight') not in types:
            raise TaskError('NATIVE_CONNECTOR_UNSUPPORTED', '连接线类型未登记。', 'NativeObjects', ['PPT-012'])
        x,y,w,h = bounds; shape = slide.shapes.add_connector(types[spec.get('connector_type', 'straight')], x,y,x+w,y+h)
    elif kind == 'image':
        from PIL import Image
        image_path=local_asset(spec['path'])
        with Image.open(image_path) as image:
            if image.format not in {'PNG','JPEG'} or image.width*image.height>40_000_000:
                raise TaskError('NATIVE_IMAGE_UNSUPPORTED','当前嵌入图片仅支持不超过 4000 万像素的 PNG/JPEG。','NativeObjects',['PPT-014'])
            image.verify()
        shape = slide.shapes.add_picture(str(image_path), *bounds)
    elif kind == 'table':
        values = spec.get('cells', [])
        if not values or not values[0] or any(len(row) != len(values[0]) for row in values):
            raise TaskError('NATIVE_TABLE_INVALID', '表格须为非空矩形。', 'NativeObjects', ['PPT-016'])
        shape = slide.shapes.add_table(len(values), len(values[0]), *bounds)
        for r,row in enumerate(values):
            for c,text in enumerate(row): shape.table.cell(r,c).text = str(text)
    elif kind == 'chart':
        types = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED,
                 'line': XL_CHART_TYPE.LINE, 'pie': XL_CHART_TYPE.PIE, 'area': XL_CHART_TYPE.AREA}
        if spec.get('chart_type', 'column') not in types:
            raise TaskError('NATIVE_CHART_UNSUPPORTED', '图表类型未登记；未回退成图片。', 'NativeObjects', ['PPT-017'])
        shape = slide.shapes.add_chart(types[spec.get('chart_type', 'column')], *bounds, chart_data(spec['data']))
    elif kind == 'group':
        shape = slide.shapes.add_group_shape()
        for child in spec.get('children', []):
            child_shape = add_object(slide, child, _validated=True)
            shape.shapes._spTree.insert_element_before(child_shape._element, 'p:extLst')
        shape.shapes._recalculate_extents()
    else:
        raise TaskError('NATIVE_OBJECT_UNSUPPORTED', '对象类型未实现；没有静默降级。', 'NativeObjects', ['PPT-001'])
    if 'name' in spec: shape.name = spec['name']
    if 'alt_text' in spec:
        shape._element.xpath('.//p:cNvPr')[0].set('descr', spec['alt_text'])
    if 'hyperlink' in spec:
        url=urlparse(spec['hyperlink'])
        if url.scheme not in {'https','http','mailto'} or (url.scheme in {'http','https'} and (not url.hostname or url.username is not None or url.password is not None)) or (url.scheme=='mailto' and not url.path):
            raise TaskError('NATIVE_LINK_UNSAFE','只允许显式 http/https/mailto 链接，禁止执行文件和含凭据 URL。','NativeObjects',['PPT-022'])
        if not shape.has_text_frame:
            raise TaskError('NATIVE_LINK_UNSUPPORTED','该对象没有文字链接接口。','NativeObjects',['PPT-022'])
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs: run.hyperlink.address=spec['hyperlink']
    format_shape(shape, spec.get('style', {}))
    return shape
def local_asset(path):
    """Container-only scoped media; secrets and symlink escapes are rejected."""
    value = Path(path).resolve()
    roots = [Path('/workspace').resolve(), Path('/runtime').resolve()]
    relative = next((value.relative_to(root) for root in roots if value.is_relative_to(root)), None)
    if relative is None or any(x.lower() in {'secrets', '.git', '.ssh', '.env'} or x.lower().startswith('.env.') for x in relative.parts):
        raise TaskError('NATIVE_ASSET_PATH_FORBIDDEN', '媒体文件须位于工作区或运行目录，禁止读取敏感目录。', 'NativeObjects', ['SEC-001', 'PPT-014'])
    if not value.is_file():
        raise TaskError('NATIVE_ASSET_NOT_FOUND', '媒体文件不存在。', 'NativeObjects', ['PPT-014'])
    if value.stat().st_size > 32*1024*1024:
        raise TaskError('IMAGE_SIZE_LIMIT', '媒体文件超过 32 MiB 显式限制。', 'NativeObjects', ['PPT-014'])
    return value
